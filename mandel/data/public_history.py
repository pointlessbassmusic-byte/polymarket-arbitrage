"""
Public and sponsor histories for the batch-6c underlying claims (source group "misc").
=====================================================================================

One module for the free sources that are neither Yahoo/FRED nor crypto venues nor prediction markets. Every source was
probed from this host on 2026-10-15 (probe files under the session scratchpad ``triage6c/`` and ``fetch6c/misc/probes``).
Each section names the endpoint, the theories that need it, the licence decision and the ``available_at`` rule. Rows are
archive rows {source, kind, key, observed_at, available_at, value}: ``observed_at`` is the time the value refers to,
``available_at`` the earliest time a market participant could know it (never later than the fetch time).

  1. CFTC Commitments of Traders (T-405), ``publicreporting.cftc.gov`` Socrata: disaggregated F&O combined
     ``kh3c-gbw2`` (the rule's report) and futures-only ``72hh-3qpy`` for gold 088691 and silver 084691 (2006-06..),
     legacy futures-only ``6dca-aqww`` and combined ``jun7-fc8e`` (gold/silver from 1986: the long variant) and TFF
     ``yw9f-hn96`` / ``gpe5-46if`` for CME bitcoin 133741 (2018-04..; BTC is not in the disaggregated report). Positions
     are as of Tuesday; the report is released at 15:30 ET, usually on Friday. ``available_at``: Socrata's row
     ``:created_at`` (the live publication, verified for every report since 2022-07: Fri 19:30/20:30 UTC = 15:30 ET, the
     2023 ION delays and the 2025 shutdown catch-up dates), never earlier than the regular Friday 15:30 ET; rows from the
     2022 bulk loads use the CFTC publication schedule of their era (``cot_release_rule``): monthly until 1990 (end of
     the 12th of the following month; the semi-monthly rows of 1986-89 had no real-time mid-month publication: pit
     ``not_realtime``), 1990 the same bound, 1991-92 the 6th federal business day, 1992-10..2000 biweekly (two weekly
     reports every other Friday: the following week's release day), from 2001 Friday 15:30 ET (the next week's Friday
     for a Thursday/Friday as-of date, one business day later for each federal closure on the Wednesday-Friday of the
     release week, never before the 3rd federal business day), the 2013 and 2019 catch-up schedules (CFTC 6745-13 and
     notices; 7864-19) and a conservative bound for the 1995-96 shutdowns. Backcasts: disaggregated rows before
     2009-09-01 (released as history on 2009-10-20) and TFF rows before 2010-07-20 (history of December 2010) are
     available at that release (pit ``backcast``; the schedule time stays in ``release_rule``). Gold and silver were in
     the first disaggregated report (2009-09-01); CME bitcoin is in the TFF report from its first row (2018-04-10).
     available_at is never before observed_at (guard). US federal data, public domain: in git.
  2. Treasury auctions (T-442), ``api.fiscaldata.treasury.gov`` ``auctions_query`` for every Note and Bond since
     1979 (nominal, TIPS and FRN rows; filter ``inflation_index_security`` = No and ``floating_rate`` = No for the
     nominal 10/30-year set including the 9-Year x-Month / 29-Year x-Month reopenings). Two kinds per auction:
     ``auction-announcement`` (schedule fields only; available at the end of the announcement day) and
     ``auction-result`` (the full raw row; available 30 minutes after the competitive close, 13:00 ET for coupons;
     announced auctions not yet held have no result row). US Treasury data, public domain: in git.
  3. Cleveland Fed inflation nowcasts (T-360): ``nowcast_month.json`` (month-over-month CPI, core CPI, PCE, core PCE
     per target month since 2013-07, one value per business day) and ``nowcast_quarter.json``. "Updated every business
     day around 10:00 a.m. Eastern time" (the page's FAQ), so a value labelled D is available at D 12:00 ET. The file
     is regenerated daily and past daily values cannot be proven to be real-time vintages; ``fetch_clevelandfed_snapshot``
     stores the recent target months as a dated snapshot so forward vintages accumulate. The ``Actual ...`` series
     (the BLS/BEA print, plotted at the release-day label) are available at that label's 08:30 ET. Federal Reserve
     Bank of Cleveland content (a Reserve Bank, not a federal agency): local only.
  4. SPDR Gold Trust archive (T-404, T-370): ``api.spdrgoldshares.com/api/v1/historical-archive?product=gld|gldm``
     (xlsx; the old ``GLD_US_archive_EN.csv`` now serves a PDF). Daily ounces per share, NAV per share, tonnes, total
     ounces, closing price, volume since 2004-11-18; "US Holiday" rows dropped. available_at = next calendar day 09:30 ET
     (t+1, the claim's rule; the file is rewritten about 02:10 UTC after each session). The file's disclaimer forbids
     redistribution: local only.
  5. Sprott physical trusts (T-472): the fund pages' embedded Highcharts series 'Share Price', 'Net Asset Value',
     'Premium/Discount to Net Asset Value' (``[Date.UTC(y, m0, d), v]``) for PHYS (gold, 2010-02..), PSLV, SPPP and
     CEF. NAV is struck after the close: available_at = next calendar day 09:30 ET. Dates include TSX-only sessions
     (align to the NYSE calendar). Sprott website terms: local only.
  6. SSGA SPY daily holdings (T-374 alternative universe): ``holdings-daily-us-en-spy.xlsx``, a dated snapshot
     (``ssga-spy-holdings-<as-of>``): available_at = the file's Last-Modified (not before the as-of close). Local only.
  7. Wikipedia S&P 500 tables (T-374, T-394, T-473, T-334): the ``constituents`` table of "List of S&P 500
     companies" (symbol, GICS sector and sub-industry, date added, CIK) and the ``changes`` table of "Historical
     components of the S&P 500" (effective date, added / removed ticker and security, reason, the cited press
     releases). The table's date is the EFFECTIVE date (changes take effect before that day's open); the announcement
     is earlier: ``announced`` is the publication date of a cited source dated 1-30 days before the effective date
     (S&P DJI press releases), else None. available_at = the end of the announcement day when known, else the
     effective day's 09:30 ET open. The constituents snapshot is available at its revision time. Read through the page
     HTML (the MediaWiki API answered 429 to this shared host); the revision id comes from ``wgRevisionId``. CC BY-SA
     4.0 with attribution and revision in the recipe: in git (ShareAlike).
  8. fja05680/sp500 point-in-time S&P 500 membership (T-473, T-334, T-394) via jsDelivr (raw.githubusercontent.com
     fallback): "S&P 500 Historical Components & Changes (Updated).csv" (one row per change date, 1996-01-02..) and
     ``sp500_ticker_start_end.csv`` (spells; start = first listed date, end = first date without the ticker). Rows are
     joins and leaves as separate events (a join row never carries its future end date). available_at = that date's
     09:30 ET. The repository is MIT-licensed, but its 1996-2019 base list comes from Andreas Clenow's book download
     (redistribution terms not stated) and later changes from Wikipedia: local only.
  9. FINRA consolidated short interest (T-375, T-393): ``POST api.finra.org/data/group/otcMarket/name/
     consolidatedShortInterest``, one bundle per settlement date (2017-12-29.., twice monthly), exchange-listed
     market classes (OTC and OTCBB left out; checked against the date's total count). FINRA does not publish the
     dissemination date in the API: available_at = the end of the 8th NYSE business day after settlement (FINRA
     publishes about 7 business days after settlement; one day of margin). FINRA terms: local only.
 10. FINRA Reg SHO daily short volume, CNMS consolidated files (T-394): ``cdn.finra.org/equity/regsho/daily/
     CNMSshvol<YYYYMMDD>.txt`` from 2018-08-01, filtered to the S&P 500 point-in-time members since 2018-08 (fja +
     Wikipedia; ``universe='all'`` keeps every symbol), monthly bundles. available_at = the end of the trade day (UTC);
     live files are written about 17:20 ET. Files before 2021-10 were (re)written on 2021-10-27 (their Last-Modified):
     the consolidated values are sums of the per-facility files FINRA published each evening, and each file-day row
     carries ``last_modified`` and ``backfilled`` so a stricter harness can use the file time. Local only.
 11. SEC EDGAR (T-375, T-393, T-451, T-473; User-Agent 'Mandel research mandel-research@example.org'):
     ``company_tickers_exchange.json`` (current CIK/ticker/exchange map), ``cik-lookup-data.txt`` (every entity
     name -> CIK, for mapping delisted names), ``data.sec.gov/submissions`` (profile with SIC for the sector of removed
     members; XBRL filings with acceptance times) and ``data.sec.gov/api/xbrl/companyfacts`` (cover-page shares,
     F-score and book-equity concepts as reported in each filing; companyfacts, not frames, so values stay tied to their
     accession). Submissions' ``acceptanceDateTime`` was found LATER than the filing's own EDGAR header by the ET offset
     for AAPL filings 2016-2026 (header 2023-11-02 18:08:27 ET, JSON 2023-11-03T02:08:27Z) and equal to it for CELG 2019,
     never earlier in 15 samples: available_at = min(JSON acceptance, filing date 22:00 ET), both upper bounds. Bundles
     are bucketed by CIK // 50,000. SEC data is public domain; the ticker map is small and in git, the bulk buckets
     and the lookup table are kept local for repository size.
 12. StockTwits symbol streams (T-414, T-415): ``api.stocktwits.com/api/2/streams/symbol/<S>.json?max=<id>`` paged
     backwards (30 per page). Message ids are global and rise with time, so ``find_anchor`` bisects the SPY stream for
     the id at a time and every symbol pages from it. T-414: every message between 16:00 ET on the previous session
     and 09:00 ET for 14 symbols. T-415: the 30 messages before 10:00, 13:00 and 15:30 ET for the 21 IV-covered
     symbols (intensity = messages per hour). Retention is a rolling year (max=632,500,000 returned 2025-10-15 on
     2026-10-15), so the oldest days are lost daily until fetched. Messages are available at their created_at; deleted
     messages are absent; bodies are not stored. Weekly bundles; StockTwits terms: local only.
 13. Deribit option trades near expiry (T-337), ``history.deribit.com``: ``mandel.data.deribit_history`` samples only
     07:00-09:00 UTC, so this module pages the whole UTC day (``options_misc.trades_in_window``) from 2024-03-01 and keeps
     expiries 0-8 calendar days out (the expiries bracketing any Polymarket 'Bitcoin above' resolution within a week),
     with deribit_history's row builders. Monthly bundles; no licence: local only.

Resumable and safe to re-run: every fetcher skips bundles stored complete ON THIS MACHINE (``refresh=True``
re-fetches), stores a chunk as soon as it is complete (tag ``complete``; a cut-short chunk is tagged ``partial`` and
fetched again), never stores a failed or empty answer, and ``dry_run`` makes no request and returns the plan with
per-bundle estimates (calls, rows, gzip MB, minutes). ``limit`` caps the bundles (chunks) fetched per fetcher.
Entry point ``fetch_all_misc(store, dry_run=False, limit=None)``; command line ``python3 -m mandel.data.public_history
all [--dry-run] [--limit N] [--preview DIR | --root DIR]``. Local-only bundles need the ``.gitignore`` lines from
``gitignore_lines()``.

ONE WRITER AT A TIME (see ``RUN_LOCKS``): ``Archive`` re-reads INDEX.json with ``setdefault``, so two processes storing
into one archive leave stale INDEX entries. A storing run (``fetch_all_misc``, a single source or ``merge`` from the
command line) holds ``<archive root>/public_history.run.lock`` and ``crypto_history.run.lock`` (a second writer of
either module is refused) and ends with ``_index_check`` (the command line exits 1 on a mismatch). To run StockTwits
(about 38 hours) next to the rest, store it under its own ``--root DIR`` and ``merge --from DIR`` afterwards.
"""

from __future__ import annotations

import argparse
import contextlib
import csv
import fcntl
import html as htmlmod
import io
import json
import math
import re
import sys
import time
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from fnmatch import fnmatch
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable, Optional
from urllib.parse import quote
from zoneinfo import ZoneInfo

import httpx

from mandel.data.crypto_history import _stored_complete
from mandel.data.deribit_history import (DAILY_ROW_GZ_BYTES, DAILY_WINDOW_TRADES, HISTORY_LAG_MS, _daily_group, _daily_idx,
                                         _daily_sort, _daily_trade_row, _dte_keep, _merge_store, _sample_row, _scan, _Spill,
                                         daily_window)
from mandel.data.dolt_options import _last_weekday, _nth_weekday, is_trading_day, prev_session
from mandel.data.fetch import Throttle
from mandel.data.market_history import ARCHIVE_DATA_PREFIX, _et_to_utc, _load, _stored_here
from mandel.data.options_misc import (RETRY_STATUS, _archive, _client, _end_of_day, _get, _retry_after,
                                      trades_in_window)

UTC = timezone.utc
NY = ZoneInfo("America/New_York")
SEC_UA = {"User-Agent": "Mandel research mandel-research@example.org"}          # SEC fair-access policy: name + contact
WIKI_UA = {"User-Agent": "mandel-substrate/0.1 (https://github.com/pointlessbassmusic-byte/polymarket-arbitrage; "
                         "mandel-research@example.org) httpx"}                     # Wikimedia User-Agent policy
_sleep = time.sleep                     # patched in tests (this module's own waits; options_misc._get has its own)


def _now() -> datetime:                 # patched in tests
    return datetime.now(UTC)


def _stamp() -> str:
    """The fetch time, whole seconds rounded up (one ISO format for every timestamp of this module)."""
    return _iso_ceil(_now())


def _today() -> str:
    return _now().date().isoformat()


def _iso(dt: datetime) -> str:
    return dt.astimezone(UTC).replace(microsecond=0).isoformat()


def _iso_ceil(dt: datetime) -> str:
    """Whole seconds, rounded up (an available_at must never move earlier)."""
    dt = dt.astimezone(UTC)
    return _iso(dt + timedelta(seconds=1)) if dt.microsecond else _iso(dt)


def _parse_iso(s: str) -> datetime:
    d = datetime.fromisoformat(str(s).replace("Z", "+00:00"))
    return d if d.tzinfo else d.replace(tzinfo=UTC)


def _et(day: str, hhmm: str) -> str:
    return _et_to_utc(day[:10], hhmm)


def _next_morning(day: str) -> str:
    """09:30 ET on the next calendar day: a value struck after the close is known at the next session's open (t+1)."""
    return _et((date.fromisoformat(day[:10]) + timedelta(days=1)).isoformat(), "09:30")


def _not_after(rule: str, fetched: str) -> str:
    """available_at never later than the fetch time (a value in hand was known by then)."""
    return rule if _parse_iso(rule) <= _parse_iso(fetched) else fetched


def _row(source: str, kind: str, key: str, observed_at: str, available_at: str, value: dict) -> dict:
    return {"source": source, "kind": kind, "key": key, "observed_at": observed_at, "available_at": available_at, "value": value}


def _put(archive, name: str, rows: Iterable[dict], recipe: str, tags: Iterable[str], schema: str, in_git: bool, complete: bool = True):
    tg = sorted(set(tags) | {"complete" if complete else "partial"})
    return archive.store(name, kind="data", rows=rows, recipe=recipe, tags=tg, schema=schema, in_git=in_git)


def _have(archive, name: str) -> bool:
    """Stored complete and present on this machine (a local-only bundle on a fresh clone has an INDEX entry only)."""
    return archive is not None and _stored_here(archive, name) and _stored_complete(archive, name)


def _stored_info(e) -> dict:
    return {"rows": e.rows, "bytes": e.bytes, "in_git": e.in_git}


def _fail(info: dict) -> dict:
    return {k: info[k] for k in ("url", "status", "error", "body", "retry_after", "attempts") if k in info}


def _num(x):
    if x is None or isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return x
    s = str(x).strip().replace(",", "")
    if s in ("", "null", "None", "NaN", "-", "--"):
        return None
    try:
        v = float(s)
        return int(v) if v.is_integer() and "." not in s and "e" not in s.lower() else v
    except ValueError:
        return None


def _post_json(client: httpx.Client, url: str, throttle: Throttle, body: dict, headers: Optional[dict] = None, retries: int = 4,
               backoff: float = 1.0, max_wait: float = 120.0) -> tuple[Optional[httpx.Response], dict]:
    """JSON POST twin of options_misc._get: retries 429/5xx and transport errors with backoff, honours Retry-After (a wait
    longer than ``max_wait`` is final); any other non-2xx answer is final."""
    info: dict = {"url": url, "attempts": 0}
    for attempt in range(retries):
        throttle.wait()
        info["attempts"] = attempt + 1
        try:
            r = client.post(url, json=body, headers={"Accept": "application/json", **(headers or {})})
        except httpx.HTTPError as e:
            info["error"] = repr(e)[:200]
            _sleep(backoff * 2 ** attempt)
            continue
        info["status"] = r.status_code
        if r.status_code in (200, 204):
            info.pop("error", None)
            return r, info
        if r.status_code in RETRY_STATUS:
            ra = _retry_after(r)
            if ra is not None and ra > max_wait:
                info.update(retry_after=ra, body=r.text[:300])
                return None, info
            _sleep(max(backoff * 2 ** attempt, ra or 0.0))
            continue
        info["body"] = r.text[:300]
        return None, info
    return None, info


# ============================================================================================ families, licences ==

LOCAL = "The repo is public, so this bundle is local disk only (in_git=False); rebuild it from this recipe."
LICENCE = {
    "cftc": "CFTC Commitments of Traders: US federal government data (CFTC), public domain",
    "treasury": "US Treasury Fiscal Data (Bureau of the Fiscal Service): US government data, public domain",
    "clevelandfed": "Federal Reserve Bank of Cleveland website content (a Reserve Bank, not a federal agency): no open licence",
    "spdr": "SPDR Gold Trust archive (State Street / World Gold Trust Services): 'Reproduction or redistribution ... expressly prohibited'",
    "sprott": "Sprott Asset Management website data: Sprott's terms of use, no redistribution licence",
    "ssga": "State Street Global Advisors holdings file: (c) State Street Corporation, all rights reserved",
    "wikipedia": "Wikipedia text and tables, CC BY-SA 4.0 (https://creativecommons.org/licenses/by-sa/4.0/), attribution in the recipe",
    "fja": ("fja05680/sp500 (MIT licence for the repository), but the 1996-2019 base list is Andreas Clenow's 'Trading Evolved' "
            "download (redistribution terms not stated) and later changes come from Wikipedia (CC BY-SA): mixed provenance"),
    "finra": "FINRA data (FINRA API / Reg SHO files): FINRA terms of use, no open redistribution licence",
    "sec": "SEC EDGAR: US federal government data, public domain",
    "stocktwits": "StockTwits API content: StockTwits terms of service, user content, no redistribution licence",
    "deribit": "Deribit public history API: no stated data licence (no redistribution right)",
}

# key: (bundle pattern, function, theories, in_git, licence key, available_at rule, gz bytes per row, seconds per call)
FAMILIES: dict[str, tuple] = {
    "cftc": ("cftc-cot-*", "fetch_cftc_cot", ["T-405"], True, "cftc",
             "Socrata :created_at (live publication, 2022-07..), never before the report week's Friday 15:30 ET; bulk-loaded rows: "
             "Friday 15:30 ET + 1 business day per federal closure Wed-Fri, CFTC 2019 catch-up schedule, shutdown bounds", 450, 1.2),
    "treasury": ("treasury-auctions", "fetch_treasury_auctions", ["T-442"], True, "treasury",
                 "announcement: end of the announcement day (UTC); result: competitive close (13:00 ET for coupons) + 30 min", 125, 1.4),
    "clevelandfed": ("clevelandfed-inflation-nowcast", "fetch_clevelandfed_nowcast", ["T-360"], False, "clevelandfed",
                     "nowcast labelled D: D 12:00 ET (updated each business day around 10:00 ET); actual print: release-day label 08:30 ET",
                     10, 2.5),
    "clevelandfed_snap": ("clevelandfed-nowcast-snapshot-*", "fetch_clevelandfed_snapshot", ["T-360"], False, "clevelandfed",
                          "capture time (a dated vintage of the recent target months)", 10, 2.5),
    "spdr": ("spdr-*-archive", "fetch_spdr_archive", ["T-404", "T-370"], False, "spdr",
             "next calendar day 09:30 ET (t+1; NAV and holdings are published after the close)", 60, 2.0),
    "sprott": ("sprott-*-nav", "fetch_sprott_trusts", ["T-472"], False, "sprott",
               "next calendar day 09:30 ET (NAV struck after the close)", 20, 2.0),
    "ssga": ("ssga-*-holdings-*", "fetch_ssga_holdings", ["T-374"], False, "ssga",
             "file Last-Modified (not before the as-of close); capture time when absent", 42, 2.0),
    "wiki_const": ("wikipedia-sp500-constituents", "fetch_wikipedia_sp500", ["T-374", "T-473", "T-394"], True, "wikipedia",
                   "the page revision's time", 47, 3.5),
    "wiki_changes": ("wikipedia-sp500-changes", "fetch_wikipedia_sp500", ["T-334", "T-394", "T-473"], True, "wikipedia",
                     "end of the announcement day when a cited S&P index announcement is dated 1-30 days before the effective date, "
                     "else the effective day's 09:30 ET open", 79, 3.5),
    "fja_comp": ("fja-sp500-components", "fetch_fja_sp500", ["T-473", "T-334", "T-394"], False, "fja",
                 "the change date's 09:30 ET (effective before the open, announced earlier)", 34, 2.0),
    "fja_spells": ("fja-sp500-ticker-spells", "fetch_fja_sp500", ["T-473", "T-334"], False, "fja",
                   "join / leave date 09:30 ET (separate rows; a join row never carries its end date)", 15, 1.5),
    "finra_si": ("finra-short-interest-*", "fetch_finra_short_interest", ["T-375", "T-393"], False, "finra",
                 "end (23:59:59 UTC) of the 8th NYSE business day after the settlement date", 49, 1.7),
    "finra_cnms": ("finra-regsho-cnms-*", "fetch_finra_cnms", ["T-394"], False, "finra",
                   "end of the trade day (23:59:59 UTC); file-day rows carry Last-Modified and a backfilled flag", 20, 0.7),
    "sec_tickers": ("sec-company-tickers", "fetch_sec_company_tickers", ["T-375", "T-393", "T-451"], True, "sec",
                    "file Last-Modified (current map, ex post)", 28, 1.0),
    "sec_lookup": ("sec-cik-lookup", "fetch_sec_cik_lookup", ["T-393", "T-375", "T-451", "T-473"], False, "sec",
                   "capture time (reference table of every EDGAR entity name)", 19, 20.0),
    "sec_sub": ("sec-submissions-*", "fetch_sec_submissions", ["T-451", "T-473", "T-393"], False, "sec",
                "filing: min(JSON acceptanceDateTime, filing date 22:00 ET); company profile: capture time", 42, 0.75),
    "sec_facts": ("sec-companyfacts-*", "fetch_sec_companyfacts", ["T-451", "T-393", "T-375"], False, "sec",
                  "min(acceptance of the fact's accession, filed date 22:00 ET)", 22, 0.7),
    "st_preopen": ("stocktwits-preopen-*", "fetch_stocktwits_preopen", ["T-414"], False, "stocktwits",
                   "message: created_at; window row: window end (09:00 ET)", 32, 1.95),
    "st_intensity": ("stocktwits-intensity-*", "fetch_stocktwits_intensity", ["T-415"], False, "stocktwits",
                     "the sample time (only messages created before it)", 120, 1.25),
    "st_anchors": ("stocktwits-anchors", "fetch_stocktwits_preopen", ["T-414", "T-415"], False, "stocktwits",
                   "the anchor's target time (fetch metadata)", 60, 1.1),
    "deribit_near": ("deribit-option-trades-near-*", "fetch_deribit_near_trades", ["T-337"], False, "deribit",
                     "trade time; sample-day row: window end", DAILY_ROW_GZ_BYTES, 0.47),
}


def _family(name: str) -> Optional[tuple]:
    hits = [(k, f) for k, f in FAMILIES.items() if fnmatch(name, f[0])]
    return max(hits, key=lambda kf: len(kf[1][0]))[1] if hits else None


def gitignore_line(name: str) -> Optional[str]:
    """The exact .gitignore line covering a local-only bundle of this module (None for an in-git bundle)."""
    f = _family(name)
    if f is None or f[3]:
        return None
    return f"{ARCHIVE_DATA_PREFIX}{f[0]}.jsonl.gz"


def gitignore_lines() -> list[str]:
    return [f"{ARCHIVE_DATA_PREFIX}{f[0]}.jsonl.gz" for f in FAMILIES.values() if not f[3]]


def _plan(fam: str, name: str, calls: float, rows: float, stored: bool = False, sec_call: Optional[float] = None, **extra) -> dict:
    pat, fn, theories, in_git, lic, rule, gz_row, sec = FAMILIES[fam]
    sec = sec if sec_call is None else sec_call
    calls = 0 if stored else int(math.ceil(calls))
    return {"name": name, "function": fn, "theories": theories, "in_git": in_git, "licence": LICENCE[lic],
            "gitignore_line": "" if in_git else f"{ARCHIVE_DATA_PREFIX}{pat}.jsonl.gz", "calls": calls, "rows": int(rows),
            "gz_mb": round(rows * gz_row / 1e6, 3), "minutes": round(calls * sec / 60, 2), "available_at_rule": rule,
            "already_stored": stored, **extra}


def _recipe(core: str, lic: str, in_git: bool) -> str:
    return f"{core} Licence: {LICENCE[lic]}." + ("" if in_git else " " + LOCAL)


def _setup(store, report, client, min_interval, timeout: float = 120):
    rep = report if report is not None else {}
    own = client is None
    return _archive(store), rep, (client or _client(timeout)), own, Throttle(min_interval)


# ============================================================================== US federal closures (COT rule) ==

# Executive-order and special federal closures (DC offices), plus Inauguration Day (a DC-area federal holiday).
EXTRA_FEDERAL_CLOSURES = {
    "1989-01-20", "1993-01-20", "1994-04-27", "1997-01-20", "2001-01-20", "2001-12-24", "2004-06-11", "2005-01-20",
    "2007-01-02", "2007-12-24", "2008-12-26", "2009-01-20", "2012-10-29", "2012-10-30", "2012-12-24", "2013-01-21",
    "2014-12-26", "2015-12-24", "2017-01-20", "2018-12-05", "2018-12-24", "2019-12-24", "2020-12-24", "2021-01-20",
    "2024-12-24", "2025-01-09", "2025-01-20", "2025-12-24", "2025-12-26",
}


def _fed_observed(d: date) -> date:
    return d - timedelta(days=1) if d.weekday() == 5 else d + timedelta(days=1) if d.weekday() == 6 else d


def federal_closures(year: int) -> set[date]:
    """US federal holidays (observed) plus EXTRA_FEDERAL_CLOSURES for one calendar year."""
    h: set[date] = set()
    if date(year, 1, 1).weekday() != 5:                 # a Saturday New Year is observed on Dec 31 of the year before
        h.add(_fed_observed(date(year, 1, 1)))
    if date(year + 1, 1, 1).weekday() == 5:
        h.add(date(year, 12, 31))
    if year >= 1986:
        h.add(_nth_weekday(year, 1, 0, 3))
    h.add(_nth_weekday(year, 2, 0, 3))
    h.add(_last_weekday(year, 5, 0))
    if year >= 2021:
        h.add(_fed_observed(date(year, 6, 19)))
    h.add(_fed_observed(date(year, 7, 4)))
    h.add(_nth_weekday(year, 9, 0, 1))
    h.add(_nth_weekday(year, 10, 0, 2))
    h.add(_fed_observed(date(year, 11, 11)))
    h.add(_nth_weekday(year, 11, 3, 4))
    h.add(_fed_observed(date(year, 12, 25)))
    h |= {date.fromisoformat(d) for d in EXTRA_FEDERAL_CLOSURES if d.startswith(str(year))}
    return h


def _federal_business_day(d: date) -> bool:
    return d.weekday() < 5 and d not in federal_closures(d.year)


# ===================================================================================== 1. CFTC COT (T-405) ==

CFTC_BASE = "https://publicreporting.cftc.gov/resource/"
COT_DATASETS = {"disagg-fo": ("kh3c-gbw2", "Disaggregated - Combined (futures and options)"),
                "disagg-f": ("72hh-3qpy", "Disaggregated - Futures Only"),
                "legacy-fo": ("jun7-fc8e", "Legacy - Combined (futures and options)"),
                "legacy-f": ("6dca-aqww", "Legacy - Futures Only"),
                "tff-fo": ("yw9f-hn96", "Traders in Financial Futures - Combined"),
                "tff-f": ("gpe5-46if", "Traders in Financial Futures - Futures Only")}
COT_CODES = {"088691": "gold (COMEX)", "084691": "silver (COMEX)", "133741": "bitcoin (CME)"}
COT_SERIES = (("disagg-fo", "088691"), ("disagg-f", "088691"), ("disagg-fo", "084691"), ("disagg-f", "084691"),
              ("legacy-f", "088691"), ("legacy-fo", "088691"), ("legacy-f", "084691"), ("legacy-fo", "084691"),
              ("tff-fo", "133741"), ("tff-f", "133741"))
COT_ROWS_EST = {("disagg", "088691"): 1061, ("disagg", "084691"): 1061, ("legacy-f", "088691"): 1937, ("legacy-f", "084691"): 1937,
                ("legacy-fo", "088691"): 1600, ("legacy-fo", "084691"): 1600, ("tff", "133741"): 444}
COT_PAGE = 5000
COT_LIVE_MAX_DAYS = 60                   # a :created_at this close after the report date is the live publication
# Publication history. CFTC 'About the COT Reports' (read 2026-10-15; fetch6c/misc/check/h.html): until 1990 the reports
# were 'compiled on an end-of-month basis and published on the 11th or 12th calendar day of the following month'; then
# 'published more often - switching to mid-month and month-end in 1990, to every two weeks in 1992, and to weekly in 2000'
# and 'moving the publication to the sixth business day after the "as of" date in 1990 and then to the third business
# day after the "as of" date in 1992'. CSI Data's COT manual (csidata.com, read 2026-10-15): '2 weekly reports have been
# issued on alternate Fridays since 10-6-92' (each release carried the latest week and the week before). Socrata's legacy
# rows are semi-monthly (15th and month-end, any weekday) until 1992-09-30 and weekly (Tuesday) from 1992-10-06.
COT_MONTHLY_BEFORE = "1990-01-01"        # monthly publication: mid-month rows before this were not published in real time
COT_SEMI_MONTHLY_FROM = "1991-01-01"     # the 1990 switch month is not stated: 1990 rows keep the (later) monthly schedule
COT_BIWEEKLY_FROM = "1992-10-01"         # weekly as-of dates, two published together every other Friday
COT_WEEKLY_FROM = "2001-01-01"           # weekly 'in 2000' (month not stated): all of 2000 keeps the biweekly bound
COT_MONTHLY_DAY = 12                     # 'the 11th or 12th calendar day of the following month' (the later one)
COT_SEMI_MONTHLY_BDAYS = 6               # 1990-92: 'the sixth business day after the as-of date'
COT_MIN_BDAYS = 3                        # 1992-: 'the third business day after the as-of date' (Tuesday -> Friday)
# CFTC press release 7864-19 (read 2026-10-15): the report of Mon 2018-12-24 on Fri 2019-02-01, then one report on Tuesday
# and one on Friday each week in chronological order until current (the 2019-03-05 report on its normal Friday, 03-08).
COT_2019_CATCHUP = {"2018-12-24": "2019-02-01", "2018-12-31": "2019-02-05", "2019-01-08": "2019-02-08", "2019-01-15": "2019-02-12",
                    "2019-01-22": "2019-02-15", "2019-01-29": "2019-02-19", "2019-02-05": "2019-02-22", "2019-02-12": "2019-02-26",
                    "2019-02-19": "2019-03-01", "2019-02-26": "2019-03-05"}
# October 2013 shutdown. CFTC release 6745-13 (2013-10-23): the first delayed report (as of 10-01, due 10-04) on 2013-10-25,
# two more in the week of 10-28, at least two more in the week of 11-04, and the normal schedule by the 11-08 release. CFTC
# 'COT reports update' notices (relayed by Mondo Visione): 'the current reports for the week of October 29, 2013 are now
# available' on 2013-11-06 and 'for the week of November 5, 2013' on 2013-11-08 (the normal Friday). Days are known by date
# only: end of that day (UTC). The 10-08..10-29 reports count from 11-06, the first day all of them are known to be out.
COT_2013_CATCHUP = {"2013-10-01": "2013-10-25", "2013-10-08": "2013-11-06", "2013-10-15": "2013-11-06", "2013-10-22": "2013-11-06",
                    "2013-10-29": "2013-11-06"}
# Shutdowns whose catch-up schedule was not found: every report dated in [first, last] counts from the bound (conservative).
COT_SHUTDOWN_BOUNDS = (("1995-11-14", "1996-01-23", "1996-02-29", "1995-96 government shutdowns and January 1996 closures "
                                                                  "(catch-up schedule not found)"),)
# Reports whose history was released after the report began: rows dated before ``live_from`` became public at ``history``.
# Disaggregated (Disaggregated Explanatory Notes, fetch6c/misc/check/cftc_disagg_expl.html): first published 2009-09-04 for
# 22 markets, the remaining physical markets from 2009-12-04; 'On October 20, 2009, CFTC made available ... historical
# disaggregated data ... dated back to June 13, 2006', classified with current trader categories (a backcast). COMEX gold
# 088691 and silver 084691 (and copper 085692, lumber 058643) are in the first report: the CFTC archive page
# cotarchives/2009/futures/other_lf090109.htm 'Positions as of September 1, 2009' lists them; platinum and palladium appear
# first on the 2009-12-01 page (fetch6c/misc/probes2/). Markets not known to be in the first 22 count from 2009-12-04.
# TFF: first report as of 2010-07-20 (cotarchives/2010/futures/financial_lf072010.htm; none for 07-13; CFTC release
# 5857-10 of 2010-07-22), four years of history made available in December 2010 (Hedgeweek, 2010-12-07). CME bitcoin
# 133741 first appears in the TFF report of 2018-04-10 (financial_lf041018.htm; absent on 04-03), so its rows are live.
COT_BACKCAST = {
    "disagg": {"live_from": "2009-09-01", "history": "2009-10-20", "first_codes": frozenset({"088691", "084691", "085692", "058643"}),
               "added": ("2009-12-01", "2009-12-04"),
               "why": "CFTC historical disaggregated release 2009-10-20 (backcast with current trader classifications)"},
    "tff": {"live_from": "2010-07-20", "history": "2010-12-07", "first_codes": None, "added": None,
            "why": "CFTC Traders in Financial Futures history release, December 2010 (backcast; Hedgeweek 2010-12-07)"},
}
COT_PIT = ("live", "schedule", "bound", "backcast", "not_realtime")
COT_SCHEMA = ("kind cot: value = the raw Socrata row (strings, every column of the report) + {report_date, dataset, report, "
              "created_at, updated_at (Socrata row metadata), release_rule (the publication schedule's release for this as-of "
              "date, ISO UTC), available_basis, pit (live = Socrata publication time; schedule = documented schedule; bound = "
              "conservative bound where the schedule is not fully documented; backcast = history released later than the "
              "schedule, available_at = that release; not_realtime = not published in real time, available_at is only a "
              "lower bound: exclude from point-in-time tests)}; observed_at = report date 17:00 ET (positions as of that "
              "close); available_at as in the recipe, never before observed_at")


def cot_bundle(report: str, code: str) -> str:
    return f"cftc-cot-{report}-{code}"


def _add_fed_bdays(d: date, n: int) -> date:
    """The n-th federal business day after ``d``."""
    for _ in range(n):
        d += timedelta(days=1)
        while not _federal_business_day(d):
            d += timedelta(days=1)
    return d


def _cot_regular_friday(r: date) -> date:
    """The release Friday of the weekly schedule before holiday shifts: the as-of week's Friday for a Monday-Wednesday
    as-of date, the next week's Friday for a Thursday or Friday as-of date (holiday weeks: never the as-of day itself)."""
    fri = r + timedelta(days=(4 - r.weekday()) % 7)
    return fri + timedelta(days=7) if r.weekday() in (3, 4) else fri


def _cot_weekly_day(r: date) -> tuple[date, int]:
    """(release day, closures that moved it) under the weekly schedule: the regular Friday, one federal business day later
    per federal closure on the Wednesday-Friday of that week after the as-of date (CFTC: 'Federal holidays may delay
    release by one or two days'; a Monday or Tuesday holiday does not move it), never before the third federal business
    day after the as-of date."""
    fri = _cot_regular_friday(r)
    closed = sum(1 for k in (2, 3, 4) for d in [fri - timedelta(days=4 - k)] if d > r and not _federal_business_day(d))
    rel = _add_fed_bdays(fri, closed) if closed else fri
    return max(rel, _add_fed_bdays(r, COT_MIN_BDAYS)), closed


def cot_release_rule(report_date: str) -> tuple[str, str, str]:
    """(ISO UTC release, basis, pit) by the CFTC's publication schedule for this as-of date (any report):
      * before 1990 (monthly publication): end of the 12th calendar day of the following month (next federal business
        day); mid-month rows were not part of that publication (pit not_realtime), month-end rows are a bound;
      * 1990 (the switch month to semi-monthly is not stated): the same, later, monthly schedule (bound);
      * 1991 .. 1992-09 (semi-monthly): end of the 6th federal business day after the as-of date (schedule);
      * 1992-10 .. 2000 (weekly data, two weeks published every other Friday): the release day the FOLLOWING week's
        report would have, end of day (bound: which week of a pair was the latest is not documented);
      * 2001 on (weekly): 15:30 ET on the regular Friday with holiday shifts and the 3-business-day minimum
        (``_cot_weekly_day``); the 2013 and 2019 shutdown catch-up schedules (schedule);
      * the 1995-96 shutdown bound."""
    rd = report_date[:10]
    r = date.fromisoformat(rd)
    if rd in COT_2019_CATCHUP:
        return _et(COT_2019_CATCHUP[rd], "15:30"), "CFTC press release 7864-19 catch-up schedule (2019 shutdown)", "schedule"
    if rd in COT_2013_CATCHUP:
        return (_end_of_day(COT_2013_CATCHUP[rd]), "CFTC release 6745-13 and 2013-11-06 notice: 2013 shutdown catch-up "
                f"(out by {COT_2013_CATCHUP[rd]}, end of day)", "schedule")
    if rd < COT_SEMI_MONTHLY_FROM:
        nm = date(r.year + (r.month == 12), r.month % 12 + 1, COT_MONTHLY_DAY)
        while not _federal_business_day(nm):
            nm += timedelta(days=1)
        mid = r.day <= 20
        if rd < COT_MONTHLY_BEFORE:
            basis = ("monthly publication until 1990 ('11th or 12th calendar day of the following month'): end of the "
                     f"{COT_MONTHLY_DAY}th (next federal business day)")
            if mid:
                return (_end_of_day(nm.isoformat()), basis + "; a mid-month row was not in the monthly publication (not real "
                        "time; lower bound only)", "not_realtime")
            return _end_of_day(nm.isoformat()), basis, "bound"
        return (_end_of_day(nm.isoformat()), "1990: semi-monthly publication on the 6th business day began during 1990 (month "
                f"not stated): the later monthly schedule, end of the {COT_MONTHLY_DAY}th of the following month", "bound")
    if rd < COT_BIWEEKLY_FROM:
        return (_end_of_day(_add_fed_bdays(r, COT_SEMI_MONTHLY_BDAYS).isoformat()),
                f"1990-92 semi-monthly schedule: end of the {COT_SEMI_MONTHLY_BDAYS}th federal business day after the as-of date",
                "schedule")
    if rd < COT_WEEKLY_FROM:
        d, _ = _cot_weekly_day(r + timedelta(days=7))
        out, basis, pit = (_end_of_day(d.isoformat()), "1992-2000 biweekly publication (two weekly reports every other Friday): "
                           "bound = the following week's release day, end of day", "bound")
    else:
        d, closed = _cot_weekly_day(r)
        out, pit = _et(d.isoformat(), "15:30"), "schedule"
        basis = ("rule: Friday 15:30 ET" if not closed else f"rule: Friday 15:30 ET + {closed} business day(s) for federal closures "
                 "that week")
        if r.weekday() in (3, 4):
            basis = "rule: the next week's Friday 15:30 ET (Thursday/Friday as-of date)" + (f" + {closed} business day(s)" if closed else "")
    for first, last, bound, why in COT_SHUTDOWN_BOUNDS:
        if first <= rd <= last and _parse_iso(_end_of_day(bound)) > _parse_iso(out):
            out, basis, pit = _end_of_day(bound), f"conservative bound: {why}", "bound"
    return out, basis, pit


def cot_backcast(report: Optional[str], code: Optional[str], report_date: str) -> Optional[tuple[str, str]]:
    """(history release, basis) when this report's row was released only as history (``COT_BACKCAST``), else None."""
    fam = COT_BACKCAST.get(str(report or "").split("-")[0])
    if not fam:
        return None
    rd = report_date[:10]
    if fam["first_codes"] is not None and code not in fam["first_codes"] and fam["added"] and rd < fam["added"][0]:
        return (_end_of_day(fam["added"][1]), f"market not known to be in the first {report} report: added {fam['added'][1]} "
                f"(rows before {fam['added'][0]} count from then; backcast)")
    if rd < fam["live_from"]:
        return _end_of_day(fam["history"]), fam["why"]
    return None


def cot_available_at(report_date: str, created_at: Optional[str], report: Optional[str] = None,
                     code: Optional[str] = None) -> tuple[str, str, str, str]:
    """(available_at, basis, schedule release, pit). A live Socrata creation time wins, never before the regular release
    (15:30 ET on the regular Friday); a backcast row counts from its history release, never before the schedule."""
    rule, basis, pit = cot_release_rule(report_date)
    r = date.fromisoformat(report_date[:10])
    bc = cot_backcast(report, code, report_date)
    if bc:
        return max(rule, bc[0], key=_parse_iso), bc[1], rule, "backcast"
    if created_at and report_date[:10] >= COT_WEEKLY_FROM:
        ca = _parse_iso(created_at)
        lag = (ca.date() - r).days
        if 2 <= lag <= COT_LIVE_MAX_DAYS:
            floor = _et(_cot_regular_friday(r).isoformat(), "15:30")
            return max(_iso_ceil(ca), floor, key=_parse_iso), "Socrata :created_at (live publication)", rule, "live"
    return rule, basis, rule, pit


def parse_cot_rows(raw: list[dict], report: str, code: str, fetched: str, stats: Optional[dict] = None) -> list[dict]:
    """Archive rows of one COT series. Guard: available_at is never before observed_at (raised and counted in
    ``stats['raised_to_observed']``; the schedules above never need it)."""
    ds = COT_DATASETS[report][0]
    out, seen = [], set()
    st = stats if stats is not None else {}
    for x in raw:
        rd = str(x.get("report_date_as_yyyy_mm_dd") or "")[:10]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", rd) or rd in seen:
            continue
        seen.add(rd)
        ca, ua = x.get(":created_at"), x.get(":updated_at")
        av, basis, rule, pit = cot_available_at(rd, ca, report, code)
        obs = _et(rd, "17:00")
        if _parse_iso(av) < _parse_iso(obs):
            av, basis = obs, basis + "; raised to observed_at (guard)"
            st["raised_to_observed"] = st.get("raised_to_observed", 0) + 1
        st[pit] = st.get(pit, 0) + 1
        v = {k: val for k, val in x.items() if not k.startswith(":")}
        v.update(report_date=rd, dataset=ds, report=report, created_at=ca, updated_at=ua, release_rule=rule, available_basis=basis,
                 pit=pit)
        out.append(_row("cftc", "cot", f"{report}:{code}:{rd}", obs, _not_after(av, fetched), v))
    return sorted(out, key=lambda r: r["value"]["report_date"])


def fetch_cftc_cot(store=None, series: Iterable[tuple[str, str]] = COT_SERIES, dry_run: bool = False, limit: Optional[int] = None,
                   skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 1.0,
                   retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``cftc-cot-<report>-<code>``: every row of one contract in one COT report, all columns, paged 5,000 at a time."""
    archive = _archive(store)
    rep = report if report is not None else {}
    todo, plan = [], []
    for rpt, code in series:
        name = cot_bundle(rpt, code)
        have = _have(archive, name)
        est = COT_ROWS_EST.get((rpt.split("-")[0], code)) or COT_ROWS_EST.get((rpt, code)) or 1000
        if have and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            plan.append(_plan("cftc", name, 0, est, stored=True))
            continue
        todo.append((rpt, code))
    if limit is not None:
        todo = todo[:limit]
    for rpt, code in todo:
        est = COT_ROWS_EST.get((rpt.split("-")[0], code)) or COT_ROWS_EST.get((rpt, code)) or 1000
        plan.append(_plan("cftc", cot_bundle(rpt, code), math.ceil((est + 1) / COT_PAGE), est))
    rep["_plan"] = plan
    if dry_run:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        for rpt, code in todo:
            name = cot_bundle(rpt, code)
            ds, title = COT_DATASETS[rpt]
            raw, offset, info = [], 0, {}
            while True:
                r, info = _get(c, CFTC_BASE + ds + ".json", th, params={
                    "$select": "*,:created_at,:updated_at", "cftc_contract_market_code": code,
                    "$order": "report_date_as_yyyy_mm_dd,id", "$limit": COT_PAGE, "$offset": offset}, retries=retries, backoff=backoff)
                if r is None:
                    break
                try:
                    page = r.json()
                except ValueError:
                    info["error"] = "not JSON"
                    r = None
                    break
                raw.extend(page)
                if len(page) < COT_PAGE:
                    break
                offset += COT_PAGE
            fetched = _stamp()
            if r is None or not raw:
                rep[name] = {"error": _fail(info) or "no rows"}
                continue
            stats: dict = {}
            rows = parse_cot_rows(raw, rpt, code, fetched, stats)
            live = stats.get("live", 0)
            rep[name] = {"rows": len(rows), "first": rows[0]["value"]["report_date"], "last": rows[-1]["value"]["report_date"],
                         "live_rows": live, "pit": {k: stats[k] for k in COT_PIT if stats.get(k)},
                         "raised_to_observed": stats.get("raised_to_observed", 0)}
            if archive is None:
                rep[name]["rows"] = rows
                continue
            pit_txt = ", ".join(f"{k} {n}" for k, n in rep[name]["pit"].items())
            e = _put(archive, name, rows, _recipe(
                f"GET {CFTC_BASE}{ds}.json?$select=*,:created_at,:updated_at&cftc_contract_market_code={code}&$order="
                f"report_date_as_yyyy_mm_dd,id&$limit={COT_PAGE}&$offset=<k> ({title}; {COT_CODES.get(code, code)}) via "
                f"mandel.data.public_history.fetch_cftc_cot; {min_interval}s pacing, {retries} tries with backoff on 429/5xx. "
                f"available_at: the row's Socrata :created_at when within {COT_LIVE_MAX_DAYS} days of the report date (the live "
                f"publication, 2022-07 on), never before the regular Friday 15:30 ET; other rows (bulk load of 2022) use the CFTC "
                "publication schedule of their era (cot_release_rule): before 1990 monthly (end of the 12th of the following "
                "month; mid-month rows were not published in real time: pit not_realtime), 1990 the same later bound, 1991-92 "
                "the 6th federal business day, 1992-10..2000 biweekly (two weekly reports every other Friday: the following "
                "week's release day, end of day), from 2001 Friday 15:30 ET (the next week's Friday for a Thursday/Friday as-of "
                "date) plus one federal business day per federal closure on the Wednesday-Friday of the release week and never "
                "before the 3rd federal business day; the 2013 (CFTC 6745-13 and notices) and 2018-19 (7864-19) catch-up "
                "schedules; a conservative bound for the 1995-96 shutdowns. Backcasts: disaggregated rows before 2009-09-01 "
                "were released as history on 2009-10-20 and TFF rows before 2010-07-20 in December 2010: available_at = that "
                "release (pit backcast; the schedule time stays in value.release_rule). available_at is never before "
                f"observed_at. Rows by pit: {pit_txt}.", "cftc", True),
                ["cot", "cftc", "positioning", "weekly", rpt, code, "T-405"], COT_SCHEMA, in_git=True)
            rep[name]["stored"] = _stored_info(e)
            if log:
                log(f"cftc {name}: {len(rows)} weeks")
    finally:
        if own:
            c.close()
    return rep


# ============================================================================== 2. Treasury auctions (T-442) ==

FISCAL_AUCTIONS = "https://api.fiscaldata.treasury.gov/services/api/fiscal_service/v1/accounting/od/auctions_query"
AUCTION_BUNDLE = "treasury-auctions"
AUCTION_FIRST = "1979-01-01"
AUCTION_TYPES = ("Note", "Bond")
AUCTION_PAGE = 1000
AUCTION_RESULT_LAG_MIN = 30
ANNOUNCE_FIELDS = ("cusip", "security_type", "security_term", "original_security_term", "security_term_week_year",
                   "security_term_day_month", "series", "auction_date", "issue_date", "maturity_date", "dated_date",
                   "announcemt_date", "offering_amt", "reopening", "inflation_index_security", "floating_rate", "callable",
                   "cash_management_bill_cmb", "closing_time_comp", "closing_time_noncomp", "original_issue_date",
                   "original_dated_date", "tint_cusip_1", "tint_cusip_2")
RESULT_MARKERS = ("high_yield", "high_price", "high_discnt_rate", "high_investment_rate", "high_discnt_margin", "bid_to_cover_ratio")
AUCTION_SCHEMA = ("kind auction-announcement: value = the schedule fields (cusip, security_type, security_term, original_security_term, "
                  "auction_date, issue_date, maturity_date, announcemt_date, offering_amt, reopening, inflation_index_security, "
                  "floating_rate, closing times, ...) + nominal_coupon (Note/Bond, not TIPS, not FRN) + available_basis; "
                  "available_at = end of the announcement day (UTC) | kind auction-result: value = the full raw auctions_query row "
                  "('null' -> None) + nominal_coupon; available_at = auction day competitive close (ET) + 30 min; observed_at = "
                  "the competitive close for both kinds")


def _closing_time(s: Optional[str]) -> Optional[str]:
    m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})\s*([AP]M)\s*", str(s or ""))
    if not m:
        return None
    h = int(m.group(1)) % 12 + (12 if m.group(3) == "PM" else 0)
    return f"{h:02d}:{m.group(2)}"


def auction_rows(raw: list[dict], fetched: str) -> list[dict]:
    out, seen = [], set()
    for x in raw:
        x = {k: (None if v in ("null", "") else v) for k, v in x.items()}
        ad, cusip = str(x.get("auction_date") or "")[:10], x.get("cusip")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", ad) or (cusip, ad) in seen:
            continue
        seen.add((cusip, ad))
        nominal = x.get("security_type") in ("Note", "Bond") and x.get("inflation_index_security") == "No" and x.get("floating_rate") == "No"
        ct = _closing_time(x.get("closing_time_comp"))
        close = _et(ad, ct) if ct else _end_of_day(ad)
        ann = str(x.get("announcemt_date") or "")[:10]
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", ann) and ann <= ad:
            a_av, basis = min(_end_of_day(ann), close, key=_parse_iso), "end of the announcement day"
        else:
            a_av, basis = close, "auction close (no announcement date)"
        key = f"{cusip}:{ad}"
        sched = {k: x.get(k) for k in ANNOUNCE_FIELDS if k in x}
        out.append(_row("treasury-fiscaldata", "auction-announcement", key, close, _not_after(a_av, fetched),
                        {**sched, "nominal_coupon": nominal, "available_basis": basis}))
        held = any(x.get(k) is not None for k in RESULT_MARKERS)
        if held and _parse_iso(close) <= _parse_iso(fetched):
            r_av = (_parse_iso(close) + timedelta(minutes=AUCTION_RESULT_LAG_MIN)).isoformat() if ct else close
            out.append(_row("treasury-fiscaldata", "auction-result", key, close, _not_after(r_av, fetched), {**x, "nominal_coupon": nominal}))
    return sorted(out, key=lambda r: (r["value"]["auction_date"], str(r["value"].get("cusip")), r["kind"]))


def fetch_treasury_auctions(store=None, first: str = AUCTION_FIRST, types: Iterable[str] = AUCTION_TYPES, dry_run: bool = False,
                            limit: Optional[int] = None, skip_stored: bool = True, refresh: bool = False,
                            client: Optional[httpx.Client] = None, min_interval: float = 1.0, retries: int = 4, backoff: float = 2.0,
                            report: Optional[dict] = None, log=None) -> dict:
    """``treasury-auctions``: every Note/Bond auction since ``first`` (about 2,800 auctions, 3 pages)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    have = _have(archive, AUCTION_BUNDLE) and skip_stored and not refresh
    rep["_plan"] = [_plan("treasury", AUCTION_BUNDLE, 3, 5584, stored=have)]
    if have:
        rep[AUCTION_BUNDLE] = {"skipped": "stored"}
    if dry_run or have or (limit is not None and limit < 1):
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    filt = f"security_type:in:({','.join(types)}),auction_date:gte:{first}"
    raw, page, total_pages, info = [], 1, None, {}
    try:
        while total_pages is None or page <= total_pages:
            r, info = _get(c, FISCAL_AUCTIONS, th, params={"filter": filt, "sort": "auction_date,cusip", "page[size]": AUCTION_PAGE,
                                                           "page[number]": page}, retries=retries, backoff=backoff)
            if r is None:
                break
            try:
                d = r.json()
            except ValueError:
                info["error"] = "not JSON"
                r = None
                break
            raw.extend(d.get("data") or [])
            total_pages = int((d.get("meta") or {}).get("total-pages") or 1)
            page += 1
    finally:
        if own:
            c.close()
    fetched = _stamp()
    if r is None or not raw:
        rep[AUCTION_BUNDLE] = {"error": _fail(info) or "no rows"}
        return rep
    rows = auction_rows(raw, fetched)
    n_res = sum(1 for x in rows if x["kind"] == "auction-result")
    rep[AUCTION_BUNDLE] = {"auctions": len(rows) - n_res, "results": n_res, "first": rows[0]["value"]["auction_date"],
                           "last": rows[-1]["value"]["auction_date"], "pages": page - 1}
    if archive is None:
        rep[AUCTION_BUNDLE]["rows"] = rows
        return rep
    e = _put(archive, AUCTION_BUNDLE, rows, _recipe(
        f"GET {FISCAL_AUCTIONS}?filter={filt}&sort=auction_date,cusip&page[size]={AUCTION_PAGE}&page[number]=<k> via "
        f"mandel.data.public_history.fetch_treasury_auctions; {min_interval}s pacing. Two rows per auction: the schedule "
        "(available at the end of the announcement day, UTC) and, once held, the full result row (available 30 minutes after "
        "the competitive close in New York time, 13:00 ET for coupons). Nominal 10/30-year set: nominal_coupon true and "
        "security_term starting 9-Year/10-Year/29-Year/30-Year (reopenings carry x-Year y-Month terms).", "treasury", True),
        ["treasury", "auctions", "calendar", "rates", "T-442"], AUCTION_SCHEMA, in_git=True)
    rep[AUCTION_BUNDLE]["stored"] = _stored_info(e)
    return rep


# ======================================================================== 3. Cleveland Fed nowcasts (T-360) ==

CLE_BASE = "https://www.clevelandfed.org/-/media/files/webcharts/inflationnowcasting/"
CLE_FILES = {"month": ("nowcast_month.json", "mom"), "quarter": ("nowcast_quarter.json", "qoq-saar")}
CLE_BUNDLE = "clevelandfed-inflation-nowcast"
CLE_UPDATE_ET = "12:00"                  # FAQ: 'every business day around 10:00 a.m. Eastern time' + 2 hours of margin
CLE_ACTUAL_ET = "08:30"                  # BLS CPI and BEA PCE release time
CLE_SNAPSHOT_TARGETS = 3
CLE_SCHEMA = ("kind nowcast: value = {file (month|quarter), target (YYYY-MM or YYYY:Qn), label_date, series (CPI Inflation, Core CPI "
              "Inflation, PCE Inflation, Core PCE Inflation), measure (mom = month-over-month percent; qoq-saar = quarterly annualised "
              "percent), value, file_stamp, vintage_verified false}; observed_at = available_at = label day 12:00 ET (capped at "
              "capture) | kind actual: same fields for the 'Actual ...' series (the print, plotted at the release-day label); "
              "available_at = that label day 08:30 ET")


def _label_date(label: str, target: str) -> Optional[str]:
    """'08/20' for target '2013-7' or '2026:Q4' -> the ISO date nearest the target period (labels carry no year)."""
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})", label.strip())
    if not m:
        return None
    mm, dd = int(m.group(1)), int(m.group(2))
    t = re.match(r"(\d{4})[-:]Q?(\d{1,2})", target)
    if not t:
        return None
    y, k = int(t.group(1)), int(t.group(2))
    mid = date(y, (k - 1) * 3 + 2, 15) if ":Q" in target else date(y, k, 15)
    best = None
    for yy in (y - 1, y, y + 1):
        try:
            d = date(yy, mm, dd)
        except ValueError:
            continue
        if best is None or abs((d - mid).days) < abs((best - mid).days):
            best = d
    return best.isoformat() if best else None


def parse_nowcast(payload: list, file: str, measure: str, fetched: str, targets: Optional[set] = None) -> list[dict]:
    out = []
    for el in payload or []:
        ch = el.get("chart") or {}
        raw_target = str(ch.get("subcaption") or "")
        if targets is not None and raw_target not in targets:
            continue
        m = re.fullmatch(r"(\d{4})-(\d{1,2})", raw_target)
        target = f"{m.group(1)}-{int(m.group(2)):02d}" if m else raw_target
        cats = ((el.get("categories") or [{}])[0] or {}).get("category") or []
        labels = [c.get("label") for c in cats if not c.get("vline")]
        for ds in el.get("dataset") or []:
            sname = str(ds.get("seriesname") or "")
            actual = sname.lower().startswith("actual")
            for lab, pt in zip(labels, ds.get("data") or []):
                v = _num((pt or {}).get("value"))
                day = _label_date(lab or "", raw_target)
                if v is None or day is None:
                    continue
                av = _et(day, CLE_ACTUAL_ET if actual else CLE_UPDATE_ET)
                out.append(_row("clevelandfed", "actual" if actual else "nowcast", f"{file}:{target}:{sname}:{day}", _not_after(av, fetched),
                                _not_after(av, fetched), {"file": file, "target": target, "label_date": day, "series": sname,
                                                          "measure": measure, "value": v, "file_stamp": ch.get("_comment"),
                                                          "vintage_verified": False}))
    return sorted(out, key=lambda r: (r["value"]["file"], r["value"]["target"], r["value"]["label_date"], r["value"]["series"]))


def _cle_fetch(c, th, retries, backoff) -> tuple[dict, dict]:
    got, infos = {}, {}
    for file, (fname, measure) in CLE_FILES.items():
        r, info = _get(c, CLE_BASE + fname, th, retries=retries, backoff=backoff)
        infos[file] = _fail(info) if r is None else {"bytes": len(r.content)}
        if r is not None:
            try:
                got[file] = (json.loads(r.content), measure)
            except ValueError:
                infos[file] = {"error": "not JSON"}
    return got, infos


def fetch_clevelandfed_nowcast(store=None, dry_run: bool = False, limit: Optional[int] = None, skip_stored: bool = True,
                               refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 2.0, retries: int = 4,
                               backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``clevelandfed-inflation-nowcast``: every daily nowcast and actual in the two current files (2 calls, ~13 MB)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    have = _have(archive, CLE_BUNDLE) and skip_stored and not refresh
    rep["_plan"] = [_plan("clevelandfed", CLE_BUNDLE, 2, 39420, stored=have)]
    if have:
        rep[CLE_BUNDLE] = {"skipped": "stored"}
    if dry_run or have or (limit is not None and limit < 1):
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        got, infos = _cle_fetch(c, th, retries, backoff)
    finally:
        if own:
            c.close()
    fetched = _stamp()
    if len(got) < len(CLE_FILES):
        rep[CLE_BUNDLE] = {"error": infos}
        return rep
    rows = [r for file, (payload, measure) in got.items() for r in parse_nowcast(payload, file, measure, fetched)]
    targets = sorted({r["value"]["target"] for r in rows if r["value"]["file"] == "month"})
    rep[CLE_BUNDLE] = {"rows": len(rows), "month_targets": len(targets), "first": targets[0] if targets else None,
                       "last": targets[-1] if targets else None}
    if archive is None:
        rep[CLE_BUNDLE]["rows"] = rows
        return rep
    e = _put(archive, CLE_BUNDLE, rows, _recipe(
        f"GET {CLE_BASE}nowcast_month.json and nowcast_quarter.json (FusionCharts JSON: one element per target period, "
        "category labels MM/DD without year, vlines marking releases) via mandel.data.public_history.fetch_clevelandfed_nowcast. "
        "Label year = the year that puts the label nearest the target period. available_at: nowcasts at the label day 12:00 ET "
        "(the FAQ: updated every business day around 10:00 a.m. ET); 'Actual' series at the label day 08:30 ET (the print is "
        "plotted at its release day). The file is regenerated daily, so past daily values are not proven real-time vintages "
        "(vintage_verified false); fetch_clevelandfed_snapshot stores dated vintages going forward.", "clevelandfed", False),
        ["inflation", "nowcast", "macro", "clevelandfed", "T-360"], CLE_SCHEMA, in_git=False)
    rep[CLE_BUNDLE]["stored"] = _stored_info(e)
    return rep


def fetch_clevelandfed_snapshot(store=None, day: Optional[str] = None, targets: int = CLE_SNAPSHOT_TARGETS, dry_run: bool = False,
                                limit: Optional[int] = None, skip_stored: bool = True, refresh: bool = False,
                                client: Optional[httpx.Client] = None, min_interval: float = 2.0, retries: int = 4, backoff: float = 2.0,
                                report: Optional[dict] = None, log=None) -> dict:
    """``clevelandfed-nowcast-snapshot-<day>``: the last ``targets`` target months/quarters as captured now (available_at =
    capture time): run daily to accumulate verifiable real-time vintages."""
    archive = _archive(store)
    rep = report if report is not None else {}
    day = day or _today()
    name = f"clevelandfed-nowcast-snapshot-{day}"
    have = _have(archive, name) and skip_stored and not refresh
    rep["_plan"] = [_plan("clevelandfed_snap", name, 2, 960, stored=have)]
    if have:
        rep[name] = {"skipped": "stored"}
    if dry_run or have or (limit is not None and limit < 1):
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        got, infos = _cle_fetch(c, th, retries, backoff)
    finally:
        if own:
            c.close()
    fetched = _stamp()
    if len(got) < len(CLE_FILES):
        rep[name] = {"error": infos}
        return rep
    rows = []
    for file, (payload, measure) in got.items():
        keep = {str((el.get("chart") or {}).get("subcaption")) for el in payload[-targets:]}
        for r in parse_nowcast(payload, file, measure, fetched, targets=keep):
            r["available_at"] = fetched
            r["value"]["vintage_captured_at"] = fetched
            rows.append(r)
    rep[name] = {"rows": len(rows)}
    if archive is None:
        rep[name]["rows"] = rows
        return rep
    e = _put(archive, name, rows, _recipe(
        f"GET {CLE_BASE}nowcast_month.json and nowcast_quarter.json via mandel.data.public_history.fetch_clevelandfed_snapshot, "
        f"keeping the last {targets} target periods of each file; available_at = capture time {fetched} (a dated vintage).",
        "clevelandfed", False), ["inflation", "nowcast", "snapshot", "clevelandfed", "T-360", day], CLE_SCHEMA, in_git=False)
    rep[name]["stored"] = _stored_info(e)
    return rep


# =========================================================================== 4. SPDR gold trust (T-404, T-370) ==

SPDR_URL = "https://api.spdrgoldshares.com/api/v1/historical-archive"
SPDR_PRODUCTS = {"gld": 5509, "gldm": 2090}            # rows after dropping 'US Holiday' (gld measured 2026-10-15)
SPDR_FIELDS = {"date": "day", "closing price": "closing_price", "ounces of gold per share": "ounces_per_share",
               "nav/share at 10:30am nyt": "nav_per_share", "indicative price per share at 4:15pm nyt": "indicative_price_415pm",
               "mid point of bid/ask spread at 4:15pm nyt": "mid_415pm", "daily share volume": "volume",
               "total ounces of gold in the trust": "total_ounces", "tonnes of gold": "tonnes",
               "total net asset value in the trust": "total_nav"}
SPDR_SCHEMA = ("kind trust-day: value = {product, day, closing_price, ounces_per_share, nav_per_share (NAV/Share at 10:30am NYT), "
               "indicative_price_415pm, mid_415pm, premium_discount_415pm (percent), volume, total_ounces, tonnes, total_nav}; "
               "'US Holiday' rows dropped; observed_at = day 16:00 ET; available_at = next calendar day 09:30 ET")


def _slug(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", h.lower()).strip("_")


def parse_spdr_xlsx(content: bytes, product: str) -> tuple[list[dict], int]:
    """(daily rows, holiday rows dropped) from the archive workbook's '<...> Historical Archive' sheet."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = next((s for s in wb.sheetnames if "historical archive" in s.lower()), wb.sheetnames[-1])
        header, out, holidays = None, [], 0
        for row in wb[sheet].iter_rows(values_only=True):
            if header is None:
                if row and str(row[0] or "").strip().lower() == "date":
                    header = [str(h or "").strip() for h in row]
                continue
            if not row or row[0] is None:
                continue
            if any(isinstance(x, str) and "holiday" in x.lower() for x in row[1:]):
                holidays += 1
                continue
            d = row[0]
            if isinstance(d, datetime):
                day = d.date().isoformat()
            else:
                try:
                    day = datetime.strptime(str(d).strip(), "%d-%b-%Y").date().isoformat()
                except ValueError:
                    continue
            v = {"product": product.upper(), "day": day}
            for h, x in zip(header[1:], row[1:]):
                key = SPDR_FIELDS.get(h.lower()) or ("premium_discount_415pm" if h.lower().startswith("premium/discount") else _slug(h))
                v[key] = _num(x)
            out.append(v)
        return sorted(out, key=lambda v: v["day"]), holidays
    finally:
        wb.close()


def fetch_spdr_archive(store=None, products: Iterable[str] = tuple(SPDR_PRODUCTS), dry_run: bool = False, limit: Optional[int] = None,
                       skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 2.0,
                       retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``spdr-<product>-archive`` for GLD (2004-11-18..) and GLDM (2018-06-26..): one xlsx call each."""
    archive = _archive(store)
    rep = report if report is not None else {}
    todo, plan = [], []
    for p in products:
        name = f"spdr-{p}-archive"
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            plan.append(_plan("spdr", name, 0, SPDR_PRODUCTS.get(p, 3000), stored=True))
        else:
            todo.append(p)
    todo = todo[:limit] if limit is not None else todo
    plan += [_plan("spdr", f"spdr-{p}-archive", 1, SPDR_PRODUCTS.get(p, 3000)) for p in todo]
    rep["_plan"] = plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        for p in todo:
            name = f"spdr-{p}-archive"
            r, info = _get(c, SPDR_URL, th, params={"product": p, "exchange": "NYSE", "lang": "en"}, retries=retries, backoff=backoff)
            fetched = _stamp()
            if r is None or "spreadsheetml" not in r.headers.get("content-type", ""):
                rep[name] = {"error": _fail(info) if r is None else f"not an xlsx: {r.headers.get('content-type')}"}
                continue
            try:
                vals, hol = parse_spdr_xlsx(r.content, p)
            except Exception as e:                      # a malformed workbook is a failed fetch, never stored
                rep[name] = {"error": f"parse: {e!r}"[:200]}
                continue
            if not vals:
                rep[name] = {"error": "no rows"}
                continue
            rows = [_row("spdr-gold", "trust-day", f"{p}:{v['day']}", _et(v["day"], "16:00"), _not_after(_next_morning(v["day"]), fetched), v)
                    for v in vals]
            lm = r.headers.get("last-modified")
            rep[name] = {"rows": len(rows), "holidays_dropped": hol, "first": vals[0]["day"], "last": vals[-1]["day"], "last_modified": lm}
            if archive is None:
                rep[name]["rows"] = rows
                continue
            e = _put(archive, name, rows, _recipe(
                f"GET {SPDR_URL}?product={p}&exchange=NYSE&lang=en (xlsx, sheet 'US {p.upper()} Historical Archive'; Last-Modified "
                f"{lm}) via mandel.data.public_history.fetch_spdr_archive; {hol} 'US Holiday' rows dropped. available_at = next "
                "calendar day 09:30 ET (t+1: holdings and NAV are published after the close; the file is rewritten about 02:10 UTC).",
                "spdr", False), ["gold", "etf", "holdings", "nav", "spdr", p, "T-404", "T-370"], SPDR_SCHEMA, in_git=False)
            rep[name]["stored"] = _stored_info(e)
    finally:
        if own:
            c.close()
    return rep


# ============================================================================ 5. Sprott physical trusts (T-472) ==

SPROTT_URL = "https://sprott.com/investment-strategies/exchange-listed-products/physical-bullion-funds/{page}/"
SPROTT_FUNDS = {"phys": ("gold", "PHYS"), "pslv": ("silver", "PSLV"), "sppp": ("platinum-and-palladium", "SPPP"),
                "cef": ("gold-and-silver", "CEF")}
SPROTT_SERIES = {"Share Price": "share_price", "Net Asset Value": "nav", "Premium/Discount to Net Asset Value": "premium_discount_pct"}
SPROTT_ROWS_EST = {"phys": 4229, "pslv": 4000, "sppp": 3500, "cef": 2200}
SPROTT_SCHEMA = ("kind trust-day: value = {fund, day, share_price, nav, premium_discount_pct} from the page's Highcharts series "
                 "'Share Price', 'Net Asset Value' and 'Premium/Discount to Net Asset Value' (USD per the page tooltips; percent); "
                 "dates include TSX-only sessions; observed_at = day 16:00 ET; available_at = next calendar day 09:30 ET")
_HC_NAME = re.compile(r"name\s*:\s*'([^']{2,120})'")
_HC_POINT = re.compile(r"\[\s*Date\.UTC\(\s*(\d{4})\s*,\s*(\d{1,2})\s*,\s*(\d{1,2})\s*\)\s*,\s*(-?[\d.]+|null)\s*\]")


def parse_highcharts(page: str) -> dict[str, dict[str, float]]:
    """{series name: {ISO day: value}} for every Highcharts series whose data are [Date.UTC(y, m0, d), v] pairs (month 0-based)."""
    out: dict[str, dict[str, float]] = {}
    names = list(_HC_NAME.finditer(page))
    for i, m in enumerate(names):
        stop = names[i + 1].start() if i + 1 < len(names) else len(page)
        j = page.find("data:", m.end(), stop)
        if j < 0:
            continue
        k = page.find("]]", j, stop)
        seg = page[j:(k + 2) if k >= 0 else stop]
        pts = {}
        for y, mo, d, v in _HC_POINT.findall(seg):
            if v == "null":
                continue
            try:
                pts[date(int(y), int(mo) + 1, int(d)).isoformat()] = float(v)
            except ValueError:
                continue
        if pts:
            out.setdefault(m.group(1), {}).update(pts)
    return out


def sprott_rows(series: dict[str, dict[str, float]], fund: str, fetched: str) -> list[dict]:
    days = sorted({d for nm in SPROTT_SERIES for d in series.get(nm, {})})
    out = []
    for d in days:
        v = {"fund": fund, "day": d, **{f: series[nm][d] for nm, f in SPROTT_SERIES.items() if d in series.get(nm, {})}}
        out.append(_row("sprott", "trust-day", f"{fund}:{d}", _et(d, "16:00"), _not_after(_next_morning(d), fetched), v))
    return out


def fetch_sprott_trusts(store=None, funds: Iterable[str] = tuple(SPROTT_FUNDS), dry_run: bool = False, limit: Optional[int] = None,
                        skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 2.0,
                        retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``sprott-<fund>-nav`` for PHYS, PSLV, SPPP and CEF: one page (~0.5-0.7 MB) each."""
    archive = _archive(store)
    rep = report if report is not None else {}
    todo, plan = [], []
    for f in funds:
        name = f"sprott-{f}-nav"
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            plan.append(_plan("sprott", name, 0, SPROTT_ROWS_EST.get(f, 3000), stored=True))
        else:
            todo.append(f)
    todo = todo[:limit] if limit is not None else todo
    plan += [_plan("sprott", f"sprott-{f}-nav", 1, SPROTT_ROWS_EST.get(f, 3000)) for f in todo]
    rep["_plan"] = plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        for f in todo:
            name = f"sprott-{f}-nav"
            page, ticker = SPROTT_FUNDS[f]
            url = SPROTT_URL.format(page=page)
            r, info = _get(c, url, th, retries=retries, backoff=backoff)
            fetched = _stamp()
            if r is None:
                rep[name] = {"error": _fail(info)}
                continue
            series = parse_highcharts(r.text)
            rows = sprott_rows(series, ticker, fetched)
            missing = [nm for nm in SPROTT_SERIES if nm not in series]
            if not rows or missing:
                rep[name] = {"error": f"series missing on the page: {missing}" if missing else "no rows"}
                continue
            rep[name] = {"rows": len(rows), "first": rows[0]["value"]["day"], "last": rows[-1]["value"]["day"]}
            if archive is None:
                rep[name]["rows"] = rows
                continue
            e = _put(archive, name, rows, _recipe(
                f"GET {url} and parse the embedded Highcharts series {list(SPROTT_SERIES)} ([Date.UTC(y, m0, d), v] pairs) via "
                "mandel.data.public_history.fetch_sprott_trusts. NAV is struck after the close: available_at = next calendar day "
                "09:30 ET. Dates include TSX-only sessions: align to the NYSE calendar.", "sprott", False),
                ["bullion", "trust", "nav", "premium", "sprott", f, "T-472"], SPROTT_SCHEMA, in_git=False)
            rep[name]["stored"] = _stored_info(e)
    finally:
        if own:
            c.close()
    return rep


# =================================================================================== 6. SSGA SPY holdings (T-374) ==

SSGA_URL = "https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/holdings-daily-us-en-{fund}.xlsx"
SSGA_SCHEMA = ("kind holding: value = {fund, as_of, name, ticker, identifier (CUSIP), sedol, weight (percent), sector, shares_held, "
               "currency}; observed_at = as-of day 16:00 ET; available_at = the file's Last-Modified (capture time when absent), not "
               "before the as-of close")


def parse_ssga_xlsx(content: bytes) -> tuple[dict, list[dict]]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        meta, header, out = {}, None, []
        for row in wb[wb.sheetnames[0]].iter_rows(values_only=True):
            if header is None:
                k = str(row[0] or "").strip().rstrip(":").lower() if row else ""
                if k in ("fund name", "ticker symbol", "holdings"):
                    meta[k] = str(row[1] or "").strip()
                elif k == "name" and row and len(row) > 1 and str(row[1] or "").strip().lower() == "ticker":
                    header = [_slug(str(h or "")) for h in row]
                continue
            if not row or row[0] is None or not any(row[1:4]):
                if out:
                    break
                continue
            v = {h: (_num(x) if h in ("weight", "shares_held") else (str(x).strip() if x is not None else None))
                 for h, x in zip(header, row) if h}
            out.append(v)
        m = re.search(r"(\d{1,2}-[A-Za-z]{3}-\d{4})", meta.get("holdings", ""))
        meta["as_of"] = datetime.strptime(m.group(1), "%d-%b-%Y").date().isoformat() if m else None
        return meta, out
    finally:
        wb.close()


def fetch_ssga_holdings(store=None, fund: str = "spy", dry_run: bool = False, limit: Optional[int] = None, skip_stored: bool = True,
                        refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 2.0, retries: int = 4,
                        backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``ssga-<fund>-holdings-<as-of>``: one dated snapshot of the daily holdings file (1 call; the as-of date is read
    from the file, so the call is made even when that snapshot is stored, and nothing is re-stored then)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    rep["_plan"] = [_plan("ssga", f"ssga-{fund}-holdings-<as-of>", 1, 505)]
    if dry_run or (limit is not None and limit < 1):
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        r, info = _get(c, SSGA_URL.format(fund=fund), th, retries=retries, backoff=backoff)
    finally:
        if own:
            c.close()
    fetched = _stamp()
    if r is None or "spreadsheetml" not in r.headers.get("content-type", ""):
        rep[f"ssga-{fund}-holdings"] = {"error": _fail(info) if r is None else f"not an xlsx: {r.headers.get('content-type')}"}
        return rep
    meta, vals = parse_ssga_xlsx(r.content)
    if not meta.get("as_of") or not vals:
        rep[f"ssga-{fund}-holdings"] = {"error": "no as-of date or no holdings"}
        return rep
    name = f"ssga-{fund}-holdings-{meta['as_of']}"
    if _have(archive, name) and skip_stored and not refresh:
        rep[name] = {"skipped": "stored"}
        return rep
    lm = r.headers.get("last-modified")
    try:
        lm_iso = _iso(datetime.strptime(lm, "%a, %d %b %Y %H:%M:%S GMT").replace(tzinfo=UTC)) if lm else None
    except ValueError:
        lm_iso = None
    close = _et(meta["as_of"], "16:00")
    av = max(_not_after(lm_iso or fetched, fetched), close, key=_parse_iso)
    rows = [_row("ssga", "holding", f"{fund}:{meta['as_of']}:{v.get('ticker') or v.get('identifier')}", close, av,
                 {"fund": fund.upper(), "as_of": meta["as_of"], **v}) for v in vals]
    rep[name] = {"rows": len(rows), "as_of": meta["as_of"], "last_modified": lm}
    if archive is None:
        rep[name]["rows"] = rows
        return rep
    e = _put(archive, name, rows, _recipe(
        f"GET {SSGA_URL.format(fund=fund)} (xlsx; '{meta.get('fund name')}', holdings as of {meta['as_of']}; Last-Modified {lm}) via "
        "mandel.data.public_history.fetch_ssga_holdings; available_at = Last-Modified (a dated snapshot).", "ssga", False),
        ["holdings", "etf", "constituents", "ssga", fund, "T-374", meta["as_of"]], SSGA_SCHEMA, in_git=False)
    rep[name]["stored"] = _stored_info(e)
    return rep


# ================================================================== 7. Wikipedia S&P 500 (T-374, T-394, T-473, T-334) ==

WIKI = "https://en.wikipedia.org/wiki/"
WIKI_PAGES = {"constituents": ("List_of_S%26P_500_companies", "constituents", "wikipedia-sp500-constituents"),
              "changes": ("Historical_components_of_the_S%26P_500", "changes", "wikipedia-sp500-changes")}
WIKI_CONST_SCHEMA = ("kind constituent: value = {symbol, security, gics_sector, gics_sub_industry, headquarters, date_added, cik, founded, "
                     "revid, revision_time}; observed_at = available_at = the revision time (a dated snapshot of current members)")
WIKI_CHANGES_SCHEMA = ("kind index-change: value = {effective_date, added_ticker, added_security, removed_ticker, removed_security, reason, "
                       "refs [{note, title, published, publisher}], announced (publication date of a cited index announcement, a "
                       "title naming an S&P index and a join/replace/add/remove, 1-30 days before the effective date, else None), "
                       "available_basis, revid}; observed_at = effective day 09:30 ET; available_at = end of "
                       "the announced day (UTC) when known, else the effective day 09:30 ET")
_MONTHS = {m: i for i, m in enumerate(["january", "february", "march", "april", "may", "june", "july", "august", "september",
                                       "october", "november", "december"], 1)}


class _TableParser(HTMLParser):
    """Rows of the HTML tables with the given ids: [[{text, links, notes, rowspan, colspan}]] (nested tables ignored)."""

    def __init__(self, ids: Iterable[str]):
        super().__init__(convert_charrefs=True)
        self.ids = set(ids)
        self.tables: dict[str, list[list[dict]]] = {}
        self._stack: list[Optional[str]] = []          # table ids (None for tables we do not want)
        self._row: Optional[list] = None
        self._cell: Optional[dict] = None
        self._skip = 0

    def _target(self) -> Optional[str]:
        return self._stack[-1] if self._stack and self._stack[-1] in self.ids else None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "table":
            self._stack.append(a.get("id") if a.get("id") in self.ids and not any(s in self.ids for s in self._stack) else None)
            if self._target():
                self.tables.setdefault(self._target(), [])
            return
        t = self._target()
        if not t:
            return
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = {"text": "", "links": [], "notes": [], "header": tag == "th",
                          "rowspan": int(a.get("rowspan") or 1) if str(a.get("rowspan") or "1").isdigit() else 1,
                          "colspan": int(a.get("colspan") or 1) if str(a.get("colspan") or "1").isdigit() else 1}
        elif tag in ("style", "script"):
            self._skip += 1
        elif tag == "sup" and self._cell is not None:
            self._skip += 1
        elif tag == "a" and self._cell is not None:
            href = a.get("href") or ""
            if href.startswith("#cite_note"):
                self._cell["notes"].append(href[1:])
            elif not self._skip:
                self._cell["links"].append(href)
        elif tag == "br" and self._cell is not None:
            self._cell["text"] += " "

    def handle_endtag(self, tag):
        if tag == "table":
            if self._stack:
                self._stack.pop()
            return
        if not self._target():
            return
        if tag in ("style", "script", "sup") and self._skip:
            self._skip -= 1
        elif tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._cell["text"] = re.sub(r"\s+", " ", self._cell["text"]).strip()
            self._row.append(self._cell)
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.tables[self._target()].append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None and not self._skip:
            self._cell["text"] += data


def table_grid(rows: list[list[dict]]) -> list[list[dict]]:
    """Expand rowspan/colspan into a rectangular grid of cells."""
    grid, pending = [], {}
    for row in rows:
        out, col, cells = [], 0, list(row)
        while cells or col in pending:
            if col in pending:
                cell, left = pending[col]
                out.append(cell)
                if left <= 1:
                    del pending[col]
                else:
                    pending[col] = (cell, left - 1)
                col += 1
                continue
            cell = cells.pop(0)
            for _ in range(cell["colspan"]):
                out.append(cell)
                if cell["rowspan"] > 1:
                    pending[col] = (cell, cell["rowspan"] - 1)
                col += 1
        grid.append(out)
    return grid


def _parse_day(s: str) -> Optional[str]:
    s = re.sub(r"\s+", " ", (s or "").replace(",", " ")).strip()
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.search(r"([A-Za-z]+)\.? (\d{1,2}) (\d{4})", s)
    if m and m.group(1).lower()[:3] in {k[:3] for k in _MONTHS}:
        mo = next(v for k, v in _MONTHS.items() if k[:3] == m.group(1).lower()[:3])
        try:
            return date(int(m.group(3)), mo, int(m.group(2))).isoformat()
        except ValueError:
            return None
    m = re.search(r"(\d{1,2}) ([A-Za-z]+) (\d{4})", s)
    if m and m.group(2).lower()[:3] in {k[:3] for k in _MONTHS}:
        mo = next(v for k, v in _MONTHS.items() if k[:3] == m.group(2).lower()[:3])
        try:
            return date(int(m.group(3)), mo, int(m.group(1))).isoformat()
        except ValueError:
            return None
    return None


class _Resp:
    """The few httpx.Response attributes the Wikipedia parsers use."""

    def __init__(self, status: int, headers: dict, content: bytes):
        self.status_code, self.headers, self.content = status, httpx.Headers(headers), content
        self.text = content.decode("utf-8", "replace")


def _wiki_get(url: str, throttle: Throttle, retries: int = 4, backoff: float = 3.0) -> tuple[Optional[_Resp], dict]:
    """GET through urllib: on 2026-10-15 Wikimedia answered 403 'Please respect our robot policy' to httpx from this host
    for every User-Agent tried, while urllib and curl with the same User-Agent got 200. Same retry rules as
    options_misc._get (429/5xx and transport errors with exponential backoff, Retry-After honoured; other answers final)."""
    import urllib.error
    import urllib.request
    info: dict = {"url": url, "attempts": 0}
    for attempt in range(retries):
        throttle.wait()
        info["attempts"] = attempt + 1
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=WIKI_UA), timeout=120) as r:
                info["status"] = r.status
                return _Resp(r.status, dict(r.headers.items()), r.read()), info
        except urllib.error.HTTPError as e:
            info["status"] = e.code
            body = e.read()[:300].decode("utf-8", "replace")
            if e.code in RETRY_STATUS:
                ra = e.headers.get("retry-after")
                try:
                    ra = float(ra) if ra else None
                except ValueError:
                    ra = None
                if ra is not None and ra > 120:
                    info.update(retry_after=ra, body=body)
                    return None, info
                _sleep(max(backoff * 2 ** attempt, ra or 0.0))
                continue
            info["body"] = body
            return None, info
        except (urllib.error.URLError, OSError) as e:
            info["error"] = repr(e)[:200]
            _sleep(backoff * 2 ** attempt)
    return None, info


def wiki_revision(page: str, headers: Optional[httpx.Headers] = None) -> dict:
    out: dict = {}
    m = re.search(r'"wgRevisionId":(\d+)', page)
    out["revid"] = int(m.group(1)) if m else None
    m = re.search(r"This page was last edited on (\d{1,2} [A-Za-z]+ \d{4}), at (\d{2}:\d{2})", page)
    if m:
        d = _parse_day(m.group(1))
        out["revision_time"] = f"{d}T{m.group(2)}:00+00:00" if d else None
    if not out.get("revision_time") and headers is not None and headers.get("last-modified"):
        try:
            out["revision_time"] = _iso(datetime.strptime(headers["last-modified"], "%a, %d %b %Y %H:%M:%S GMT").replace(tzinfo=UTC))
        except ValueError:
            pass
    return out


def wiki_references(page: str) -> dict[str, dict]:
    """{cite_note id: {title, published, publisher}} from the page's reference list."""
    out = {}
    for m in re.finditer(r'<li[^>]*id="(cite_note-[^"]+)"[^>]*>(.*?)</li>', page, flags=re.S):
        text = re.sub(r"\s+", " ", htmlmod.unescape(re.sub(r"<[^>]+>", " ", re.sub(r"(?s)<style.*?</style>", "", m.group(2))))).strip()
        title = re.search(r'"([^"]{3,300})"', text)
        before = text.split("Retrieved")[0]
        pub = None
        for cand in re.findall(r"(\d{4}-\d{2}-\d{2}|[A-Z][a-z]+\.? \d{1,2}, \d{4}|\d{1,2} [A-Z][a-z]+ \d{4})", before):
            pub = _parse_day(cand) or pub
        publisher = re.search(r"\.\s*([A-Z][A-Za-z&' .]{2,60}?)\s*\.\s*(?:[A-Z][a-z]+\.? \d{1,2}, \d{4}|\d{4}-\d{2}-\d{2})", before)
        out[m.group(1)] = {"title": title.group(1) if title else None, "published": pub,
                           "publisher": publisher.group(1).strip() if publisher else None}
    return out


def parse_wiki_constituents(page: str, fetched: str, headers=None) -> tuple[list[dict], dict]:
    rev = wiki_revision(page, headers)
    p = _TableParser(["constituents"])
    p.feed(page)
    grid = table_grid(p.tables.get("constituents", []))
    if not grid:
        return [], rev
    names = [_slug(c["text"]) for c in grid[0]]
    col = {"symbol": "symbol", "security": "security", "gics_sector": "gics_sector", "gics_sub_industry": "gics_sub_industry",
           "headquarters_location": "headquarters", "date_added": "date_added", "cik": "cik", "founded": "founded"}
    t = rev.get("revision_time") or fetched
    out = []
    for row in grid[1:]:
        if all(c["header"] for c in row):
            continue
        v = {col.get(n, n): c["text"] for n, c in zip(names, row)}
        if not v.get("symbol"):
            continue
        v["date_added"] = _parse_day(v.get("date_added") or "") or v.get("date_added")
        v.update(revid=rev.get("revid"), revision_time=rev.get("revision_time"))
        out.append(_row("wikipedia", "constituent", v["symbol"], t, _not_after(t, fetched), v))
    return out, rev


_INDEX_NEWS = re.compile(r"(?i)S&P\s*(500|MidCap|SmallCap|100|Composite\s*1500).*\b(join|joins|replace|replaces|added|add|adds|remove|removed|"
                         r"drop|dropped|deleted|change|changes|rebalanc)|\b(join|joins|replace|replaces|added|add|adds|remove|removed|drop|"
                         r"dropped|deleted)\b.*S&P\s*(500|MidCap|SmallCap|100)")


def parse_wiki_changes(page: str, fetched: str, headers=None) -> tuple[list[dict], dict]:
    rev = wiki_revision(page, headers)
    p = _TableParser(["changes"])
    p.feed(page)
    grid = table_grid(p.tables.get("changes", []))
    refs = wiki_references(page)
    out, n = [], defaultdict(int)
    for row in grid:
        if any(c["header"] for c in row) or len(row) < 6:
            continue
        eff = _parse_day(row[0]["text"])
        if not eff:
            continue
        cells = [c["text"] or None for c in row]
        notes = [nid for c in row[5:] for nid in c["notes"]]
        rlist = [{"note": nid, **refs.get(nid, {})} for nid in dict.fromkeys(notes)]
        cands = [r["published"] for r in rlist if r.get("published") and _INDEX_NEWS.search(r.get("title") or "")
                 and 1 <= (date.fromisoformat(eff) - date.fromisoformat(r["published"])).days <= 30]
        ann = min(cands) if cands else None
        opn = _et(eff, "09:30")
        av, basis = ((min(_end_of_day(ann), opn, key=_parse_iso), "end of the announced day (cited source date)") if ann
                     else (opn, "effective day 09:30 ET (announcement date not cited)"))
        key = f"{eff}:{cells[1] or ''}:{cells[3] or ''}"
        n[key] += 1
        v = {"effective_date": eff, "added_ticker": cells[1], "added_security": cells[2], "removed_ticker": cells[3],
             "removed_security": cells[4], "reason": cells[5] if len(cells) > 5 else None, "refs": rlist, "announced": ann,
             "available_basis": basis, "revid": rev.get("revid")}
        out.append(_row("wikipedia", "index-change", key + (f"#{n[key]}" if n[key] > 1 else ""), opn, _not_after(av, fetched), v))
    return sorted(out, key=lambda r: (r["value"]["effective_date"], r["key"])), rev


def _wiki_recipe(title: str, rev: dict, what: str) -> str:
    oldid = f"https://en.wikipedia.org/w/index.php?title={title}&oldid={rev.get('revid')}"
    return (f"GET {WIKI}{title} (page HTML; the MediaWiki API answered 429 to this shared host) and parse {what} via "
            f"mandel.data.public_history.fetch_wikipedia_sp500 (urllib: Wikimedia refused httpx here). Attribution: Wikipedia contributors, '{title.replace('_', ' ').replace('%26', '&')}', "
            f"revision {rev.get('revid')} of {rev.get('revision_time')}, {oldid}, licensed CC BY-SA 4.0 "
            "(https://creativecommons.org/licenses/by-sa/4.0/). Changes: table parsed into rows (cells as text, rowspans expanded, "
            "dates normalised to ISO, cited sources reduced to title/date). ShareAlike: this bundle is an adaptation redistributed "
            "under CC BY-SA 4.0; anything shared that is built from it must carry the same licence and this attribution.")


def fetch_wikipedia_sp500(store=None, tables: Iterable[str] = ("constituents", "changes"), dry_run: bool = False,
                          limit: Optional[int] = None, skip_stored: bool = True, refresh: bool = False,
                          client: Optional[httpx.Client] = None, min_interval: float = 3.0, retries: int = 4, backoff: float = 3.0,
                          report: Optional[dict] = None, log=None) -> dict:
    """``wikipedia-sp500-constituents`` and ``wikipedia-sp500-changes``: one page each (in git, CC BY-SA 4.0)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    todo, plan = [], []
    for t in tables:
        title, _, name = WIKI_PAGES[t]
        fam = "wiki_const" if t == "constituents" else "wiki_changes"
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            plan.append(_plan(fam, name, 0, 503 if t == "constituents" else 412, stored=True))
        else:
            todo.append(t)
            plan.append(_plan(fam, name, 1, 503 if t == "constituents" else 412))
    todo = todo[:limit] if limit is not None else todo
    rep["_plan"] = plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        for t in todo:
            title, tid, name = WIKI_PAGES[t]
            r, info = (_get(c, WIKI + title, th, headers=WIKI_UA, retries=retries, backoff=backoff) if client is not None
                       else _wiki_get(WIKI + title, th, retries=retries, backoff=backoff))
            fetched = _stamp()
            if r is None:
                rep[name] = {"error": _fail(info)}
                continue
            rows, rev = (parse_wiki_constituents if t == "constituents" else parse_wiki_changes)(r.text, fetched, r.headers)
            if not rows or not rev.get("revid"):
                rep[name] = {"error": f"table '{tid}' or revision id not found"}
                continue
            rep[name] = {"rows": len(rows), "revid": rev.get("revid"), "revision_time": rev.get("revision_time")}
            if t == "changes":
                rep[name]["announced_known"] = sum(1 for x in rows if x["value"]["announced"])
                rep[name]["first"], rep[name]["last"] = rows[0]["value"]["effective_date"], rows[-1]["value"]["effective_date"]
            if archive is None:
                rep[name]["rows"] = rows
                continue
            e = _put(archive, name, rows, _wiki_recipe(title, rev, f"the table id='{tid}'") + (
                " The table date is the EFFECTIVE date (changes take effect before that day's open); available_at = the end of the "
                "announced day (UTC) when a cited S&P index announcement (title naming an S&P index and a join/replace/add/remove) is "
                "dated 1-30 days before it, else the effective day 09:30 ET."
                if t == "changes" else " available_at = the revision time (a snapshot of current members)."),
                ["sp500", "index", "membership", "wikipedia", "cc-by-sa-4.0", t] + list(FAMILIES["wiki_const" if t == "constituents" else "wiki_changes"][2]),
                WIKI_CONST_SCHEMA if t == "constituents" else WIKI_CHANGES_SCHEMA, in_git=True)
            rep[name]["stored"] = _stored_info(e)
    finally:
        if own:
            c.close()
    return rep


# ============================================================================ 8. fja05680 membership (T-473, T-334) ==

FJA_JSD = "https://cdn.jsdelivr.net/gh/fja05680/sp500@master/"
FJA_RAW = "https://raw.githubusercontent.com/fja05680/sp500/master/"
FJA_COMPONENTS = "S&P 500 Historical Components & Changes (Updated).csv"
FJA_SPELLS = "sp500_ticker_start_end.csv"
FJA_COMP_SCHEMA = ("kind sp500-members: value = {date, tickers (list, the file's spelling e.g. BRK.B), n}; one row per change date "
                   "1996-01-02..; membership holds until the next row; observed_at = available_at = the date 09:30 ET")
FJA_SPELL_SCHEMA = ("kind sp500-join: value = {ticker, date, spell}; kind sp500-leave: value = {ticker, date (first date without the "
                    "ticker), start_date, spell}; observed_at = available_at = the date 09:30 ET (a join row never carries its end)")


def _fja_get(c, th, fname: str, retries: int, backoff: float) -> tuple[Optional[httpx.Response], dict]:
    q = quote(fname, safe="")
    r, info = _get(c, FJA_JSD + q, th, retries=retries, backoff=backoff)
    if r is None:
        r, info2 = _get(c, FJA_RAW + q, th, retries=retries, backoff=backoff)
        info = {"jsdelivr": _fail(info), "raw": _fail(info2)} if r is None else {**info2, "via": "raw.githubusercontent.com"}
    return r, info


def parse_fja_components(text: str, fetched: str) -> list[dict]:
    out = []
    for x in csv.DictReader(io.StringIO(text)):
        d = (x.get("date") or "").strip()
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", d):
            continue
        tk = sorted({t.strip() for t in (x.get("tickers") or "").split(",") if t.strip()})
        t = _et(d, "09:30")
        out.append(_row("fja05680", "sp500-members", d, t, _not_after(t, fetched), {"date": d, "tickers": tk, "n": len(tk)}))
    return sorted(out, key=lambda r: r["key"])


def parse_fja_spells(text: str, fetched: str) -> list[dict]:
    out, spell = [], defaultdict(int)
    for x in csv.DictReader(io.StringIO(text)):
        tk, s, e = (x.get("ticker") or "").strip(), (x.get("start_date") or "").strip(), (x.get("end_date") or "").strip()
        if not tk or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
            continue
        spell[tk] += 1
        t = _et(s, "09:30")
        out.append(_row("fja05680", "sp500-join", f"{tk}:{s}", t, _not_after(t, fetched), {"ticker": tk, "date": s, "spell": spell[tk]}))
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", e):
            t = _et(e, "09:30")
            out.append(_row("fja05680", "sp500-leave", f"{tk}:{e}", t, _not_after(t, fetched),
                            {"ticker": tk, "date": e, "start_date": s, "spell": spell[tk]}))
    return sorted(out, key=lambda r: (r["value"]["date"], r["kind"], r["value"]["ticker"]))


def fetch_fja_sp500(store=None, files: Iterable[str] = ("components", "spells"), dry_run: bool = False, limit: Optional[int] = None,
                    skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 1.5,
                    retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``fja-sp500-components`` (5.5 MB csv) and ``fja-sp500-ticker-spells``: one call each."""
    archive = _archive(store)
    rep = report if report is not None else {}
    spec = {"components": ("fja-sp500-components", FJA_COMPONENTS, "fja_comp", 2720, parse_fja_components, FJA_COMP_SCHEMA),
            "spells": ("fja-sp500-ticker-spells", FJA_SPELLS, "fja_spells", 2021, parse_fja_spells, FJA_SPELL_SCHEMA)}
    todo, plan = [], []
    for f in files:
        name, _, fam, est, _, _ = spec[f]
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            plan.append(_plan(fam, name, 0, est, stored=True))
        else:
            todo.append(f)
            plan.append(_plan(fam, name, 1, est))
    todo = todo[:limit] if limit is not None else todo
    rep["_plan"] = plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        for f in todo:
            name, fname, fam, _, parse, schema = spec[f]
            r, info = _fja_get(c, th, fname, retries, backoff)
            fetched = _stamp()
            if r is None:
                rep[name] = {"error": info}
                continue
            rows = parse(r.text, fetched)
            if not rows:
                rep[name] = {"error": "no rows"}
                continue
            dates = [x["value"]["date"] for x in rows]
            rep[name] = {"rows": len(rows), "first": min(dates), "last": max(dates), "via": info.get("via", "cdn.jsdelivr.net"),
                         "etag": r.headers.get("etag")}
            if archive is None:
                rep[name]["rows"] = rows
                continue
            e = _put(archive, name, rows, _recipe(
                f"GET {FJA_JSD}{quote(fname, safe='')} (fallback {FJA_RAW}...; github.com itself answers 403 here; ETag "
                f"{r.headers.get('etag')}, last date {max(dates)}) via mandel.data.public_history.fetch_fja_sp500. "
                "Repository fja05680/sp500 (MIT licence, (c) 2019-2020 Farrell J. Aultman). Dates are effective dates: available_at = "
                "that day's 09:30 ET (changes are effective before the open and announced earlier). Changes after the file's last "
                "date come from wikipedia-sp500-changes.", "fja", False),
                ["sp500", "index", "membership", "point-in-time", "fja05680", f, "T-473", "T-334", "T-394"], schema, in_git=False)
            rep[name]["stored"] = _stored_info(e)
    finally:
        if own:
            c.close()
    return rep


# ============================================================================ 9. FINRA short interest (T-375, T-393) ==

FINRA_SI = "https://api.finra.org/data/group/otcMarket/name/consolidatedShortInterest"
FINRA_OTC = ("OTC", "OTCBB")
FINRA_LISTED = ("NNM", "NYSE", "ARCA", "SC", "BZX", "AMEX", "TXSE", "IEXG", "NGM", "NGS", "NCM", "NAS")
FINRA_PAGE = 5000
FINRA_FIRST = "2017-12-29"
FINRA_PUBLISH_BDAYS = 8
FINRA_SI_SCHEMA = ("kind short-interest: value = the FINRA row {symbolCode, issueName, marketClassCode, settlementDate, "
                   "currentShortPositionQuantity, previousShortPositionQuantity, changePreviousNumber, changePercent, "
                   "averageDailyVolumeQuantity, daysToCoverQuantity, stockSplitFlag, revisionFlag, issuerServicesGroupExchangeCode} "
                   "(accountingYearMonthNumber dropped: equal to settlementDate); exchange-listed classes only (OTC, OTCBB out); "
                   "observed_at = settlement day 16:00 ET; available_at = end of the 8th NYSE business day after settlement (UTC)")


def nth_business_day(day: str, n: int) -> str:
    d = date.fromisoformat(day[:10])
    for _ in range(n):
        d += timedelta(days=1)
        while not is_trading_day(d):
            d += timedelta(days=1)
    return d.isoformat()


def finra_si_available(settlement: str) -> str:
    return _end_of_day(nth_business_day(settlement, FINRA_PUBLISH_BDAYS))


def finra_si_rows_est(day: str) -> int:
    """Listed rows of one settlement date, interpolated between the measured 8,326 (2017-12-29) and 12,900 (2026-09-30)."""
    a, b = date(2017, 12, 29), date(2026, 9, 30)
    f = min(max((date.fromisoformat(day[:10]) - a).days / (b - a).days, 0.0), 1.2)
    return int(8326 + (12900 - 8326) * f)


def _prev_trading_day_or_same(d: date) -> date:
    while not is_trading_day(d):
        d -= timedelta(days=1)
    return d


def finra_si_calendar(first: str = FINRA_FIRST, today: Optional[str] = None) -> list[str]:
    """Settlement dates by FINRA's rule (mid-month: the 15th, or the business day before it when the 15th is not one;
    month-end: the last business day), from ``first`` to the last one published by ``today`` (7 NYSE business days after
    settlement). The dry-run plan uses it (no request); a real run lists the dates from the API. It reproduces the API's
    211 dates 2017-12-29..2026-09-30 (fetch6c/misc/probes/finra_aapl_dates.json)."""
    today = today or _today()
    out, d = [], date.fromisoformat(first[:10]).replace(day=1)
    while d.isoformat() <= today:
        nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
        for s in (_prev_trading_day_or_same(d.replace(day=15)), _prev_trading_day_or_same(nxt - timedelta(days=1))):
            if s.isoformat() >= first[:10] and nth_business_day(s.isoformat(), FINRA_PUBLISH_BDAYS - 1) <= today:
                out.append(s.isoformat())
        d = nxt
    return out


def _finra_query(c, th, body: dict, retries: int, backoff: float) -> tuple[Optional[list], Optional[int], dict]:
    r, info = _post_json(c, FINRA_SI, th, body, retries=retries, backoff=backoff)
    if r is None:
        return None, None, info
    total = r.headers.get("record-total")
    try:
        data = r.json() if r.status_code == 200 and r.content else []
    except ValueError:
        info["error"] = "not JSON"
        return None, None, info
    return data, int(total) if total and total.isdigit() else None, info


def finra_settlement_dates(c, th, retries: int = 4, backoff: float = 2.0, symbols: Iterable[str] = ("AAPL", "SPY")) -> list[str]:
    out = set()
    for s in symbols:
        data, _, _ = _finra_query(c, th, {"limit": FINRA_PAGE, "offset": 0, "fields": ["settlementDate"],
                                          "compareFilters": [{"compareType": "EQUAL", "fieldName": "symbolCode", "fieldValue": s}]},
                                  retries, backoff)
        out |= {str(x.get("settlementDate"))[:10] for x in data or [] if x.get("settlementDate")}
    return sorted(d for d in out if d >= FINRA_FIRST)


def finra_si_rows(data: list[dict], fetched: str) -> list[dict]:
    out, seen = [], set()
    for x in data:
        sd, sym = str(x.get("settlementDate") or "")[:10], x.get("symbolCode")
        if not sym or not sd or (sd, sym) in seen or x.get("marketClassCode") in FINRA_OTC:
            continue
        seen.add((sd, sym))
        v = {k: val for k, val in x.items() if k != "accountingYearMonthNumber"}
        out.append(_row("finra", "short-interest", f"{sd}:{sym}", _et(sd, "16:00"), _not_after(finra_si_available(sd), fetched), v))
    return sorted(out, key=lambda r: r["key"])


def fetch_finra_short_interest(store=None, dates: Optional[Iterable[str]] = None, dry_run: bool = False, limit: Optional[int] = None,
                               skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None,
                               min_interval: float = 1.0, retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None,
                               log=None) -> dict:
    """``finra-short-interest-<settlement date>``: exchange-listed rows of one settlement date (~13k rows, 3 pages + 2
    counts). Settlement dates come from the AAPL and SPY histories (one call each) unless ``dates`` is given."""
    archive = _archive(store)
    rep = report if report is not None else {}
    if dry_run:                          # no request: the rule's calendar stands in for the API's date list
        ds = sorted(dates) if dates is not None else finra_si_calendar()
        todo, plan = [], []
        for d in ds:
            name = f"finra-short-interest-{d}"
            if _have(archive, name) and skip_stored and not refresh:
                rep[name] = {"skipped": "stored"}
                plan.append(_plan("finra_si", name, 0, finra_si_rows_est(d), stored=True))
            else:
                todo.append(d)
        todo = todo[:limit] if limit is not None else todo
        for i, d in enumerate(todo):
            n = finra_si_rows_est(d)
            plan.append(_plan("finra_si", f"finra-short-interest-{d}", 2 + math.ceil(n / FINRA_PAGE) + (2 if i == 0 and dates is None else 0), n))
        rep["_plan"] = plan              # (+2 on the first: the AAPL and SPY date-listing calls)
        rep["_dry_run"] = {"dates": len(ds), "already_stored": sum(1 for p in plan if p["already_stored"]), "to_fetch": len(todo)}
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        ds = sorted(dates) if dates is not None else finra_settlement_dates(c, th, retries, backoff)
        todo = []
        for d in ds:
            name = f"finra-short-interest-{d}"
            if _have(archive, name) and skip_stored and not refresh:
                rep[name] = {"skipped": "stored"}
            else:
                todo.append(d)
        todo = todo[:limit] if limit is not None else todo
        rep["_plan"] = [_plan("finra_si", f"finra-short-interest-{d}", 2 + math.ceil(finra_si_rows_est(d) / FINRA_PAGE),
                              finra_si_rows_est(d)) for d in todo]
        for d in todo:
            name = f"finra-short-interest-{d}"
            eq = {"compareType": "EQUAL", "fieldName": "settlementDate", "fieldValue": d}
            _, total, i1 = _finra_query(c, th, {"limit": 1, "offset": 0, "compareFilters": [eq]}, retries, backoff)
            _, otc, i2 = _finra_query(c, th, {"limit": 1, "offset": 0, "compareFilters": [eq],
                                              "domainFilters": [{"fieldName": "marketClassCode", "values": list(FINRA_OTC)}]}, retries, backoff)
            if total is None or otc is None:
                rep[name] = {"error": _fail(i1 if total is None else i2)}
                continue
            rows_raw, offset, ok, listed_total, mode = [], 0, True, None, "listed"
            body = {"limit": FINRA_PAGE, "compareFilters": [eq], "domainFilters": [{"fieldName": "marketClassCode", "values": list(FINRA_LISTED)}]}
            while True:
                data, tot, info = _finra_query(c, th, {**body, "offset": offset}, retries, backoff)
                if data is None:
                    ok = False
                    break
                listed_total = tot if tot is not None else listed_total
                rows_raw.extend(data)
                offset += FINRA_PAGE
                if len(data) < FINRA_PAGE or (listed_total is not None and offset >= listed_total):
                    break
            if ok and listed_total is not None and listed_total + otc != total and mode == "listed":
                mode, rows_raw, offset = "all (unknown market classes)", [], 0          # a class we do not list: page everything
                while True:
                    data, tot, info = _finra_query(c, th, {"limit": FINRA_PAGE, "offset": offset, "compareFilters": [eq]}, retries, backoff)
                    if data is None:
                        ok = False
                        break
                    rows_raw.extend(data)
                    offset += FINRA_PAGE
                    if len(data) < FINRA_PAGE or (tot is not None and offset >= tot):
                        break
            fetched = _stamp()
            if not ok:
                rep[name] = {"error": _fail(info)}
                continue
            rows = finra_si_rows(rows_raw, fetched)
            if not rows:
                rep[name] = {"error": f"no listed rows (total {total}, OTC {otc})"}
                continue
            classes = sorted({x["value"].get("marketClassCode") or "" for x in rows})
            rep[name] = {"rows": len(rows), "total": total, "otc": otc, "mode": mode, "classes": classes}
            if archive is None:
                rep[name]["rows"] = rows
                continue
            e = _put(archive, name, rows, _recipe(
                f"POST {FINRA_SI} JSON {{limit {FINRA_PAGE}, offset k, compareFilters [settlementDate EQUAL {d}], domainFilters "
                f"[marketClassCode IN {list(FINRA_LISTED)}]}} (paged; checked against the date's total {total} minus OTC/OTCBB {otc}; "
                f"mode {mode}) via mandel.data.public_history.fetch_finra_short_interest. FINRA does not give the dissemination "
                f"date: available_at = end of the {FINRA_PUBLISH_BDAYS}th NYSE business day after settlement (UTC), one day later "
                "than FINRA's usual schedule.", "finra", False), ["short-interest", "finra", "equities", d, "T-375", "T-393"],
                FINRA_SI_SCHEMA, in_git=False)
            rep[name]["stored"] = _stored_info(e)
            if log:
                log(f"finra si {d}: {len(rows)} rows")
    finally:
        if own:
            c.close()
    return rep


# ======================================================================== 10. FINRA Reg SHO CNMS daily (T-394) ==

CNMS_URL = "https://cdn.finra.org/equity/regsho/daily/CNMSshvol{ymd}.txt"
CNMS_FIRST = "2018-08-01"
CNMS_SETTLED_DAYS = 10                   # a trading day still without a file this long after is recorded as absent
CNMS_SCHEMA = ("kind short-volume: value = {day, symbol (FINRA spelling, e.g. BRK/B), short_volume, short_exempt_volume, total_volume, "
               "market (contributing facilities)}; available_at = end of the trade day (UTC) | kind file-day: value = {day, status, "
               "rows_in_file, rows_kept, last_modified, backfilled (Last-Modified more than 2 days after the trade day)}")


def norm_symbol(s: str) -> str:
    """One spelling for class shares across sources: BRK.B / BRK-B / BRK/B -> BRK/B."""
    s = (s or "").strip().upper()
    return re.sub(r"^([A-Z]{1,6})[.\-/]([A-Z])$", r"\1/\2", s)


def cnms_universe(store, since: str = CNMS_FIRST) -> set[str]:
    """S&P 500 members since ``since`` from the stored fja components and Wikipedia bundles (normalised symbols)."""
    archive = _archive(store)
    out: set[str] = set()
    rows = _load(archive, "fja-sp500-components")
    for r in rows:
        v = r.get("value") or {}
        if str(v.get("date")) >= since:
            out |= {norm_symbol(t) for t in v.get("tickers") or []}
    if rows:                                            # the members on the last date before ``since`` were members at ``since``
        before = [r for r in rows if str((r.get("value") or {}).get("date")) < since]
        if before:
            out |= {norm_symbol(t) for t in before[-1]["value"].get("tickers") or []}
    for r in _load(archive, "wikipedia-sp500-constituents"):
        out.add(norm_symbol((r.get("value") or {}).get("symbol") or ""))
    for r in _load(archive, "wikipedia-sp500-changes"):
        v = r.get("value") or {}
        if str(v.get("effective_date")) >= since:
            for k in ("added_ticker", "removed_ticker"):
                if v.get(k):
                    out.add(norm_symbol(v[k]))
    out.discard("")
    return out


def parse_cnms(text: str, day: str, universe: Optional[set], fetched: str) -> tuple[list[dict], int]:
    out, n = [], 0
    av = _not_after(_end_of_day(day), fetched)
    for line in text.splitlines()[1:]:
        p = line.split("|")
        if len(p) < 6 or not p[0].isdigit():
            continue
        n += 1
        sym = p[1].strip()
        if universe is not None and norm_symbol(sym) not in universe:
            continue
        out.append(_row("finra-regsho", "short-volume", f"{day}:{sym}", _et(day, "16:00"), av,
                        {"day": day, "symbol": sym, "short_volume": _num(p[2]), "short_exempt_volume": _num(p[3]),
                         "total_volume": _num(p[4]), "market": p[5].strip()}))
    return out, n


def _months(first: str, last: str) -> list[str]:
    y, m = int(first[:4]), int(first[5:7])
    out = []
    while f"{y:04d}-{m:02d}" <= last[:7]:
        out.append(f"{y:04d}-{m:02d}")
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def _month_end(mo: str) -> str:
    y, m = int(mo[:4]), int(mo[5:7])
    return (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)).isoformat()


def _trading_days(first: str, last: str) -> list[str]:
    d, e, out = date.fromisoformat(first), date.fromisoformat(last), []
    while d <= e:
        if is_trading_day(d):
            out.append(d.isoformat())
        d += timedelta(days=1)
    return out


def fetch_finra_cnms(store=None, first: str = CNMS_FIRST, last: Optional[str] = None, universe=None, dry_run: bool = False,
                     limit: Optional[int] = None, skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None,
                     min_interval: float = 0.6, retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``finra-regsho-cnms-<YYYY-MM>``: one CNMS file per NYSE trading day, rows kept for ``universe`` (default: the S&P 500
    point-in-time members since 2018-08 from the stored fja and Wikipedia bundles; ``'all'`` keeps every symbol)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    last = last or (date.fromisoformat(_today()) - timedelta(days=1)).isoformat()
    months = _months(first, last)
    todo, plan = [], []
    for mo in months:
        name = f"finra-regsho-cnms-{mo}"
        days = _trading_days(max(first, mo + "-01"), min(last, _month_end(mo)))
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            continue
        todo.append((mo, days))
    todo = todo[:limit] if limit is not None else todo
    per_day = 625 if universe != "all" else 11000          # kept rows per file: 620 of 7,676 on 2018-08-01 (S&P universe)
    plan = [_plan("finra_cnms", f"finra-regsho-cnms-{mo}", len(days), len(days) * (per_day + 1)) for mo, days in todo]
    rep["_plan"] = plan
    if dry_run or not todo:
        return rep
    if universe == "all":
        uni = None
    elif universe is not None:
        uni = {norm_symbol(s) for s in universe}
    else:
        uni = cnms_universe(store, first)
        if not uni:
            rep["_error"] = "no universe: store fja-sp500-components / wikipedia-sp500-* first, or pass universe= ('all' keeps every symbol)"
            return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    today = date.fromisoformat(_today())
    try:
        for mo, days in todo:
            name = f"finra-regsho-cnms-{mo}"
            rows, absent, failed = [], [], []
            for d in days:
                url = CNMS_URL.format(ymd=d.replace("-", ""))
                r, info = _get(c, url, th, retries=retries, backoff=backoff)
                fetched = _stamp()
                if r is None:
                    if info.get("status") in (403, 404):
                        absent.append(d)
                        status = "absent"
                    else:
                        failed.append(d)
                        continue
                    kept, n, lm = [], 0, None
                else:
                    kept, n = parse_cnms(r.text, d, uni, fetched)
                    lm, status = r.headers.get("last-modified"), "ok"
                try:
                    lm_day = datetime.strptime(lm, "%a, %d %b %Y %H:%M:%S GMT").date() if lm else None
                except ValueError:
                    lm_day = None
                rows.extend(kept)
                rows.append(_row("finra-regsho", "file-day", f"{d}:file", _et(d, "16:00"), _not_after(_end_of_day(d), fetched),
                                 {"day": d, "status": status, "rows_in_file": n, "rows_kept": len(kept), "last_modified": lm,
                                  "backfilled": bool(lm_day and (lm_day - date.fromisoformat(d)).days > 2)}))
            settled = all((today - date.fromisoformat(d)).days >= CNMS_SETTLED_DAYS for d in absent)
            month_days = _trading_days(mo + "-01", _month_end(mo))
            complete = not failed and settled and bool(days) and days[-1] == month_days[-1]
            rep[name] = {"days": len(days), "absent": absent, "failed": failed, "rows": sum(1 for x in rows if x["kind"] == "short-volume")}
            if not any(x["kind"] == "short-volume" for x in rows):
                rep[name]["error"] = "no rows (nothing stored)"
                continue
            rows.sort(key=lambda x: (x["value"]["day"], x["kind"] != "file-day", x["key"]))
            if archive is None:
                rep[name]["rows"] = rows
                continue
            e = _put(archive, name, rows, _recipe(
                f"GET {CNMS_URL} for every NYSE trading day of {mo} (403/404 = no file) via mandel.data.public_history.fetch_finra_cnms; "
                f"rows kept for {'every symbol' if uni is None else f'{len(uni)} symbols: the S&P 500 point-in-time members since {first} (fja + Wikipedia, class shares as X/Y)'}. "
                "available_at = end of the trade day (UTC; live files are written about 17:20 ET). Files before 2021-10 carry "
                "Last-Modified 2021-10-27 (re-posted): their values are sums of the per-facility files FINRA published each evening; "
                "file-day rows flag them (backfilled).", "finra", False),
                ["short-volume", "regsho", "finra", "daily", mo, "T-394"], CNMS_SCHEMA, in_git=False, complete=complete)
            rep[name]["stored"] = _stored_info(e)
            if log:
                log(f"cnms {mo}: {rep[name]['rows']} rows, absent {len(absent)}, failed {len(failed)}")
    finally:
        if own:
            c.close()
    return rep


# ===================================================================== 11. SEC EDGAR (T-375, T-393, T-451, T-473) ==

SEC_TICKERS = "https://www.sec.gov/files/company_tickers_exchange.json"
SEC_LOOKUP = "https://www.sec.gov/Archives/edgar/cik-lookup-data.txt"
SEC_SUBMISSIONS = "https://data.sec.gov/submissions/"
SEC_FACTS = "https://data.sec.gov/api/xbrl/companyfacts/"
SEC_BUCKET = 50_000
SEC_LISTED_EXCHANGES = ("Nasdaq", "NYSE", "CBOE", "TXSE")
SEC_ANNUAL_FORMS = frozenset({"10-K", "10-K/A", "10-KT", "10-KT/A", "10-K405", "10-K405/A", "20-F", "20-F/A", "40-F", "40-F/A"})
SEC_ALL_FORM_CONCEPTS = frozenset({("dei", "EntityCommonStockSharesOutstanding"), ("dei", "EntityPublicFloat"),
                                   ("us-gaap", "CommonStockSharesOutstanding"), ("us-gaap", "WeightedAverageNumberOfSharesOutstandingBasic")})
SEC_ANNUAL_CONCEPTS = {
    "us-gaap": ("NetIncomeLoss", "ProfitLoss", "NetIncomeLossAvailableToCommonStockholdersBasic", "Assets", "AssetsCurrent",
                "LiabilitiesCurrent", "Liabilities", "LongTermDebtNoncurrent", "LongTermDebt", "NetCashProvidedByUsedInOperatingActivities",
                "NetCashProvidedByUsedInOperatingActivitiesContinuingOperations", "Revenues", "RevenueFromContractWithCustomerExcludingAssessedTax",
                "SalesRevenueNet", "SalesRevenueGoodsNet", "GrossProfit", "CostOfRevenue", "CostOfGoodsAndServicesSold", "CostOfGoodsSold",
                "StockholdersEquity", "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest", "PreferredStockValue",
                "ProceedsFromIssuanceOfCommonStock", "WeightedAverageNumberOfDilutedSharesOutstanding"),
    "ifrs-full": ("ProfitLoss", "ProfitLossAttributableToOwnersOfParent", "Assets", "CurrentAssets", "CurrentLiabilities", "Liabilities",
                  "CashFlowsFromUsedInOperatingActivities", "Revenue", "GrossProfit", "CostOfSales", "Equity", "EquityAttributableToOwnersOfParent"),
}
SEC_SUB_PER_CIK = 1.3                     # calls per CIK (older-filings pages for long filing histories: AAPL 2); estimate only
SEC_SUB_ROWS_PER_CIK = 70                 # company + XBRL filing rows per CIK (ABT 160, AAPL 145, CELG 51 on 2026-10-15); estimate
SEC_FACT_ROWS_PER_CIK = 700               # facts kept per CIK (ABT 1,179, AAPL 1,347, CELG 938; small filers fewer); estimate
SEC_LISTED_CIKS = 10000                   # universe: 6,072 non-OTC current CIKs + ~4,000 name-matched delisted ones; estimate only
SEC_TICKERS_SCHEMA = ("kind ticker: value = {cik, name, ticker, exchange}; observed_at = available_at = the file's Last-Modified "
                      "(a current map: tickers of delisted companies are absent)")
SEC_LOOKUP_SCHEMA = "kind cik-name: value = {name (EDGAR conformed or former name), cik}; available_at = capture time (reference table)"
SEC_SUB_SCHEMA = ("kind company: value = {cik, name, entityType, sic, sicDescription, category, tickers, exchanges, fiscalYearEnd, "
                  "stateOfIncorporation, formerNames, status} (current profile; available_at = capture) | kind filing: value = {cik, accn, "
                  "form, filing_date, report_date, acceptance_json, is_xbrl, is_inline_xbrl, primary_document} for XBRL filings; "
                  "available_at = min(acceptance_json, filing date 22:00 ET) | kind coverage: value = {bucket, ciks, status per cik}")
SEC_FACTS_SCHEMA = ("kind fact: value = {cik, taxonomy, tag, unit, start, end, val, accn, fy, fp, form, filed, frame}: dei cover shares "
                    "and public float, us-gaap CommonStockSharesOutstanding and basic weighted shares from every form; the F-score and "
                    "book-equity concepts from annual forms (10-K, 20-F, 40-F and amendments) only; available_at = min(acceptance of the "
                    "accession from sec-submissions, filed date 22:00 ET) | kind coverage: value = {bucket, ciks, status per cik}")


def sec_bucket(cik: int) -> str:
    return f"b{int(cik) // SEC_BUCKET:03d}"


def _sec_filed_bound(day: str) -> str:
    return _et(day, "22:00")                # EDGAR accepts until 22:00 ET; later filings get the next business day's date


def _lm_iso(r: httpx.Response) -> Optional[str]:
    lm = r.headers.get("last-modified")
    try:
        return _iso(datetime.strptime(lm, "%a, %d %b %Y %H:%M:%S GMT").replace(tzinfo=UTC)) if lm else None
    except ValueError:
        return None


def fetch_sec_company_tickers(store=None, dry_run: bool = False, limit: Optional[int] = None, skip_stored: bool = True,
                              refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 0.5, retries: int = 4,
                              backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    archive = _archive(store)
    rep = report if report is not None else {}
    name = "sec-company-tickers"
    have = _have(archive, name) and skip_stored and not refresh
    rep["_plan"] = [_plan("sec_tickers", name, 1, 10450, stored=have)]
    if have:
        rep[name] = {"skipped": "stored"}
    if dry_run or have or (limit is not None and limit < 1):
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        r, info = _get(c, SEC_TICKERS, th, headers=SEC_UA, retries=retries, backoff=backoff)
    finally:
        if own:
            c.close()
    fetched = _stamp()
    if r is None:
        rep[name] = {"error": _fail(info)}
        return rep
    d = r.json()
    t = _not_after(_lm_iso(r) or fetched, fetched)
    rows = [_row("sec", "ticker", f"{x[0]}:{x[2]}", t, t, dict(zip(d["fields"], x))) for x in d.get("data") or []]
    rep[name] = {"rows": len(rows)}
    if archive is None:
        rep[name]["rows"] = rows
        return rep
    e = _put(archive, name, rows, _recipe(f"GET {SEC_TICKERS} (User-Agent '{SEC_UA['User-Agent']}') via "
                                          "mandel.data.public_history.fetch_sec_company_tickers; current filers only.", "sec", True),
             ["sec", "edgar", "cik", "tickers", "reference"], SEC_TICKERS_SCHEMA, in_git=True)
    rep[name]["stored"] = _stored_info(e)
    return rep


def fetch_sec_cik_lookup(store=None, dry_run: bool = False, limit: Optional[int] = None, skip_stored: bool = True, refresh: bool = False,
                         client: Optional[httpx.Client] = None, min_interval: float = 0.5, retries: int = 4, backoff: float = 2.0,
                         report: Optional[dict] = None, log=None) -> dict:
    """``sec-cik-lookup``: every EDGAR entity name (current and former) -> CIK (~1.06 M lines, 40 MB, one call)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    name = "sec-cik-lookup"
    have = _have(archive, name) and skip_stored and not refresh
    rep["_plan"] = [_plan("sec_lookup", name, 1, 1063400, stored=have)]
    if have:
        rep[name] = {"skipped": "stored"}
    if dry_run or have or (limit is not None and limit < 1):
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval, timeout=600)
    try:
        r, info = _get(c, SEC_LOOKUP, th, headers=SEC_UA, retries=retries, backoff=backoff)
    finally:
        if own:
            c.close()
    fetched = _stamp()
    if r is None:
        rep[name] = {"error": _fail(info)}
        return rep
    text = r.text

    def rows():
        for i, line in enumerate(text.splitlines()):
            m = re.match(r"^(.*):(\d{10}):\s*$", line)
            if m:
                yield _row("sec", "cik-name", f"{int(m.group(2))}:{i}", fetched, fetched, {"name": m.group(1).strip(), "cik": int(m.group(2))})
    rep[name] = {"last_modified": r.headers.get("last-modified")}
    if archive is None:
        rep[name]["rows"] = list(rows())
        return rep
    e = _put(archive, name, rows(), _recipe(f"GET {SEC_LOOKUP} (Last-Modified {r.headers.get('last-modified')}) via "
                                            "mandel.data.public_history.fetch_sec_cik_lookup; one row per name:cik line. Public domain; kept "
                                            "local for repository size (about 20 MB gzip).", "sec", True),
             ["sec", "edgar", "cik", "names", "reference"], SEC_LOOKUP_SCHEMA, in_git=False)
    rep[name]["stored"] = _stored_info(e)
    rep[name]["rows"] = e.rows
    return rep


_NAME_DROP = re.compile(r"\b(THE|INC|INCORPORATED|CORP|CORPORATION|CO|COMPANY|LTD|LIMITED|PLC|LLC|LP|L P|NV|N V|SA|AG|HOLDINGS?|GROUP|"
                        r"COMMON STOCK|CLASS [A-Z]|ORDINARY SHARES?|SHARES?|COM|NEW|DEL|DE|MD|NY|NJ|PA|TX|CA)\b")


def normalize_company_name(s: str) -> str:
    s = re.sub(r"/[A-Z]{2,3}/", " ", (s or "").upper())
    s = re.sub(r"&", " AND ", s)
    s = re.sub(r"[^A-Z0-9 ]+", " ", s)
    s = _NAME_DROP.sub(" ", s)
    return re.sub(r"\s+", " ", s).strip()


def match_names_to_ciks(names: Iterable[str], lookup_rows: Iterable[dict]) -> dict[str, list[int]]:
    """{name: [CIKs]} by exact match of normalised names against the sec-cik-lookup rows (empty list when unmatched)."""
    want = {n: normalize_company_name(n) for n in names if n}
    idx: dict[str, set[int]] = defaultdict(set)
    targets = set(want.values())
    for r in lookup_rows:
        v = r.get("value") or {}
        k = normalize_company_name(v.get("name") or "")
        if k in targets:
            idx[k].add(int(v["cik"]))
    return {n: sorted(idx.get(k, ())) for n, k in want.items()}


def _iter_rows(archive, name: str):
    """Stream a stored bundle's rows (nothing when it is not on this machine)."""
    if archive is None:
        return
    try:
        archive.get(name)
        yield from archive.load(name)
    except (KeyError, FileNotFoundError, OSError, ValueError):
        return


FINRA_COMPANY_CLASSES = frozenset({"NNM", "NYSE", "SC", "AMEX", "NGM", "NGS", "NCM", "NAS", "TXSE"})
_NOT_COMPANY = re.compile(r"(?i)\b(ETF|ETN|FUND|FD|WARRANTS?|WTS?|RIGHTS?|UNITS?|PREFERRED|PFD|NOTES?|DEBENTURES?|BONDS?|TRUST UNITS|"
                          r"INDEX|ISHARES|PROSHARES|DIREXION|INVESCO|SPDR|VANGUARD)\b")


def sec_universe(store, extra_ciks: Iterable[int] = (), match_delisted: bool = True) -> tuple[list[int], dict]:
    """CIKs to fetch: non-OTC current filers (sec-company-tickers), the Wikipedia constituents' CIKs, ``extra_ciks``, and
    with ``match_delisted`` the CIKs whose EDGAR name matches a removed S&P 500 member (wikipedia-sp500-changes) or a FINRA
    short-interest issue on a NYSE/Nasdaq/NYSE American class whose symbol is not a current ticker and whose name is not a
    fund, ETF/ETN, warrant, right, unit, note or preferred (sec-cik-lookup, normalised exact match; streamed)."""
    archive = _archive(store)
    ciks, prov = set(), {}
    tickers = set()
    for r in _iter_rows(archive, "sec-company-tickers"):
        v = r.get("value") or {}
        tickers.add(norm_symbol(v.get("ticker") or ""))
        if v.get("exchange") in SEC_LISTED_EXCHANGES:
            ciks.add(int(v["cik"]))
    prov["listed_current"] = len(ciks)
    for r in _load(archive, "wikipedia-sp500-constituents"):
        c = re.sub(r"\D", "", str((r.get("value") or {}).get("cik") or ""))
        if c:
            ciks.add(int(c))
    ciks |= {int(c) for c in extra_ciks}
    if match_delisted:
        names = set()
        for r in _load(archive, "wikipedia-sp500-changes"):
            v = r.get("value") or {}
            if v.get("removed_security"):
                names.add(v["removed_security"])
        idx = archive.entries(tag="short-interest") if archive is not None else []
        for e in idx:
            for r in _iter_rows(archive, e.name):
                v = r.get("value") or {}
                nm = v.get("issueName") or ""
                if (nm and v.get("marketClassCode") in FINRA_COMPANY_CLASSES and not _NOT_COMPANY.search(nm)
                        and norm_symbol(v.get("symbolCode") or "") not in tickers):
                    names.add(nm)
        if names and _stored_here(archive, "sec-cik-lookup"):
            got = match_names_to_ciks(names, _iter_rows(archive, "sec-cik-lookup"))
            matched = {c for v in got.values() for c in v}
            prov["delisted_names"], prov["delisted_matched_ciks"] = len(names), len(matched - ciks)
            ciks |= matched
    return sorted(ciks), prov


def sec_filing_rows(sub: dict, cik: int, fetched: str, extra_pages: Iterable[dict] = ()) -> list[dict]:
    out, seen = [], set()
    for block in [((sub.get("filings") or {}).get("recent") or {})] + list(extra_pages):
        accs = block.get("accessionNumber") or []
        for i, acc in enumerate(accs):
            if acc in seen:
                continue
            g = lambda k: (block.get(k) or [None] * len(accs))[i]
            if not g("isXBRL"):
                continue
            seen.add(acc)
            fd = str(g("filingDate") or "")[:10]
            if not fd:
                continue
            aj = g("acceptanceDateTime")
            bound = _sec_filed_bound(fd)
            av = min(_iso_ceil(_parse_iso(aj)), bound, key=_parse_iso) if aj else bound
            out.append(_row("sec", "filing", f"{cik}:{acc}", av, _not_after(av, fetched),
                            {"cik": cik, "accn": acc, "form": g("form"), "filing_date": fd, "report_date": g("reportDate") or None,
                             "acceptance_json": aj, "is_xbrl": g("isXBRL"), "is_inline_xbrl": g("isInlineXBRL"),
                             "primary_document": g("primaryDocument")}))
    return out


def sec_company_row(sub: dict, cik: int, fetched: str) -> dict:
    keys = ("name", "entityType", "sic", "sicDescription", "category", "tickers", "exchanges", "fiscalYearEnd", "stateOfIncorporation", "formerNames")
    return _row("sec", "company", str(cik), fetched, fetched, {"cik": cik, **{k: sub.get(k) for k in keys}})


def sec_fact_rows(cf: dict, cik: int, accept: dict[str, str], fetched: str) -> list[dict]:
    out = []
    for tax, concepts in (cf.get("facts") or {}).items():
        for tag, body in concepts.items():
            allform = (tax, tag) in SEC_ALL_FORM_CONCEPTS
            if not allform and tag not in SEC_ANNUAL_CONCEPTS.get(tax, ()):
                continue
            for unit, facts in (body.get("units") or {}).items():
                for f in facts:
                    form = f.get("form") or ""
                    if not allform and form not in SEC_ANNUAL_FORMS:
                        continue
                    filed = str(f.get("filed") or "")[:10]
                    if not filed:
                        continue
                    bound = _sec_filed_bound(filed)
                    acc = accept.get(f.get("accn"))
                    av = min(acc, bound, key=_parse_iso) if acc else bound
                    v = {"cik": cik, "taxonomy": tax, "tag": tag, "unit": unit, "start": f.get("start"), "end": f.get("end"),
                         "val": f.get("val"), "accn": f.get("accn"), "fy": f.get("fy"), "fp": f.get("fp"), "form": form,
                         "filed": filed, "frame": f.get("frame")}
                    out.append(_row("sec", "fact", f"{cik}:{tax}:{tag}:{unit}:{f.get('accn')}:{f.get('start')}:{f.get('end')}", av,
                                    _not_after(av, fetched), v))
    return out


def _bucket_rows(archive, name: str) -> tuple[list[dict], dict]:
    rows = _load(archive, name)
    cov = {}
    for r in rows:
        if r.get("kind") == "coverage":
            cov.update((r.get("value") or {}).get("status") or {})
    return [r for r in rows if r.get("kind") != "coverage"], cov


def _sec_buckets(ciks: Iterable[int]) -> dict[str, list[int]]:
    out: dict[str, list[int]] = defaultdict(list)
    for c in sorted(set(int(x) for x in ciks)):
        out[sec_bucket(c)].append(c)
    return dict(sorted(out.items()))


def _sec_bucket_fetch(kind: str, store, ciks, dry_run, limit, skip_stored, refresh, client, min_interval, retries, backoff, report, log):
    """Shared driver of fetch_sec_submissions / fetch_sec_companyfacts: per CIK bucket, fetch the CIKs not yet covered and
    store the bucket (old rows kept); a bucket is complete when every requested CIK answered (200 or 404)."""
    fam, prefix = ("sec_sub", "sec-submissions") if kind == "submissions" else ("sec_facts", "sec-companyfacts")
    archive = _archive(store)
    rep = report if report is not None else {}
    if ciks is None:
        ciks, prov = sec_universe(store)
        rep["_universe"] = prov
    buckets = _sec_buckets(ciks)
    todo, plan = [], []
    for b, cs in buckets.items():
        name = f"{prefix}-{b}"
        cov = _bucket_rows(archive, name)[1] if (archive is not None and _stored_here(archive, name)) else {}
        missing = [c for c in cs if cov.get(str(c)) not in ("ok", "404")] if (skip_stored and not refresh) else cs
        per = SEC_SUB_PER_CIK if kind == "submissions" else 1.0
        rows_est = len(cs) * (SEC_SUB_ROWS_PER_CIK if kind == "submissions" else SEC_FACT_ROWS_PER_CIK)
        if not missing:
            rep[name] = {"skipped": "stored"}
            plan.append(_plan(fam, name, 0, rows_est, stored=True))
            continue
        todo.append((b, cs, missing))
        plan.append(_plan(fam, name, len(missing) * per, rows_est, ciks=len(cs), to_fetch=len(missing)))
    if not buckets:
        per = SEC_SUB_PER_CIK if kind == "submissions" else 1.0
        plan.append(_plan(fam, f"{prefix}-b<NNN>", SEC_LISTED_CIKS * per, SEC_LISTED_CIKS * (SEC_SUB_ROWS_PER_CIK if kind == "submissions" else SEC_FACT_ROWS_PER_CIK),
                          bundles=43, note="CIK universe not stored yet (run sec_tickers, sec_lookup, wikipedia, finra_si first): estimate"))
    if limit is not None:
        todo = todo[:limit]
    rep["_plan"] = plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval)
    try:
        for b, cs, missing in todo:
            name = f"{prefix}-{b}"
            old, cov = _bucket_rows(archive, name) if archive is not None else ([], {})
            if refresh:
                old = [r for r in old if int((r.get("value") or {}).get("cik") or -1) not in set(missing)]
            accept = {}
            if kind == "facts":
                for r in _load(archive, f"sec-submissions-{b}"):
                    if r.get("kind") == "filing":
                        accept[r["value"]["accn"]] = r["available_at"]
            new, aborted = [], None
            for cik in missing:
                url = (f"{SEC_SUBMISSIONS}CIK{cik:010d}.json" if kind == "submissions" else f"{SEC_FACTS}CIK{cik:010d}.json")
                r, info = _get(c, url, th, headers=SEC_UA, retries=retries, backoff=backoff)
                fetched = _stamp()
                if r is None:
                    if info.get("status") == 404:
                        cov[str(cik)] = "404"
                        continue
                    if info.get("status") == 403 or info.get("retry_after"):
                        aborted = _fail(info)                    # SEC blocks abusive clients: stop the run, keep what we have
                        break
                    cov[str(cik)] = None
                    continue
                try:
                    d = r.json()
                except ValueError:
                    continue
                if kind == "submissions":
                    pages = []
                    for fpage in ((d.get("filings") or {}).get("files") or []):
                        if str(fpage.get("filingTo") or "") >= "2009-01-01":
                            r2, _ = _get(c, SEC_SUBMISSIONS + fpage["name"], th, headers=SEC_UA, retries=retries, backoff=backoff)
                            if r2 is not None:
                                try:
                                    pages.append(r2.json())
                                except ValueError:
                                    pass
                    new.append(sec_company_row(d, cik, fetched))
                    new.extend(sec_filing_rows(d, cik, fetched, pages))
                else:
                    new.extend(sec_fact_rows(d, cik, accept, fetched))
                cov[str(cik)] = "ok"
            done = [c_ for c_ in cs if cov.get(str(c_)) in ("ok", "404")]
            complete = len(done) == len(cs) and aborted is None
            rep[name] = {"ciks": len(cs), "fetched": sum(1 for c_ in missing if cov.get(str(c_)) == "ok"),
                         "not_found": sum(1 for c_ in missing if cov.get(str(c_)) == "404"), "new_rows": len(new)}
            if aborted:
                rep[name]["aborted"] = aborted
            fetched = _stamp()
            covrow = _row("mandel-sec", "coverage", f"{b}:coverage", fetched, fetched, {"bucket": b, "ciks": cs, "status": cov})
            rows = sorted(old + new, key=lambda r: r["key"]) + [covrow]
            if archive is None:
                rep[name]["rows"] = rows
            elif new or any(cov.get(str(c_)) for c_ in missing):
                e = _put(archive, name, rows, _recipe(
                    (f"GET {SEC_SUBMISSIONS}CIK##########.json (+ older 'files' pages reaching 2009 or later)" if kind == "submissions"
                     else f"GET {SEC_FACTS}CIK##########.json, keeping {sorted(SEC_ALL_FORM_CONCEPTS)} from every form and "
                          f"{ {k: len(v) for k, v in SEC_ANNUAL_CONCEPTS.items()} } F-score/book-equity concepts from annual forms")
                    + f" for the CIKs {cs[0]}..{cs[-1]} of bucket {b} (CIK // {SEC_BUCKET}), User-Agent '{SEC_UA['User-Agent']}', "
                    f"{min_interval}s pacing, via mandel.data.public_history.fetch_sec_{'submissions' if kind == 'submissions' else 'companyfacts'}. "
                    "available_at = min(acceptance time, filing date 22:00 ET): data.sec.gov's acceptanceDateTime was later than the "
                    "EDGAR header by the ET offset for AAPL 2016-2026 and equal to it for CELG 2019 (never earlier). Public domain; "
                    "kept local for repository size.", "sec", True), ["sec", "edgar", kind, b] + list(FAMILIES[fam][2]),
                    SEC_SUB_SCHEMA if kind == "submissions" else SEC_FACTS_SCHEMA, in_git=False, complete=complete)
                rep[name]["stored"] = _stored_info(e)
            if log:
                log(f"sec {kind} {b}: {rep[name]}")
            if aborted:
                rep["_aborted"] = aborted
                break
    finally:
        if own:
            c.close()
    return rep


def fetch_sec_submissions(store=None, ciks: Optional[Iterable[int]] = None, dry_run: bool = False, limit: Optional[int] = None,
                          skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None,
                          min_interval: float = 0.5, retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``sec-submissions-b<NNN>``: company profiles (SIC etc.) and XBRL filings with acceptance times, per CIK bucket.
    ``ciks`` defaults to ``sec_universe(store)`` (store sec-company-tickers, sec-cik-lookup and the Wikipedia/FINRA bundles
    first)."""
    return _sec_bucket_fetch("submissions", store, ciks, dry_run, limit, skip_stored, refresh, client, min_interval, retries,
                             backoff, report, log)


def fetch_sec_companyfacts(store=None, ciks: Optional[Iterable[int]] = None, dry_run: bool = False, limit: Optional[int] = None,
                           skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None,
                           min_interval: float = 0.5, retries: int = 4, backoff: float = 2.0, report: Optional[dict] = None, log=None) -> dict:
    """``sec-companyfacts-b<NNN>``: the shares, F-score and book-equity facts per CIK bucket (acceptance times from the
    same bucket's sec-submissions bundle when stored; otherwise the filed-date bound)."""
    return _sec_bucket_fetch("facts", store, ciks, dry_run, limit, skip_stored, refresh, client, min_interval, retries,
                             backoff, report, log)


# =========================================================================== 12. StockTwits backfill (T-414, T-415) ==

ST_API = "https://api.stocktwits.com/api/2/streams/symbol/{sym}.json"
ST_PREOPEN_SYMBOLS = ("SPY", "QQQ", "TSLA", "NVDA", "GME", "AMC", "AAPL", "PLTR", "MSTR", "COIN", "BTC.X", "ETH.X", "SOL.X", "DOGE.X")
ST_INTENSITY_SYMBOLS = ("SPY", "QQQ", "IWM", "AAPL", "AMZN", "GOOG", "GLD", "USO", "EEM", "TLT", "XLB", "XLE", "XLF", "XLI", "XLK",
                        "XLP", "XLU", "XLV", "XLY", "BTC.X", "ETH.X")
ST_SAMPLE_TIMES = ("10:00", "13:00", "15:30")
ST_WINDOW = ("16:00", "09:00")           # previous NYSE session 16:00 ET .. the day's 09:00 ET
ST_ANCHOR_SYMBOL = "SPY"
ST_RETENTION_DAYS = 360                  # max=632,500,000 answered 2025-10-15 messages on 2026-10-15: about a rolling year
ST_SEED = ((632_500_000, "2025-10-15T23:18:17+00:00"), (640_000_000, "2025-12-23T18:29:25+00:00"))
ST_MAX_PAGES = 400
# Pages per weekday pre-open window. Measured on Tuesday 2025-10-21 (133 pages for the 14 streams); the checker re-measured
# 2025-10-29 (229 pages, 448 s) and the four heaviest streams on 2026-02-11, 2026-06-16 and 2026-09-22 (SPY 39-56, mean 47;
# BTC.X 17-47, mean 28; NVDA 8-64; TSLA 6-16): about 1.3x the first measurement. SPY, BTC.X, NVDA and TSLA use those means
# (NVDA without its 64-page event day), the other streams 1.3x their 2025-10-21 count. Monday windows span the weekend
# (64 hours) and are estimated at ST_MONDAY_FACTOR times a weekday.
ST_PAGES_EST = {"SPY": 48, "QQQ": 17, "TSLA": 11, "NVDA": 13, "GME": 10, "AMC": 5, "AAPL": 8, "PLTR": 4, "MSTR": 4, "COIN": 2,
                "BTC.X": 28, "ETH.X": 9, "SOL.X": 4, "DOGE.X": 4}
ST_MONDAY_FACTOR = 2.5
ST_ANCHOR_CALLS = 9                      # secant-search calls per anchor (5-7 in the preview, 7-14 in the checker's runs)
# Seconds per call for the plan: 1.95 (FAMILIES['st_preopen']; measured 1.85 s on 2025-10-21, 1.92 s on 2025-10-29).
ST_MSGS_SCHEMA = ("kind message: value = {stream, id, created_at, sentiment (Bullish|Bearish|None), user_id, followers, symbols, likes, "
                  "body_len} (bodies not stored); available_at = created_at | kind window: value = {stream, day, window_start, "
                  "window_end, anchor_id, pages, messages, tagged, bullish, bearish, complete}; available_at = window end")
ST_INT_SCHEMA = ("kind intensity-sample: value = {stream, day, sample_time, anchor_id, n (messages created before the sample time "
                 "in the page), created_at (list, newest first), sentiment (list), oldest, newest, per_hour (n / hours from oldest to the "
                 "sample time)}; available_at = the sample time")
ST_ANCHOR_SCHEMA = ("kind anchor: value = {target, id (the first SPY message id at or after the target: every message created "
                    "before the target has a smaller id), calls, exact}; available_at = the target time (fetch metadata)")


def _st_msg(m: dict, stream: str) -> dict:
    u = m.get("user") or {}
    return {"stream": stream, "id": m.get("id"), "created_at": m.get("created_at"),
            "sentiment": (((m.get("entities") or {}).get("sentiment")) or {}).get("basic"), "user_id": u.get("id"),
            "followers": u.get("followers"), "symbols": [s.get("symbol") for s in m.get("symbols") or []],
            "likes": (m.get("likes") or {}).get("total"), "body_len": len(m.get("body") or "")}


def st_page(sym: str, max_id: Optional[int], c, th, retries: int = 4, backoff: float = 30.0) -> tuple[Optional[list], dict]:
    """Messages with id < max_id (newest first) and the cursor; None on failure. 429s wait 30 s, 60 s, ... (Retry-After wins)."""
    params = {"max": int(max_id)} if max_id else None
    r, info = _get(c, ST_API.format(sym=sym), th, params=params, retries=retries, backoff=backoff, max_wait=1800.0)
    if r is None:
        return None, info
    try:
        d = r.json()
    except ValueError:
        info["error"] = "not JSON"
        return None, info
    info["cursor"] = d.get("cursor") or {}
    return d.get("messages") or [], info


def find_anchor(target: str, c, th, known: list[tuple[int, str]], max_calls: int = 14, retries: int = 4,
                backoff: float = 30.0) -> tuple[Optional[int], dict]:
    """The id of the first SPY message at or after ``target`` (an upper bound for every message created before it), by
    secant search on (id, time) points; ``known`` collects the points seen. None when the target is before retention."""
    tt = _parse_iso(target)
    pts = sorted(set(known) | set(ST_SEED), key=lambda p: p[0])
    lo = max((p for p in pts if _parse_iso(p[1]) < tt), key=lambda p: p[0], default=None)
    hi = min((p for p in pts if _parse_iso(p[1]) >= tt), key=lambda p: p[0], default=None)
    calls = 0
    if hi is None:
        msgs, info = st_page(ST_ANCHOR_SYMBOL, None, c, th, retries, backoff)
        calls += 1
        if not msgs:
            return None, {"calls": calls, "error": _fail(info)}
        newest = (int(msgs[0]["id"]), msgs[0]["created_at"])
        known.append(newest)
        if _parse_iso(newest[1]) < tt:
            return None, {"calls": calls, "error": "target is in the future"}
        hi = newest
    if lo is None:                                      # nothing known before the target: search down from the newest point
        lo = (1, "2000-01-01T00:00:00+00:00")
    while calls < max_calls and hi[0] - lo[0] > 1:
        t0, t1 = _parse_iso(lo[1]).timestamp(), _parse_iso(hi[1]).timestamp()
        frac = (tt.timestamp() - t0) / (t1 - t0) if t1 > t0 else 0.5
        guess = int(lo[0] + min(max(frac, 0.02), 0.98) * (hi[0] - lo[0])) + 1
        msgs, info = st_page(ST_ANCHOR_SYMBOL, guess, c, th, retries, backoff)
        calls += 1
        if msgs is None:
            return None, {"calls": calls, "error": _fail(info)}
        if not msgs:                                    # nothing below the guess: before retention
            lo = (guess, lo[1])
            if guess >= hi[0] - 1:
                break
            continue
        newest, oldest = (int(msgs[0]["id"]), msgs[0]["created_at"]), (int(msgs[-1]["id"]), msgs[-1]["created_at"])
        known.extend([newest, oldest])
        if _parse_iso(newest[1]) < tt:
            lo = (guess - 1, newest[1])               # no SPY message in (newest, guess): the first one at/after target is >= guess
        elif _parse_iso(oldest[1]) >= tt:
            hi = oldest
            if not info.get("cursor", {}).get("more", True) and len(msgs) < 30:
                return None, {"calls": calls, "error": "target before the oldest retained message"}
        else:
            after = [int(m["id"]) for m in msgs if _parse_iso(m["created_at"]) >= tt]
            return min(after), {"calls": calls, "exact": True}
    return hi[0], {"calls": calls, "exact": False}


def st_window(sym: str, start: str, end_id: int, c, th, max_pages: int = ST_MAX_PAGES, retries: int = 4,
              backoff: float = 30.0, end: Optional[str] = None) -> tuple[list[dict], dict]:
    """Messages of ``sym`` created in [start, end) paging back from ``end_id``; info {pages, complete, error?}."""
    ts, te = _parse_iso(start), (_parse_iso(end) if end else None)
    out, info, mx, seen = [], {"pages": 0, "complete": False}, end_id, set()
    while info["pages"] < max_pages:
        msgs, i = st_page(sym, mx, c, th, retries, backoff)
        info["pages"] += 1
        if msgs is None:
            info["error"] = _fail(i)
            return out, info
        for m in msgs:
            t = _parse_iso(m["created_at"])
            if t >= ts and (te is None or t < te) and m["id"] not in seen:
                seen.add(m["id"])
                out.append(m)
        if not msgs or _parse_iso(msgs[-1]["created_at"]) < ts or not (i.get("cursor") or {}).get("more", True):
            info["complete"] = True
            return out, info
        mx = int((i.get("cursor") or {}).get("max") or msgs[-1]["id"])
    info["stopped"] = "max_pages"
    return out, info


def _week_over(ds: list[str]) -> bool:
    """The last of ``ds`` is the last NYSE trading day of its ISO week (the week is fully in the past)."""
    d0 = date.fromisoformat(ds[0])
    return ds[-1] == _trading_days(ds[0], (d0 + timedelta(days=6 - d0.weekday())).isoformat())[-1]


def _iso_week(day: str) -> str:
    y, w, _ = date.fromisoformat(day).isocalendar()
    return f"{y}-W{w:02d}"


def _st_days(first: Optional[str], last: Optional[str], end_time: str) -> list[str]:
    now = _now()
    lo = max(first or "0000", (now.date() - timedelta(days=ST_RETENTION_DAYS)).isoformat())
    hi = last or now.date().isoformat()
    return [d for d in _trading_days(lo, hi) if _parse_iso(_et(d, end_time)) <= now]


def _anchors(archive, c, th, targets: list[str], known: list, cache: dict, retries: int, backoff: float, rep: dict) -> None:
    for t in targets:
        if t in cache:
            continue
        aid, info = find_anchor(t, c, th, known, retries=retries, backoff=backoff)
        cache[t] = {"target": t, "id": aid, "calls": info.get("calls"), "exact": info.get("exact", False), "error": info.get("error")}
        rep.setdefault("_anchor_calls", 0)
        rep["_anchor_calls"] += info.get("calls") or 0


def _store_anchors(archive, cache: dict) -> None:
    """Store the anchor cache merged with what is on disk (anchors of earlier runs; one writer per archive, see run_lock)."""
    if archive is None or not cache:
        return
    for r in _load(archive, "stocktwits-anchors"):
        v = r.get("value") or {}
        if v.get("id") and v.get("target") not in cache:
            cache[v["target"]] = v
    rows = [_row("mandel-stocktwits", "anchor", a["target"], a["target"], a["target"], a) for a in sorted(cache.values(), key=lambda a: a["target"])
            if a.get("id")]
    if rows:
        _put(archive, "stocktwits-anchors", rows, _recipe(
            f"find_anchor: secant search on {ST_API.format(sym=ST_ANCHOR_SYMBOL)}?max=<id> (message ids are global and rise with time) via "
            "mandel.data.public_history; fetch metadata merged across runs.", "stocktwits", False),
            ["stocktwits", "social", "anchors", "T-414", "T-415"], ST_ANCHOR_SCHEMA, in_git=False)


def _load_anchor_cache(archive) -> tuple[dict, list]:
    cache, known = {}, []
    for r in _load(archive, "stocktwits-anchors"):
        v = r.get("value") or {}
        if v.get("id"):
            cache[v["target"]] = v
            known.append((int(v["id"]), v["target"]))
    return cache, known


def _st_week_store(archive, name: str, rows: list[dict], days_done: set, recipe: str, tags: list, schema: str, complete: bool):
    old = [r for r in _load(archive, name) if (r.get("value") or {}).get("day") not in days_done]
    allrows = sorted(old + rows, key=lambda r: (r["value"].get("day") or "", r["key"]))
    return _put(archive, name, allrows, recipe, tags, schema, in_git=False, complete=complete)


def _st_week_days_done(archive, name: str, kind: str) -> dict[str, bool]:
    out: dict[str, dict] = defaultdict(dict)
    for r in _load(archive, name):
        if r.get("kind") == kind:
            v = r.get("value") or {}
            out[v.get("day")][v.get("stream")] = bool(v.get("complete", True))
    return out


def fetch_stocktwits_preopen(store=None, symbols: Iterable[str] = ST_PREOPEN_SYMBOLS, first: Optional[str] = None,
                             last: Optional[str] = None, dry_run: bool = False, limit: Optional[int] = None, skip_stored: bool = True,
                             refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 1.0, retries: int = 4,
                             backoff: float = 30.0, max_pages: int = ST_MAX_PAGES, max_days: Optional[int] = None,
                             report: Optional[dict] = None, log=None) -> dict:
    """``stocktwits-preopen-<YYYY>-W<ww>``: every message of ``symbols`` from 16:00 ET on the previous NYSE session to 09:00
    ET of each trading day, oldest week first (retention erodes the start daily). A week is stored after every day it
    gains; complete when all its days and symbols are. ``limit`` caps weeks, ``max_days`` days (a preview)."""
    syms = list(symbols)
    archive = _archive(store)
    rep = report if report is not None else {}
    days = _st_days(first, last, ST_WINDOW[1])
    weeks: dict[str, list[str]] = defaultdict(list)
    for d in days:
        weeks[_iso_week(d)].append(d)
    todo, plan = [], []
    for w, ds in sorted(weeks.items()):
        name = f"stocktwits-preopen-{w}"
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            continue
        done = _st_week_days_done(archive, name, "window") if (skip_stored and not refresh) else {}
        need = [d for d in ds if not (d in done and all(done[d].get(s) for s in syms))]
        todo.append((w, ds, need))
        pages = sum(ST_PAGES_EST.get(s, 5) * (ST_MONDAY_FACTOR if date.fromisoformat(d).weekday() == 0 else 1) for d in need for s in syms)
        plan.append(_plan("st_preopen", name, pages + ST_ANCHOR_CALLS * len(need), pages * 27 + len(need) * len(syms), days=len(need)))
    todo = todo[:limit] if limit is not None else todo
    rep["_plan"] = plan[:len(todo)] if limit is not None else plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval, timeout=60)
    cache, known = _load_anchor_cache(archive)
    n_days = 0
    try:
        for w, ds, need in todo:
            name = f"stocktwits-preopen-{w}"
            stop = False
            for d in need:
                if max_days is not None and n_days >= max_days:
                    stop = True
                    break
                end_t, start_t = _et(d, ST_WINDOW[1]), _et(prev_session(d), ST_WINDOW[0])
                _anchors(archive, c, th, [end_t], known, cache, retries, backoff, rep)
                aid = cache[end_t].get("id")
                if not aid:
                    rep.setdefault("_no_anchor", []).append(d)
                    continue
                fetched = _stamp()
                rows = []
                for s in syms:
                    msgs, info = st_window(s, start_t, aid, c, th, max_pages=max_pages, retries=retries, backoff=backoff, end=end_t)
                    for m in msgs:
                        v = _st_msg(m, s)
                        v["day"] = d
                        t_ = _iso(_parse_iso(m["created_at"]))
                        rows.append(_row("stocktwits", "message", f"{s}:{m['id']}", t_, _not_after(t_, fetched), v))
                    tags = [(m.get("entities") or {}).get("sentiment") or {} for m in msgs]
                    bull = sum(1 for t in tags if t.get("basic") == "Bullish")
                    bear = sum(1 for t in tags if t.get("basic") == "Bearish")
                    rows.append(_row("mandel-stocktwits", "window", f"{s}:{d}:window", end_t, _not_after(end_t, fetched),
                                     {"stream": s, "day": d, "window_start": start_t, "window_end": end_t, "anchor_id": aid,
                                      "pages": info["pages"], "messages": len(msgs), "tagged": bull + bear, "bullish": bull,
                                      "bearish": bear, "complete": info["complete"], **({"error": info["error"]} if info.get("error") else {})}))
                    if info.get("error", {}).get("retry_after"):
                        stop = True
                        rep["_aborted"] = info["error"]
                        break
                n_days += 1
                if archive is None:
                    rep.setdefault(name, {"rows": []})["rows"].extend(rows)
                    continue
                prev = _st_week_days_done(archive, name, "window")
                day_ok = all(r["value"]["complete"] for r in rows if r["kind"] == "window") and len([r for r in rows if r["kind"] == "window"]) == len(syms)
                all_days = all((dd == d and day_ok) or (dd in prev and all(prev[dd].get(s) for s in syms)) for dd in ds)
                week_over = _week_over(ds)
                e = _st_week_store(archive, name, rows, {d}, _recipe(
                    f"GET {ST_API}?max=<id> paged backwards (30 messages per page) for {syms}, from find_anchor(09:00 ET) back to 16:00 ET "
                    f"of the previous NYSE session, via mandel.data.public_history.fetch_stocktwits_preopen; {min_interval}s pacing, 429 "
                    "backoff from 30 s. Messages are available at created_at; deleted messages are absent; bodies not stored.",
                    "stocktwits", False), ["stocktwits", "social", "sentiment", "preopen", w, "T-414"], ST_MSGS_SCHEMA,
                    complete=all_days and week_over)
                rep[name] = {"stored": _stored_info(e), "days": sorted(set(prev) | {d})}
                _store_anchors(archive, cache)
                if log:
                    log(f"stocktwits preopen {d}: {sum(1 for r in rows if r['kind'] == 'message')} messages")
                if stop:
                    break
            if stop:
                break
    finally:
        if own:
            c.close()
    return rep


def fetch_stocktwits_intensity(store=None, symbols: Iterable[str] = ST_INTENSITY_SYMBOLS, times: Iterable[str] = ST_SAMPLE_TIMES,
                               first: Optional[str] = None, last: Optional[str] = None, dry_run: bool = False, limit: Optional[int] = None,
                               skip_stored: bool = True, refresh: bool = False, client: Optional[httpx.Client] = None,
                               min_interval: float = 1.0, retries: int = 4, backoff: float = 30.0, max_days: Optional[int] = None,
                               report: Optional[dict] = None, log=None) -> dict:
    """``stocktwits-intensity-<YYYY>-W<ww>``: for each trading day, sample time and symbol, the page of 30 messages just
    before the sample time (one call each; anchors shared with the pre-open fetch)."""
    syms, tms = list(symbols), list(times)
    archive = _archive(store)
    rep = report if report is not None else {}
    days = _st_days(first, last, tms[-1])
    weeks: dict[str, list[str]] = defaultdict(list)
    for d in days:
        weeks[_iso_week(d)].append(d)
    todo, plan = [], []
    for w, ds in sorted(weeks.items()):
        name = f"stocktwits-intensity-{w}"
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            continue
        have = set()
        if skip_stored and not refresh:
            have = {r["value"]["day"] for r in _load(archive, name) if r.get("kind") == "intensity-sample"}
        need = [d for d in ds if d not in have]
        todo.append((w, ds, need))
        plan.append(_plan("st_intensity", name, len(need) * len(tms) * (len(syms) + ST_ANCHOR_CALLS), len(need) * len(tms) * len(syms),
                          days=len(need)))
    todo = todo[:limit] if limit is not None else todo
    rep["_plan"] = plan[:len(todo)] if limit is not None else plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval, timeout=60)
    cache, known = _load_anchor_cache(archive)
    n_days = 0
    try:
        for w, ds, need in todo:
            name = f"stocktwits-intensity-{w}"
            stop = False
            for d in need:
                if max_days is not None and n_days >= max_days:
                    stop = True
                    break
                targets = [_et(d, t) for t in tms]
                _anchors(archive, c, th, targets, known, cache, retries, backoff, rep)
                fetched = _stamp()
                rows, ok = [], True
                for t in targets:
                    aid = cache[t].get("id")
                    if not aid:
                        ok = False
                        continue
                    for s in syms:
                        msgs, info = st_page(s, aid, c, th, retries, backoff)
                        if msgs is None:
                            ok = False
                            if (info.get("retry_after")):
                                stop = True
                                rep["_aborted"] = _fail(info)
                                break
                            continue
                        before = [m for m in msgs if _parse_iso(m["created_at"]) < _parse_iso(t)]
                        oldest = before[-1]["created_at"] if before else None
                        hours = ((_parse_iso(t) - _parse_iso(oldest)).total_seconds() / 3600) if oldest else None
                        rows.append(_row("stocktwits", "intensity-sample", f"{s}:{t}", t, _not_after(t, fetched),
                                         {"stream": s, "day": d, "sample_time": t, "anchor_id": aid, "n": len(before),
                                          "created_at": [m["created_at"] for m in before],
                                          "sentiment": [(((m.get("entities") or {}).get("sentiment")) or {}).get("basic") for m in before],
                                          "oldest": oldest, "newest": before[0]["created_at"] if before else None,
                                          "per_hour": round(len(before) / hours, 4) if hours else None}))
                    if stop:
                        break
                n_days += 1
                if archive is None:
                    rep.setdefault(name, {"rows": []})["rows"].extend(rows)
                elif rows and ok:
                    prev = {r["value"]["day"] for r in _load(archive, name) if r.get("kind") == "intensity-sample"}
                    week_over = _week_over(ds)
                    e = _st_week_store(archive, name, rows, {d}, _recipe(
                        f"GET {ST_API}?max=<find_anchor(sample time)> once per symbol and sample time {tms} ET for {syms} via "
                        "mandel.data.public_history.fetch_stocktwits_intensity; only messages created before the sample time are kept.",
                        "stocktwits", False), ["stocktwits", "social", "intensity", w, "T-415"], ST_INT_SCHEMA,
                        complete=week_over and all(dd in prev | {d} for dd in ds))
                    rep[name] = {"stored": _stored_info(e)}
                    _store_anchors(archive, cache)
                if log:
                    log(f"stocktwits intensity {d}: {len(rows)} samples")
                if stop:
                    break
            if stop:
                break
    finally:
        if own:
            c.close()
    return rep


# ======================================================================= 13. Deribit near-expiry option trades (T-337) ==

NEAR_PREFIX = "deribit-option-trades-near-"
NEAR_FIRST = "2024-03-01"
NEAR_DTE = (0, 8)
NEAR_HOURS = (0, 24)
NEAR_KEEP_SHARE = 0.67                   # 12,264 of 18,228 trades on 2024-03-01 were 0-8 days out (preview); estimate only
NEAR_DAY_FACTOR = 7.7                    # whole-day trades / deribit_history's 07-09 UTC window mean (18,228 / 2,360 for 2024)
NEAR_SCHEMA = ("kind option-trade: value = {trade_id, timestamp (ms), instrument_name, price (coin), mark_price, iv, index_price, amount, "
               "direction} + {currency, expiry, expiry_code, strike, type} (deribit_history's daily row); trades of the whole UTC day in "
               "expiries 0-8 calendar days after it; available_at = trade time | kind sample-day: value = {day, currency, window_start, "
               "window_end, hours, dte_days, trades_seen, trades_kept, pages, complete}, available_at = window end")


def near_bundle(month: str, currency: str = "BTC") -> str:
    return f"{NEAR_PREFIX}{currency}-{month}"


def fetch_deribit_near_trades(store=None, currency: str = "BTC", first: str = NEAR_FIRST, last: Optional[str] = None,
                              dte_days: tuple = NEAR_DTE, dry_run: bool = False, limit: Optional[int] = None, skip_stored: bool = True,
                              refresh: bool = False, client: Optional[httpx.Client] = None, min_interval: float = 0.3, retries: int = 4,
                              backoff: float = 1.0, max_pages: Optional[int] = None, max_days: Optional[int] = None,
                              tmp_dir: Optional[str] = None, report: Optional[dict] = None, log=None) -> dict:
    """``deribit-option-trades-near-<cur>-<YYYY-MM>``: every option trade of each UTC day in expiries ``dte_days`` days out,
    month bundles, complete days merged into the month after every ``chunk`` (streamed, deribit_history._merge_store)."""
    archive = _archive(store)
    rep = report if report is not None else {}
    now_ms = int(_now().timestamp() * 1000)
    last = last or (_now().date() - timedelta(days=1)).isoformat()
    todo, plan = [], []
    for mo in _months(first, last):
        name = near_bundle(mo, currency)
        mlast = _month_end(mo)
        days = []
        d = date.fromisoformat(max(first, mo + "-01"))
        while d.isoformat() <= min(last, mlast):
            if daily_window(d.isoformat(), NEAR_HOURS)[1] + HISTORY_LAG_MS <= now_ms:
                days.append(d.isoformat())
            d += timedelta(days=1)
        if _have(archive, name) and skip_stored and not refresh:
            rep[name] = {"skipped": "stored"}
            continue
        stored = _scan(archive, name, ("sample-day",), _daily_group)["sample-day"] if archive is not None else {}
        need = [x for x in days if not (skip_stored and not refresh and stored.get(x))]
        if not need:
            continue
        todo.append((mo, days, need, mlast))
        per = [DAILY_WINDOW_TRADES[min(max(int(x[:4]), min(DAILY_WINDOW_TRADES)), max(DAILY_WINDOW_TRADES))] * NEAR_DAY_FACTOR for x in need]
        plan.append(_plan("deribit_near", name, sum(math.ceil(n / 1000) for n in per), sum(per) * NEAR_KEEP_SHARE + len(need), days=len(need)))
    todo = todo[:limit] if limit is not None else todo
    rep["_plan"] = plan[:len(todo)] if limit is not None else plan
    if dry_run or not todo:
        return rep
    archive, rep, c, own, th = _setup(store, rep, client, min_interval, timeout=300)
    n_days = 0
    try:
        for mo, days, need, mlast in todo:
            name = near_bundle(mo, currency)
            spill = _Spill(tmp_dir) if archive is not None else None
            done, kept_rows, aborted = set(), [], None
            try:
                for day in need:
                    if max_days is not None and n_days >= max_days:
                        break
                    s, e = daily_window(day, NEAR_HOURS)
                    kept, info = trades_in_window(currency, s, e, c, th, keep=_dte_keep(day, dte_days), max_pages=max_pages,
                                                  retries=retries, backoff=backoff)
                    observed = _stamp()
                    drows = sorted([_daily_trade_row(t, currency, observed) for t in kept]
                                   + [_sample_row(day, currency, s, e, NEAR_HOURS, dte_days, info, len(kept), observed)], key=_daily_sort)
                    n_days += 1
                    if not info["complete"]:
                        rep.setdefault("_incomplete", {})[day] = {k: info.get(k) for k in ("pages", "seen", "stopped", "error") if info.get(k)}
                        err = info.get("error") or {}
                        if err.get("retry_after"):
                            aborted = err
                            break
                        continue
                    done.add(day)
                    if spill is not None:
                        spill.write(drows)
                    else:
                        kept_rows.extend(drows)
                    if log:
                        log(f"deribit near {day}: {len(kept)} of {info.get('seen')} trades")
                if archive is None:
                    rep[name] = {"rows": kept_rows, "days": sorted(done)}
                elif done:
                    prev = _scan(archive, name, ("sample-day",), _daily_group)["sample-day"]
                    have_days = {x for x, ok in prev.items() if ok} | done
                    complete = all(x in have_days for x in days) and days and days[-1] == mlast
                    ent = _merge_store(archive, name, spill.rows(), {date.fromisoformat(x).toordinal() - date(1970, 1, 1).toordinal() for x in done},
                                       _daily_idx, _daily_sort, tags=sorted({"options", "crypto", "deribit", "trades", "near-expiry", mo, "T-337",
                                                                             "complete" if complete else "partial"}),
                                       schema=NEAR_SCHEMA, tmp_dir=tmp_dir, recipe=_recipe(
                                           f"GET https://history.deribit.com/api/v2/public/get_last_trades_by_currency_and_time?currency={currency}"
                                           "&kind=option&start_timestamp=<day 00:00 UTC ms>&end_timestamp=<next day 00:00 ms - 1>&count=1000&"
                                           "include_old=true&sorting=asc, paged forward (mandel.data.options_misc.trades_in_window), keeping "
                                           f"expiries {dte_days[0]}-{dte_days[1]} calendar days after the day (the expiries bracketing any "
                                           "Polymarket 'Bitcoin above' resolution within a week), rows built by mandel.data.deribit_history; "
                                           "via mandel.data.public_history.fetch_deribit_near_trades. available_at = trade time.", "deribit", False))
                    rep[name] = {"stored": _stored_info(ent), "days": len(have_days), "complete": bool(complete)}
            finally:
                if spill is not None:
                    spill.close()
            if aborted:
                rep["_aborted"] = aborted
                break
            if max_days is not None and n_days >= max_days:
                break
    finally:
        if own:
            c.close()
    return rep


# =================================================================================================== everything ==

SOURCES = {
    "wikipedia": fetch_wikipedia_sp500, "fja": fetch_fja_sp500, "treasury": fetch_treasury_auctions, "cftc": fetch_cftc_cot,
    "clevelandfed": fetch_clevelandfed_nowcast, "clevelandfed_snapshot": fetch_clevelandfed_snapshot, "spdr": fetch_spdr_archive,
    "sprott": fetch_sprott_trusts, "ssga": fetch_ssga_holdings, "finra_si": fetch_finra_short_interest, "finra_cnms": fetch_finra_cnms,
    "sec_tickers": fetch_sec_company_tickers, "sec_lookup": fetch_sec_cik_lookup, "sec_submissions": fetch_sec_submissions,
    "sec_companyfacts": fetch_sec_companyfacts, "stocktwits_preopen": fetch_stocktwits_preopen,
    "stocktwits_intensity": fetch_stocktwits_intensity, "deribit_near": fetch_deribit_near_trades,
}
# StockTwits first: its rolling one-year retention loses the oldest day every day, and nothing else here erodes. The
# intensity samples (~7 h) go before the pre-open windows (~30 h): the pre-open start then slips 7 hours instead of the
# intensity start slipping 30. Then dependencies first (the membership lists before CNMS; FINRA and the SEC maps before
# the SEC buckets).
DEFAULT_ORDER = ("stocktwits_intensity", "stocktwits_preopen", "wikipedia", "fja", "treasury", "cftc", "clevelandfed",
                 "clevelandfed_snapshot", "spdr", "sprott", "ssga", "sec_tickers", "sec_lookup", "finra_si", "finra_cnms",
                 "sec_submissions", "sec_companyfacts", "deribit_near")
# ONE WRITER AT A TIME. ``Archive`` keeps an in-memory index and re-reads INDEX.json with ``setdefault`` on every store,
# so a second process storing into the same archive writes back its own stale entry for every bundle it has seen
# (reproduced by the checker: a StockTwits week indexed at 1 row while its file held 2; ``verify`` fails). Storing runs
# therefore hold an exclusive, non-blocking lock in the archive root: ``public_history.run.lock`` and also
# ``crypto_history.run.lock`` (the sibling module's lock, so a crypto_history run and a public_history run refuse each
# other); both are gitignored by ``mandel/**/*.lock``. To run StockTwits in parallel with the rest, give it its own
# archive root (``--root DIR``) and ``merge`` it into the real archive afterwards.
RUN_LOCKS = ("public_history.run.lock", "crypto_history.run.lock")


@contextlib.contextmanager
def run_lock(archive, names: Iterable[str] = RUN_LOCKS):
    """Exclusive, non-blocking locks ``<archive root>/<name>`` for one storing run: a second writer on the same archive
    raises RuntimeError at once instead of interleaving its INDEX writes (see ONE WRITER AT A TIME). No-op without an
    archive root (store=None)."""
    root = getattr(archive, "root", None)
    if root is None:
        yield None
        return
    held = []
    try:
        for n in names:
            path = Path(root) / n
            fh = open(path, "a+", encoding="utf-8")
            try:
                fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            except OSError:
                fh.close()
                raise RuntimeError(f"another writer holds {path}: run one storing process per archive (two writers leave stale "
                                   "INDEX.json entries); run StockTwits under its own --root and merge afterwards") from None
            held.append(fh)
        yield [Path(root) / n for n in names]
    finally:
        for fh in reversed(held):
            fcntl.flock(fh.fileno(), fcntl.LOCK_UN)
            fh.close()


def index_check(archive, before: dict[str, str]) -> dict:
    """Compare this process's index entry of every bundle of this module it stored (sha256 changed since ``before``) with
    a fresh read of INDEX.json. A mismatch means another process wrote INDEX.json from a stale view during the run."""
    stored = {e.name: e for e in archive.entries() if before.get(e.name) != e.sha256 and _family(e.name)}
    if not stored:
        return {"ok": True, "bundles": 0}
    disk = type(archive)(root=archive.root, ledger=archive.ledger)
    bad = []
    for n, e in sorted(stored.items()):
        try:
            d = disk.get(n)
        except KeyError:
            bad.append({"bundle": n, "index": None, "stored_rows": e.rows})
            continue
        if (d.sha256, d.rows, d.bytes) != (e.sha256, e.rows, e.bytes):
            bad.append({"bundle": n, "index_rows": d.rows, "index_created_at": d.created_at, "stored_rows": e.rows,
                        "stored_created_at": e.created_at})
    out = {"ok": not bad, "bundles": len(stored)}
    if bad:
        out["mismatched"] = bad
        out["hint"] = ("another process wrote INDEX.json from a stale view during this run; when nothing else is writing, "
                       "re-run these bundles' sources with --refresh and check Archive().verify(name)")
    return out


@contextlib.contextmanager
def _single_writer(store, rep: dict, dry_run: bool, log=None):
    """run_lock + index_check around a storing run (nothing for a dry run or store=None)."""
    archive = _archive(store)
    if dry_run or archive is None or getattr(archive, "root", None) is None:
        yield
        return
    with run_lock(archive):
        before = {e.name: e.sha256 for e in archive.entries()}
        try:
            yield
        finally:
            rep["_index_check"] = index_check(archive, before)
            if log and not rep["_index_check"]["ok"]:
                log(f"INDEX CHECK FAILED: {json.dumps(rep['_index_check'], default=str)[:2000]}")


def fetch_all_misc(store, dry_run: bool = False, limit: Optional[int] = None, sources: Optional[Iterable[str]] = None,
                   refresh: bool = False, report: Optional[dict] = None, log=None) -> dict:
    """Every bundle of the misc source group in ONE storing process, resumable (bundles stored complete on this machine are
    skipped). Sources run in ``DEFAULT_ORDER`` (StockTwits first, then dependencies first). A storing run holds the archive's
    run locks (a second writer raises RuntimeError) and reports '_index_check'. Returns {source: report, '_plan': [...],
    '_totals': {...}, '_index_check': {...}}."""
    rep = report if report is not None else {}
    chosen = [s for s in DEFAULT_ORDER if sources is None or s in set(sources)]
    with _single_writer(store, rep, dry_run, log):
        for s in chosen:
            try:
                rep[s] = SOURCES[s](store, dry_run=dry_run, limit=limit, refresh=refresh, log=log)
            except Exception as e:                             # one source failing never stops the others
                rep[s] = {"_error": repr(e)[:300]}
                if log:
                    log(f"{s}: {e!r}")
    plan = [p for s in chosen for p in (rep.get(s) or {}).get("_plan", [])]
    rep["_plan"] = plan
    rep["_totals"] = {"bundles": len(plan), "calls": sum(p["calls"] for p in plan), "rows": sum(p["rows"] for p in plan),
                      "gz_mb": round(sum(p["gz_mb"] for p in plan), 1), "minutes": round(sum(p["minutes"] for p in plan), 1),
                      "local_only_families": len(gitignore_lines())}
    return rep


def _root_archive(root):
    """The scratch archive of ``--root DIR`` / ``--preview DIR``: DIR/archive with its own ledger DIR/ledger.jsonl."""
    from mandel.archive import Archive
    from mandel.ledger import Ledger
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    return Archive(root=root / "archive", ledger=Ledger(root / "ledger.jsonl"))


def merge_archive(src_root, store, pattern: str = "*", report: Optional[dict] = None, log=None) -> dict:
    """Copy this module's bundles matching ``pattern`` from the archive of ``src_root`` (the DIR of a ``--root DIR`` run:
    DIR/archive, DIR/ledger.jsonl) into
    ``store``, under the run locks, with an index check. A bundle missing from ``store`` or with different content is
    stored (recipe, tags, schema and in_git kept) unless the target already holds more rows (reported as a conflict, not
    overwritten); identical bundles are skipped. ``stocktwits-anchors`` is merged as the union of both by target time."""
    if not (Path(src_root) / "archive" / "INDEX.json").exists():
        raise FileNotFoundError(f"no archive under {src_root} (expected {Path(src_root) / 'archive' / 'INDEX.json'})")
    src = _root_archive(src_root)
    dst = _archive(store)
    rep = report if report is not None else {}
    with _single_writer(dst, rep, False, log):
        for e in sorted(src.entries(kind="data"), key=lambda x: x.name):
            if not fnmatch(e.name, pattern) or _family(e.name) is None:
                continue
            try:
                d = dst.get(e.name)
            except KeyError:
                d = None
            if d is not None and d.sha256 == e.sha256:
                rep[e.name] = {"skipped": "identical"}
                continue
            if e.name == "stocktwits-anchors":
                byt = {r["value"]["target"]: r for r in (_load(dst, e.name) if d is not None else [])}
                byt.update({r["value"]["target"]: r for r in src.load(e.name)})
                rows = [byt[k] for k in sorted(byt)]
            elif d is not None and d.rows > e.rows:
                rep[e.name] = {"conflict": f"target holds {d.rows} rows, source {e.rows}: not overwritten"}
                continue
            else:
                rows = src.load(e.name)
            ent = dst.store(e.name, kind=e.kind, rows=rows, recipe=e.recipe, tags=e.tags, parked_reason=e.parked_reason,
                            schema=e.schema, in_git=e.in_git)
            rep[e.name] = {"stored": _stored_info(ent), "replaced": d is not None}
            if log:
                log(f"merged {e.name}: {ent.rows} rows")
    return rep


def main(argv: Optional[list[str]] = None) -> int:
    p = argparse.ArgumentParser(prog="python3 -m mandel.data.public_history",
                                description="CFTC, Treasury, Cleveland Fed, SPDR, Sprott, SSGA, Wikipedia, fja05680, FINRA, SEC, StockTwits "
                                            "and Deribit near-expiry histories for the batch-6c underlying claims.")
    p.add_argument("source", choices=["all", "merge"] + list(SOURCES),
                   help="'all' (DEFAULT_ORDER, one process), one source, or 'merge' (copy a --root run into the archive)")
    p.add_argument("--dry-run", action="store_true", help="no request: print the plan (calls, rows, gzip MB, minutes)")
    p.add_argument("--preview", metavar="DIR", help="store into a scratch archive under DIR (DIR/archive, DIR/ledger.jsonl), not the real one")
    p.add_argument("--root", metavar="DIR", help="store into the archive DIR/archive (ledger DIR/ledger.jsonl) instead of the real one, "
                                                 "e.g. StockTwits in parallel with the rest; then 'merge --from DIR'")
    p.add_argument("--from", dest="src", metavar="DIR", help="with 'merge': the DIR of a --root run")
    p.add_argument("--pattern", default="*", help="with 'merge': bundle name pattern (default every bundle of this module)")
    p.add_argument("--limit", type=int, help="bundles (chunks) per fetcher")
    p.add_argument("--max-days", type=int, help="StockTwits / Deribit: days per run (a preview)")
    p.add_argument("--refresh", action="store_true", help="re-fetch stored bundles")
    p.add_argument("--sources", help="with 'all': comma-separated subset of the sources (default: every source, DEFAULT_ORDER)")
    p.add_argument("--first", help="first day (StockTwits, CNMS, Deribit)")
    p.add_argument("--last", help="last day (StockTwits, CNMS, Deribit)")
    a = p.parse_args(argv)
    from mandel.archive import Archive
    if a.preview and a.root:
        p.error("--preview and --root are exclusive")
    store = _root_archive(a.preview or a.root) if (a.preview or a.root) else Archive()
    log = (lambda s: print(s, file=sys.stderr, flush=True))
    try:
        if a.source == "merge":
            if not a.src:
                p.error("merge needs --from DIR (the DIR of a --root run)")
            rep = merge_archive(a.src, store, pattern=a.pattern, log=log)
        elif a.source == "all":
            srcs = [x.strip() for x in a.sources.split(",")] if a.sources else None
            unknown = sorted(set(srcs or ()) - set(SOURCES))
            if unknown:
                p.error(f"unknown sources {unknown}; choose from {sorted(SOURCES)}")
            rep = fetch_all_misc(store, dry_run=a.dry_run, limit=a.limit, sources=srcs, refresh=a.refresh, log=log)
        else:
            kw = {}
            if a.max_days is not None and a.source in ("stocktwits_preopen", "stocktwits_intensity", "deribit_near"):
                kw["max_days"] = a.max_days
            if a.first and a.source in ("stocktwits_preopen", "stocktwits_intensity", "deribit_near", "finra_cnms"):
                kw["first"] = a.first
            if a.last and a.source in ("stocktwits_preopen", "stocktwits_intensity", "deribit_near", "finra_cnms"):
                kw["last"] = a.last
            rep = {}
            with _single_writer(store, rep, a.dry_run, log):         # the single-source path: same lock and index check
                rep.update(SOURCES[a.source](store, dry_run=a.dry_run, limit=a.limit, refresh=a.refresh, log=log, **kw))
    except RuntimeError as e:                                            # another writer holds the run lock
        print(f"refused: {e}", file=sys.stderr)
        return 2
    if a.preview:
        (Path(a.preview) / f"{a.source}.report.json").write_text(json.dumps(rep, indent=1, default=str)[:5_000_000], encoding="utf-8")
    summary = {k: v for k, v in rep.items() if k.startswith("_") and k != "_plan"}
    summary["plan"] = rep.get("_plan", [])[:60]
    summary["results"] = {k: ({kk: vv for kk, vv in v.items() if kk not in ("rows", "_plan")} if isinstance(v, dict) else v)
                          for k, v in rep.items() if not k.startswith("_")}
    print(json.dumps(summary, indent=1, default=str)[:30000])
    if a.dry_run:
        print("\n".join(gitignore_lines()))
    return 1 if not (rep.get("_index_check") or {"ok": True})["ok"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
