"""
Data unlocks: free sources that turn blocked theories into testable ones.
=========================================================================

  * Fama-French daily factors (3-factor, 5-factor 2x3, momentum) and 49 industry portfolios, from Ken French's data library.
  * FOMC meeting calendar 1994 onward from federalreserve.gov (historical pages for <= 2020, the calendars page after):
    decision day = last day of each meeting; scheduled vs unscheduled; conference calls and cancelled meetings excluded.
  * Deribit: BTC-PERPETUAL funding history (hourly, paged), index delivery prices (daily settlement TWAP), daily closes of
    quarterly BTC futures reconstructed by name (last Friday of Mar/Jun/Sep/Dec) through the TradingView chart endpoint.
  * CBOE VIX futures settlement history, one CSV per expiry (Wednesday 30 days before the next month's third Friday).
  * AQR Betting-Against-Beta daily factor (US), from AQR's published workbook.
All rows are point-in-time stamped: observed_at is the date the value refers to, available_at when it could first be known.
"""

from __future__ import annotations

import csv
import io
import re
import time
import zipfile
from datetime import date, datetime, timedelta, timezone
from typing import Iterable, Optional

import httpx

from mandel.data.fetch import Throttle, now_iso
from mandel.data.store import Observation, ObservationStore, iso

UA = {"User-Agent": "mandel-substrate research bot (contact: https://github.com/pointlessbassmusic-byte/polymarket-arbitrage/issues)"}
FRENCH = "https://mba.tuck.dartmouth.edu/pages/faculty/ken.french/ftp/"
MONTHS = {m: i for i, m in enumerate(["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"], 1)}


def _client():
    return httpx.Client(headers=UA, timeout=120, follow_redirects=True)


# ------------------------------------------------------------------ Fama-French --

def parse_french_daily(text: str) -> tuple[list[str], list[dict]]:
    """The first block of YYYYMMDD (daily) or YYYYMM (monthly, day = 'YYYY-MM') rows after a header line that starts with ','.
    Returns (columns, rows) with values in percent; -99.99 / -999 (French's missing codes) are dropped."""
    lines = text.splitlines()
    cols, rows, in_block = None, [], False
    for ln in lines:
        if not in_block:
            if ln.startswith(",") and ln.count(",") >= 1:
                cols = [c.strip() for c in ln.split(",")[1:]]; in_block = True
            continue
        parts = [p.strip() for p in ln.split(",")]
        if not parts or not re.fullmatch(r"\d{8}|\d{6}", parts[0] or ""):
            if rows:
                break
            continue
        d = f"{parts[0][:4]}-{parts[0][4:6]}-{parts[0][6:]}" if len(parts[0]) == 8 else f"{parts[0][:4]}-{parts[0][4:6]}"
        try:
            vals = [float(x) for x in parts[1:len(cols) + 1]]
        except ValueError:
            continue
        rows.append({"day": d, **{c: v for c, v in zip(cols, vals) if v > -99.0}})
    return cols or [], rows


def fetch_french(store: ObservationStore, files: dict | None = None) -> dict:
    files = files or {"ff3-daily": "F-F_Research_Data_Factors_daily_CSV.zip", "ff5-daily": "F-F_Research_Data_5_Factors_2x3_daily_CSV.zip",
                      "momentum-daily": "F-F_Momentum_Factor_daily_CSV.zip", "industry49-daily": "49_Industry_Portfolios_daily_CSV.zip"}
    out = {}
    with _client() as c:
        th = Throttle(1.0)
        for name, f in files.items():
            th.wait()
            r = c.get(FRENCH + f)
            if r.status_code != 200:
                out[name] = f"HTTP {r.status_code}"; continue
            z = zipfile.ZipFile(io.BytesIO(r.content))
            cols, rows = parse_french_daily(z.read(z.namelist()[0]).decode("latin-1"))
            def _ts(dd):
                return dd + "T21:00:00+00:00" if len(dd) == 10 else dd + "-28T21:00:00+00:00"
            e = store.write_bundle(f"french-{name}", (Observation("ken-french", "factor", name, observed_at=_ts(x["day"]),
                                                                  available_at=_ts(x["day"]), value=x) for x in rows),
                                   recipe=f"{FRENCH}{f} (first daily block; percent returns; the library is revised monthly, so available_at is optimistic for recent rows)",
                                   tags=["factors", "fama-french", "history"], schema=f"value = {{day, {', '.join(cols[:8])}{'...' if len(cols) > 8 else ''}}} in percent")
            out[name] = {"rows": e.rows, "columns": len(cols), "first": rows[0]["day"] if rows else None, "last": rows[-1]["day"] if rows else None}
    return out


# ----------------------------------------------------------------------- FOMC --

_H5 = re.compile(r"<h5[^>]*>\s*([^<]*?)\s*-\s*(\d{4})\s*</h5>")


def _month(tok: str) -> int:
    tok = tok.strip()[:3]
    return next(v for k, v in MONTHS.items() if k.startswith(tok))


def parse_meeting_heading(text: str, year: int) -> Optional[dict]:
    """'June 30-July 1 Meeting', 'April/May 30-1 Meeting', 'Jan/Feb 31-1 Meeting', 'March 17-18 (cancelled) Meeting',
    'March 15 (unscheduled) Meeting', 'October 16 (unscheduled)', 'January 9 Conference Call'."""
    t = text.strip()
    unscheduled = "unscheduled" in t.lower()
    if "Conference Call" in t or "cancelled" in t.lower() or ("Meeting" not in t and not unscheduled):
        return None
    core = re.sub(r"\(.*?\)", "", t).replace("Meeting", "").strip()
    m = re.fullmatch(r"([A-Z][a-z]+)(?:/([A-Z][a-z]+))?\s+(\d{1,2})(?:\s*-\s*(?:([A-Z][a-z]+)\s+)?(\d{1,2}))?", core)
    if not m:
        return None
    m1 = _month(m.group(1)); d1 = int(m.group(3))
    m2 = _month(m.group(4) or m.group(2) or m.group(1)); d2 = int(m.group(5) or m.group(3))
    start = date(year, m1, d1); end = date(year, m2, d2)
    return {"start": start.isoformat(), "decision_day": end.isoformat(), "scheduled": not unscheduled, "label": t}


def parse_calendar_page(html: str) -> list[dict]:
    """fomccalendars.htm (recent years): panels '<YYYY> FOMC Meetings' with month / date cells like '27-28', 'Apr/May 30-1', '15*'."""
    out = []
    panels = re.split(r'<h4><a id="\d+">(\d{4}) FOMC Meetings</a></h4>', html)
    for i in range(1, len(panels), 2):
        year = int(panels[i]); body = panels[i + 1]
        starts = [m.start() for m in re.finditer(r"fomc-meeting__month", body)] + [len(body)]
        for a, b in zip(starts[:-1], starts[1:]):
            block = body[a:b]
            mm = re.search(r"<strong>([^<]+)</strong>", block); dd = re.search(r'fomc-meeting__date[^>]*>([^<]+)</div>', block)
            if not mm or not dd:
                continue
            mm, dd = mm.group(1), dd.group(1)
            months = [x.strip() for x in mm.split("/")]
            days = re.findall(r"\d{1,2}", dd)
            if not days or "cancel" in block.lower() or "notation vote" in block.lower():
                continue
            end = date(year, _month(months[-1]), int(days[-1])); start = date(year, _month(months[0]), int(days[0]))
            out.append({"start": start.isoformat(), "decision_day": end.isoformat(), "scheduled": "unscheduled" not in block.lower(),
                        "label": f"{mm} {dd}".strip(), "sep": "*" in dd})
    return out


def fetch_fomc(store: ObservationStore, first_year: int = 1994) -> dict:
    meetings = []
    with _client() as c:
        th = Throttle(1.0)
        th.wait(); cur = c.get("https://www.federalreserve.gov/monetarypolicy/fomccalendars.htm").text
        recent = parse_calendar_page(cur)
        recent_years = {int(m["decision_day"][:4]) for m in recent}
        for y in range(first_year, min(recent_years) if recent_years else 2021):
            th.wait()
            t = c.get(f"https://www.federalreserve.gov/monetarypolicy/fomchistorical{y}.htm").text
            for head, yy in _H5.findall(t):
                m = parse_meeting_heading(head, int(yy))
                if m:
                    meetings.append(m)
        meetings += recent
    seen, uniq = set(), []
    for m in sorted(meetings, key=lambda x: x["decision_day"]):
        if m["decision_day"] in seen:
            continue
        # a scheduled meeting listed under two consecutive dates on the source page (e.g. 2003-09-15/16) is one meeting: keep the later day
        if uniq and m["scheduled"] and uniq[-1]["scheduled"] and (date.fromisoformat(m["decision_day"]) - date.fromisoformat(uniq[-1]["decision_day"])).days < 3:
            uniq[-1] = {**m, "start": uniq[-1]["start"], "label": uniq[-1]["label"] + " / " + m["label"]}
        else:
            uniq.append(m)
        seen.add(m["decision_day"])
    today = now_iso()[:10]
    e = store.write_bundle("fomc-meetings", (Observation("federalreserve.gov", "fomc", m["decision_day"], observed_at=m["decision_day"] + "T18:00:00+00:00",
                                                         available_at=(m["start"] if m["scheduled"] else m["decision_day"]) + "T00:00:00+00:00", value={**m, "future": m["decision_day"] > today})
                                             for m in uniq),
                           recipe="federalreserve.gov fomchistorical<year>.htm h5 headings (<= 2020) + fomccalendars.htm panels (2021+); conference calls and cancelled meetings excluded",
                           tags=["calendar", "fomc", "macro"], schema="value = {start, decision_day, scheduled, label, sep?, future}; available_at = meeting start for scheduled meetings (dates are announced a year ahead; conservative), decision day for unscheduled")
    by_year = {}
    for m in uniq:
        if m["scheduled"]:
            by_year[m["decision_day"][:4]] = by_year.get(m["decision_day"][:4], 0) + 1
    return {"meetings": e.rows, "scheduled_per_year": by_year, "unscheduled": sum(not m["scheduled"] for m in uniq)}


# --------------------------------------------------------------------- Deribit --

def fetch_deribit_history(store: ObservationStore, start_ms: int = 1_556_668_800_000) -> dict:
    """Funding history (hourly, paged 30 days at a time), delivery prices (paged), quarterly futures daily closes by name."""
    out = {}
    with _client() as c:
        th = Throttle(0.25)
        rows, t0, end = [], start_ms, int(time.time() * 1000)
        while t0 < end:
            th.wait()
            r = c.get("https://www.deribit.com/api/v2/public/get_funding_rate_history", params={"instrument_name": "BTC-PERPETUAL", "start_timestamp": t0, "end_timestamp": min(t0 + 30 * 86_400_000, end)})
            if r.status_code == 200:
                rows += r.json().get("result") or []
            t0 += 30 * 86_400_000
        seen, fund = set(), []
        for x in sorted(rows, key=lambda x: x["timestamp"]):
            if x["timestamp"] not in seen:
                seen.add(x["timestamp"]); fund.append(x)
        e = store.write_bundle("deribit-funding-BTC-PERPETUAL", (Observation("deribit", "funding", "BTC-PERPETUAL", observed_at=iso(x["timestamp"] / 1000), available_at=iso(x["timestamp"] / 1000), value=x) for x in fund),
                               recipe="deribit /public/get_funding_rate_history BTC-PERPETUAL, 30-day pages from 2019-05", tags=["crypto", "funding", "history"], schema="value = {timestamp, index_price, prev_index_price, interest_1h, interest_8h}")
        out["funding_hours"] = e.rows
        deliv, off, total = [], 0, None
        while total is None or off < total:
            th.wait()
            r = c.get("https://www.deribit.com/api/v2/public/get_delivery_prices", params={"index_name": "btc_usd", "offset": off, "count": 100})
            res = (r.json().get("result") or {}) if r.status_code == 200 else {}
            data = res.get("data") or []
            total = res.get("records_total", 0) if total is None else total
            if not data:
                break
            deliv += data; off += len(data)
        e = store.write_bundle("deribit-delivery-prices-btc", (Observation("deribit", "delivery", x["date"], observed_at=x["date"] + "T08:00:00+00:00", available_at=x["date"] + "T08:30:00+00:00", value=x) for x in deliv),
                               recipe="deribit /public/get_delivery_prices index_name=btc_usd (paged); settlement = 30-minute TWAP to 08:00 UTC", tags=["crypto", "settlement", "history"], schema="value = {date, delivery_price}")
        out["delivery_days"] = e.rows
        names = []
        d = date(2019, 3, 1)
        while d <= date.today() + timedelta(days=400):
            if d.month in (3, 6, 9, 12):
                last = date(d.year + (d.month == 12), (d.month % 12) + 1, 1) - timedelta(days=1)
                while last.weekday() != 4:
                    last -= timedelta(days=1)
                names.append((f"BTC-{last.day}{last.strftime('%b').upper()}{str(last.year)[2:]}", last))
            d = date(d.year + (d.month == 12), (d.month % 12) + 1, 1)
        fut = []
        for name, exp in names:
            th.wait()
            r = c.get("https://www.deribit.com/api/v2/public/get_tradingview_chart_data", params={"instrument_name": name, "start_timestamp": int(datetime(exp.year - 1, exp.month, 1, tzinfo=timezone.utc).timestamp() * 1000),
                                                                                                  "end_timestamp": int(datetime(exp.year, exp.month, exp.day, 9, tzinfo=timezone.utc).timestamp() * 1000), "resolution": "1D"})
            res = (r.json().get("result") or {}) if r.status_code == 200 else {}
            for ts, cl in zip(res.get("ticks") or [], res.get("close") or []):
                fut.append({"instrument": name, "expiry": exp.isoformat(), "ts_ms": ts, "close": cl})
        e = store.write_bundle("deribit-quarterly-futures-btc", (Observation("deribit", "future", f"{x['instrument']}:{x['ts_ms']}", observed_at=iso(x["ts_ms"] / 1000), available_at=iso(x["ts_ms"] / 1000 + 86400), value=x) for x in fut),
                               recipe="deribit /public/get_tradingview_chart_data resolution=1D for quarterly names BTC-<last Friday of Mar/Jun/Sep/Dec> since 2019", tags=["crypto", "futures", "history"], schema="value = {instrument, expiry, ts_ms (bar open), close}")
        out["futures_rows"] = e.rows; out["futures_contracts"] = len({x["instrument"] for x in fut})
    return out


# -------------------------------------------------------------- CBOE VX futures --

def vx_expiries(first: date = date(2013, 1, 1), last: Optional[date] = None) -> list[date]:
    """Monthly VIX futures final settlement: the Wednesday 30 days before the third Friday of the following month (ignoring holiday shifts;
    the fetcher also tries the day before when the CSV is missing)."""
    last = last or date.today() + timedelta(days=270)
    out, d = [], date(first.year, first.month, 1)
    while d <= last:
        nm = date(d.year + (d.month == 12), (d.month % 12) + 1, 1)
        tf = nm + timedelta(days=(4 - nm.weekday()) % 7) + timedelta(days=14)
        out.append(tf - timedelta(days=30)); d = nm
    return out


def fetch_vx(store: ObservationStore, first: date = date(2013, 1, 1)) -> dict:
    rows, got, missing = [], 0, []
    with _client() as c:
        th = Throttle(0.4)
        for exp in vx_expiries(first):
            ok = False
            for cand in (exp, exp - timedelta(days=1)):
                th.wait()
                r = c.get(f"https://cdn.cboe.com/data/us/futures/market_statistics/historical_data/VX/VX_{cand.isoformat()}.csv")
                if r.status_code == 200 and r.text.startswith("Trade Date"):
                    for x in csv.DictReader(io.StringIO(r.text)):
                        try:
                            rows.append({"expiry": cand.isoformat(), "day": x["Trade Date"], "settle": float(x["Settle"]), "close": float(x["Close"] or 0) or None,
                                         "volume": float(x.get("Total Volume") or 0), "open_interest": float(x.get("Open Interest") or 0)})
                        except (ValueError, KeyError):
                            continue
                    ok = True; got += 1; break
            if not ok:
                missing.append(exp.isoformat())
    e = store.write_bundle("cboe-vx-futures", (Observation("cboe", "vx", f"{x['expiry']}:{x['day']}", observed_at=x["day"] + "T21:15:00+00:00", available_at=x["day"] + "T21:30:00+00:00", value=x) for x in rows),
                           recipe="https://cdn.cboe.com/data/us/futures/market_statistics/historical_data/VX/VX_<expiry>.csv for each monthly expiry", tags=["volatility", "futures", "history"],
                           schema="value = {expiry, day, settle, close, volume, open_interest}")
    return {"contracts": got, "rows": e.rows, "missing": missing[-6:], "n_missing": len(missing)}


# ------------------------------------------------------------------- AQR BAB --

def fetch_aqr_bab(store: ObservationStore) -> dict:
    import openpyxl
    with _client() as c:
        r = c.get("https://www.aqr.com/-/media/AQR/Documents/Insights/Data-Sets/Betting-Against-Beta-Equity-Factors-Daily.xlsx")
    if r.status_code != 200:
        return {"error": f"HTTP {r.status_code}"}
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    ws = wb["BAB Factors"] if "BAB Factors" in wb.sheetnames else wb[wb.sheetnames[0]]
    header, rows = None, []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and row[0] == "DATE":
                header = list(row)
            continue
        if not row or row[0] is None:
            continue
        d = row[0]
        if isinstance(d, str):
            try:
                d = datetime.strptime(d.strip(), "%m/%d/%Y").date()
            except ValueError:
                continue
        if isinstance(d, datetime):
            d = d.date()
        if not isinstance(d, date):
            continue
        usa = row[header.index("USA")] if "USA" in header else None
        if isinstance(usa, (int, float)):
            rows.append({"day": d.isoformat(), "bab_usa": float(usa)})
    e = store.write_bundle("aqr-bab-daily", (Observation("aqr", "factor", "BAB-USA", observed_at=x["day"] + "T21:00:00+00:00", available_at=x["day"] + "T21:00:00+00:00", value=x) for x in rows),
                           recipe="AQR Betting-Against-Beta-Equity-Factors-Daily.xlsx, sheet 'BAB Factors', column USA (decimal daily returns; AQR reconstructs the full history at every update, so past values are revised: not point in time)",
                           tags=["factors", "aqr", "history"], schema="value = {day, bab_usa}")
    return {"rows": e.rows, "first": rows[0]["day"] if rows else None, "last": rows[-1]["day"] if rows else None, "sheets": wb.sheetnames[:6]}


# ------------------------------------------------------------ macro release dates --

def fetch_macro_releases(store: ObservationStore, series=(("CPIAUCSL", "cpi"), ("PAYEMS", "employment-situation"), ("GDP", "gdp"))) -> dict:
    """Release dates as ALFRED vintage dates (the days each data release first existed): CPI, Employment Situation (payrolls), GDP.
    Works around bls.gov refusing this host. A vintage can also be a revision-only date; those are rare and kept (flagged by gap)."""
    out = {}
    with _client() as c:
        th = Throttle(1.0)
        for sid, name in series:
            th.wait()
            t = c.get("https://alfred.stlouisfed.org/series/downloaddata", params={"seid": sid}).text
            dates = sorted(set(re.findall(r'<option[^>]*value="(\d{4}-\d{2}-\d{2})"', t)))
            rows = []
            for i, d in enumerate(dates):
                gap = (date.fromisoformat(d) - date.fromisoformat(dates[i - 1])).days if i else None
                rows.append({"day": d, "series": sid, "release": name, "days_since_previous": gap})
            e = store.write_bundle(f"macro-release-{name}", (Observation("alfred", "release", f"{sid}:{x['day']}", observed_at=x["day"] + "T12:30:00+00:00", available_at=x["day"] + "T12:30:00+00:00", value=x) for x in rows),
                                   recipe=f"https://alfred.stlouisfed.org/series/downloaddata?seid={sid} vintage-date options (release days; 08:30 ET)", tags=["calendar", "macro", name],
                                   schema="value = {day, series, release, days_since_previous}")
            out[name] = {"rows": e.rows, "first": dates[0] if dates else None, "last": dates[-1] if dates else None}
    return out


# --------------------------------------------------------------------- G10 FX --

FX_SPOT = {"EUR": ("DEXUSEU", True), "GBP": ("DEXUSUK", True), "AUD": ("DEXUSAL", True), "NZD": ("DEXUSNZ", True),
           "JPY": ("DEXJPUS", False), "CHF": ("DEXSZUS", False), "CAD": ("DEXCAUS", False), "NOK": ("DEXNOUS", False), "SEK": ("DEXSDUS", False)}
FX_RATE3M = {"USD": "IR3TIB01USM156N", "EUR": "IR3TIB01EZM156N", "GBP": "IR3TIB01GBM156N", "AUD": "IR3TIB01AUM156N", "NZD": "IR3TIB01NZM156N",
             "JPY": "IR3TIB01JPM156N", "CHF": "IR3TIB01CHM156N", "CAD": "IR3TIB01CAM156N", "NOK": "IR3TIB01NOM156N", "SEK": "IR3TIB01SEM156N"}


def fetch_fx(store: ObservationStore) -> dict:
    """FRED daily FX spot (all expressed as USD per unit of foreign currency) and OECD 3-month interbank rates (monthly, percent)."""
    from mandel.data.options import fred_series
    out, spot_rows, rate_rows = {}, [], []
    with _client() as c:
        th = Throttle(0.8)
        for ccy, (sid, usd_per) in FX_SPOT.items():
            for r in fred_series(sid, c, th):
                v = r["value"] if usd_per else (1.0 / r["value"] if r["value"] else None)
                if v:
                    spot_rows.append({"day": r["day"], "ccy": ccy, "usd_per_unit": v, "fred_id": sid})
        for ccy, sid in FX_RATE3M.items():
            for r in fred_series(sid, c, th):
                rate_rows.append({"month": r["day"][:7], "ccy": ccy, "rate3m": r["value"], "fred_id": sid})
    e1 = store.write_bundle("fx-spot-g10", (Observation("fred", "fx", f"{x['ccy']}:{x['day']}", observed_at=x["day"] + "T17:00:00+00:00", available_at=x["day"] + "T21:00:00+00:00", value=x) for x in spot_rows),
                            recipe="FRED " + ", ".join(s for s, _ in FX_SPOT.values()) + " (noon NY buying rates; inverted where FRED quotes foreign per USD)", tags=["fx", "history"], schema="value = {day, ccy, usd_per_unit, fred_id}")
    e2 = store.write_bundle("fx-rates-3m", (Observation("fred", "rate", f"{x['ccy']}:{x['month']}", observed_at=x["month"] + "-01T00:00:00+00:00", available_at=x["month"] + "-28T00:00:00+00:00", value=x) for x in rate_rows),
                            recipe="FRED OECD " + ", ".join(FX_RATE3M.values()) + " (monthly averages; available_at conservatively at month end)", tags=["fx", "rates", "history"], schema="value = {month, ccy, rate3m (percent), fred_id}")
    out["spot_rows"] = e1.rows; out["rate_rows"] = e2.rows; out["currencies"] = sorted({x["ccy"] for x in spot_rows})
    return out


# --------------------------------------------------------------- hourly bars --

HOURLY = ["SPY", "QQQ", "IWM", "TLT", "GLD", "BTC-USD", "ETH-USD", "IBIT", "^VIX"]


def fetch_hourly(store: ObservationStore, symbols=HOURLY) -> dict:
    """Yahoo 1-hour bars (Yahoo serves the last ~730 days). Each run merges with the stored bundle so history accumulates past the window."""
    out = {}
    a = store.archive
    with _client() as c:
        th = Throttle(1.0)
        for sym in symbols:
            th.wait()
            r = c.get(f"https://query1.finance.yahoo.com/v8/finance/chart/{sym}", params={"range": "730d", "interval": "1h"})
            try:
                res = r.json()["chart"]["result"][0]; ts = res["timestamp"]; q = res["indicators"]["quote"][0]
            except Exception:
                out[sym] = f"HTTP {r.status_code}"; continue
            new = {int(t): {"symbol": sym, "t": int(t), "open": q["open"][i], "high": q["high"][i], "low": q["low"][i], "close": q["close"][i], "volume": q["volume"][i]}
                   for i, t in enumerate(ts) if q["close"][i] is not None}
            name = "hourly-" + "".join(ch if ch.isalnum() or ch in "-_." else "-" for ch in sym)
            try:
                for row in a.load(name):
                    new.setdefault(int(row["value"]["t"]), row["value"])
            except Exception:
                pass
            rows = [new[k] for k in sorted(new)]
            e = store.write_bundle(name, (Observation("yahoo", "bar1h", sym, observed_at=iso(x["t"]), available_at=iso(x["t"] + 3600), value=x) for x in rows),
                                   recipe=f"yahoo /v8/finance/chart/{sym}?range=730d&interval=1h, merged with earlier captures (bar time = bar start)", tags=["prices", "hourly"],
                                   schema="value = {symbol, t (bar start, unix), open, high, low, close, volume}")
            out[sym] = {"rows": e.rows, "first": iso(rows[0]["t"])[:10] if rows else None}
    return out
