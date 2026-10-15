"""
Options misc: four free research and history sources behind T-188, T-518, T-222 and T-260.
==========================================================================================

1. ``box-rates-bdg`` (T-188). van Binsbergen, Diamond and Grotteria, "Risk-free interest rates", Journal of Financial
   Economics 143(1), 2022 (doi:10.1016/j.jfineco.2021.06.012); dataset doi:10.35065/b1dk-rp56 on LBS Research Online:
   ``https://lbsresearch.london.edu/id/eprint/2149/1/box_gov_07302019.xlsx``, sheet ``box_gov_l``: year, month, day,
   box_6m, box_12m, box_18m (SPX box-implied rates) and gov_6m, gov_12m, gov_18m (the authors' Treasury fits). The
   ``Notes`` sheet: continuously compounded, annualised, in percent (1.5 = 0.015), linearly interpolated to constant
   maturities. 3,548 trading days, 2004-01-02 .. 2018-03-19.
   This is a PUBLISHED RESEARCH SERIES, not a live feed: the file is dated 2019-07-30 (created 2019-06-27), the
   observations end in March 2018 and are never updated. It was not public in real time, so every row carries the
   same ``available_at``: the date the file was deposited on LBS Research Online, 2021-11-23 13:58 UK time, read from
   the landing page at fetch time (the constant below is the fallback). The file's own timestamps and the NBER w26138
   release (Aug 2019) are earlier, but nothing shows the data file itself was public then; the deposit date is the
   earliest provable one. Licence: CC BY-NC-ND 4.0 (attribution, non-commercial, no derivatives), so the bundle stays
   out of git by default (``in_git=False``); a trading use may not count as non-commercial.

2. ``opec-announcements`` (T-518). Känzig (2021 AER, doi:10.1257/aer.20190964) oil supply news data vintages,
   CC BY 4.0 (cite the paper): ``https://raw.githubusercontent.com/dkaenzig/oilsupplynews/master/
   oilSupplyNewsShocks_2025M12.xlsx``, sheet ``Daily``: column ``Date/Contract`` lists the OPEC announcement (meeting)
   dates, 1983-07-19 .. 2025-12-01 (169 dates); the other columns are the published oil futures surprises (Front,
   1M..12M) and their first principal component (PC, the oil supply surprise). The vintage 2025M12 was committed on
   2026-06-05 and corrected on 2026-06-09 (git log). Two row kinds:
     * ``opec-announcement``: the event flag. ``available_at`` = the announcement date itself, 00:00 UTC (OPEC
       conference dates are scheduled in advance; the file does not mark the unscheduled ones, e.g. the Sunday
       2023-04-02 voluntary cut, so the date itself is the only rule that is safe for all of them). Weekend dates carry
       ``effective_weekday`` (the Monday after); exchange holidays are not adjusted.
     * ``oil-supply-surprise``: the values, ``available_at`` = when that vintage was committed to GitHub (the PC is
       re-estimated on every vintage and differs from the pre-Covid sheet, so the values are not point in time before).

3. ``deribit-option-trades-catalysts`` (T-222). Deribit's public history API (no key; per-IP rate limits):
   ``GET https://history.deribit.com/api/v2/public/get_last_trades_by_currency_and_time?currency=BTC|ETH&kind=option
   &start_timestamp=<ms>&end_timestamp=<ms>&count=1000&include_old=true&sorting=asc``, paged forward (the next page
   starts at the last timestamp seen; trades are de-duplicated by trade_id) while ``has_more``. Fields: trade_id,
   trade_seq, timestamp, instrument_name, price (in coin), mark_price, iv, index_price, amount, direction,
   tick_direction. Trades exist from 2016 for BTC (mark_price null in 2017) and 2019-03 for ETH.
   The registry entry for T-222 names only categories ("ETF decisions, halving, protocol upgrades"); the calendar
   ``CATALYSTS`` below is the unblock sweep's candidate list, frozen here, with each event time checked against chain
   data (mempool.space block times; Ethereum block times from a public RPC; beacon-chain epochs from genesis time).
   Window per (catalyst, currency): [event date - 22 days, event date + 2 days) in UTC days, i.e. from 22 days before
   through the day after. Only trades in the target expiry are kept: the first expiry after the event time that was
   already listed by the end of the entry day (event date - 21 days). Listing times are the instruments'
   ``creation_timestamp`` from ``get_instruments`` (expired=true, ~77 MB per currency, plus the live list). Trade rows:
   ``available_at`` = trade time (the tape is public on execution). One ``catalyst-window`` row per window records the
   event, target expiry, listing time and completeness; its ``available_at`` is ``known_since`` where that date is
   verified (statute or protocol), else the event time itself, so an unverified catalyst never looks known in advance.
   Entries carry ``known_since_hint`` (unverified) for the T-222 harness to check before it freezes its sample.

4. ``deribit-box-monthly`` (T-260). Tardis.dev HTTP data feed, keyless only for the first day of each month
   (any other day answers 401 "only ... the first day of each month"):
   ``GET https://api.tardis.dev/v1/data-feeds/deribit?from=YYYY-MM-01&offset=0&filters=[{"channel":"ticker",
   "symbols":[...]}]`` returns minute 00:00-00:01 UTC as NDJSON lines ``<local ISO timestamp> <Deribit message>``.
   Deribit sends a ticker snapshot for every subscribed instrument when Tardis resubscribes at 00:00, so the legs are
   near-synchronous (the 2024-03-01 probe had all four legs within 0.31 s). Legs per capture day: every BTC expiry 30-90
   days out that was listed then (instrument ``creation_timestamp``), the ``n_side`` strikes at or below and above the
   spot hint (BTC-PERPETUAL 1-hour bar ending at 00:00 UTC, Deribit chart API) that have both a call and a put listed;
   any pair of those strikes forms a box. Kept per instrument: the first ticker message (the snapshot) and the last in
   the minute, raw. ``available_at`` = Tardis's local receive timestamp. Deribit options are on Tardis from 2019-03-30,
   so captures run 2019-04-01 onward (about 91 days to 2026-10). Licence (Tardis.dev FAQ and Terms clause 9, read
   2026-10-15): internal research use only, including free first-of-month data; no redistribution of raw data (only
   aggregates of >= 10 minutes that cannot be reverse-engineered), no use to train or evaluate general-purpose AI
   models, and external service providers (named example: a cloud AI service) must be bound in writing before they get
   the data or results derived from it. So this bundle is never in git, and the raw quotes belong in local Python
   harnesses, not in an LLM context.
   Keyless access is also capped by a data-transfer budget per "account", which for keyless calls is the caller's IP:
   on 2026-10-15 this host (a shared egress proxy) got HTTP 429 code 277 "This account has reached its data transfer
   limit" with retryAfterSeconds 45818 (~12.7 h). The fetcher stops at the first such answer rather than retrying; run it
   from a host with its own IP. A day costs about 24 ticker messages per leg per minute (about 300 KB raw, 40 KB gzip).

Deribit's API documentation states no data licence (its Terms page needs JavaScript and could not be read); its
usage policy is fair use with per-IP limits for public calls and a special 1 request/second limit on get_instruments.
The trades bundle is treated like the deribit-book-* captures: local disk only.

Each ``fetch_*`` takes ``store`` (an Archive, an ObservationStore whose ``.archive`` is used, or None to only return the
rows), ``limit`` and ``dry_run``; a failed or partial fetch never overwrites a stored bundle, and the two Deribit
bundles merge new windows / capture days into what is already stored.
"""

from __future__ import annotations

import io
import json
import re
import time
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from typing import Callable, Iterable, Optional
from zoneinfo import ZoneInfo

import httpx

from mandel.data.fetch import Throttle, now_iso

UA = {"User-Agent": "mandel-substrate research bot (contact: https://github.com/pointlessbassmusic-byte/polymarket-arbitrage/issues)"}
RETRY_STATUS = (429, 500, 502, 503, 504)
MONTHS3 = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
DAY_MS = 86_400_000
_sleep = time.sleep                                    # patched in tests


def _client(timeout: float = 300) -> httpx.Client:
    return httpx.Client(headers=UA, timeout=timeout, follow_redirects=True)


def _archive(store):
    return getattr(store, "archive", store)


def _retry_after(r: httpx.Response) -> Optional[float]:
    """Seconds from the Retry-After header, or Tardis's JSON ``retryAfterSeconds``."""
    try:
        if r.headers.get("retry-after"):
            return float(r.headers["retry-after"])
        j = r.json()
        return float(j["retryAfterSeconds"]) if isinstance(j, dict) and "retryAfterSeconds" in j else None
    except (ValueError, KeyError, TypeError):
        return None


def _get(client: httpx.Client, url: str, throttle: Throttle, params: Optional[dict] = None, headers: Optional[dict] = None,
         retries: int = 4, backoff: float = 1.0, max_wait: float = 120.0) -> tuple[Optional[httpx.Response], dict]:
    """GET with pacing. Retries 429/5xx and transport errors with exponential backoff, waiting at least Retry-After; a
    Retry-After longer than ``max_wait`` (Tardis answers 429 code 277 "data transfer limit" with ~12 hours) is final and
    recorded as ``retry_after``. Any other non-200 answer is final. Returns (response or None, info)."""
    info: dict = {"url": url, "attempts": 0}
    for attempt in range(retries):
        throttle.wait()
        info["attempts"] = attempt + 1
        try:
            r = client.get(url, params=params, headers=headers)
        except httpx.HTTPError as e:
            info["error"] = repr(e)[:200]
            _sleep(backoff * 2 ** attempt)
            continue
        info["status"] = r.status_code
        if r.status_code == 200:
            info.pop("error", None)
            return r, info
        if r.status_code in RETRY_STATUS:
            ra = _retry_after(r)
            if ra is not None and ra > max_wait:
                info.update(retry_after=ra, body=r.text[:300])
                return None, info
            _sleep(max(backoff * 2 ** attempt, ra or 0.0))
            continue
        info["body"] = r.text[:300]
        return None, info
    return None, info


def _iso_ms(ms: float) -> str:
    return datetime.fromtimestamp(ms / 1000, tz=timezone.utc).isoformat(timespec="milliseconds")


def _ms(day_or_iso: str) -> int:
    s = day_or_iso if "T" in day_or_iso else day_or_iso + "T00:00:00+00:00"
    d = datetime.fromisoformat(s.replace("Z", "+00:00"))
    d = d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    return int(round(d.timestamp() * 1000))


def _cap(available_at: str, observed_at: str) -> str:
    """available_at never later than observed_at."""
    return available_at if _ms(available_at) <= _ms(observed_at) else observed_at


def _load_existing(archive, name: str) -> list[dict]:
    try:
        archive.get(name)
        return list(archive.load(name))
    except (KeyError, FileNotFoundError, OSError, ValueError):
        return []


# =========================================================================== 1. box rates (T-188) ==

BDG_XLSX = "https://lbsresearch.london.edu/id/eprint/2149/1/box_gov_07302019.xlsx"
BDG_PAGE = "https://lbsresearch.london.edu/id/eprint/2149/"
BDG_SHEET = "box_gov_l"
BDG_COLUMNS = ("box_6m", "box_12m", "box_18m", "gov_6m", "gov_12m", "gov_18m")
BDG_DEPOSITED = "2021-11-23T13:58:00+00:00"            # landing page "Date Deposited: 23 Nov 2021 13:58" (UK time = UTC in November)
BDG_LICENCE = "https://creativecommons.org/licenses/by-nc-nd/4.0/"
BDG_CITATION = ("van Binsbergen, J.H., Diamond, W. and Grotteria, M. (2021) Daily Data: Risk-free Interest Rates [Dataset], "
                "doi:10.35065/b1dk-rp56; paper: Journal of Financial Economics 143(1), 2022, doi:10.1016/j.jfineco.2021.06.012")
BOX_SCHEMA = ("value = {day, box_6m, box_12m, box_18m, gov_6m, gov_12m, gov_18m}: SPX box-implied rates and the authors' Treasury "
              "fits, continuously compounded, annualised, percent, constant maturities; one row per trading day 2004-01-02..2018-03-19; "
              "available_at = the LBS deposit date for every row (published research series, not public in real time)")


def parse_box_rates(content: bytes, sheet: str = BDG_SHEET) -> list[dict]:
    """Rows {day, box_6m, ..., gov_18m} sorted by day, one per day (the last wins). Rows without a valid date or without
    any numeric rate are dropped; a missing rate is left out of its row."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        header: Optional[list] = None
        out: dict[str, dict] = {}
        for row in ws.iter_rows(values_only=True):
            if header is None:
                if row and str(row[0]).strip().lower() == "year":
                    header = [str(c).strip().lower() if c is not None else "" for c in row]
                continue
            try:
                day = date(int(row[0]), int(row[1]), int(row[2])).isoformat()
            except (TypeError, ValueError, IndexError):
                continue
            v: dict = {"day": day}
            for col in BDG_COLUMNS:
                i = header.index(col) if col in header else None
                x = row[i] if i is not None and i < len(row) else None
                if isinstance(x, (int, float)) and not isinstance(x, bool):
                    v[col] = float(x)
            if len(v) > 1:
                out[day] = v
        return [out[d] for d in sorted(out)]
    finally:
        wb.close()


def parse_bdg_page(html: str) -> dict:
    """Deposit time (UK local -> UTC ISO) and the licence URL from the LBS Research Online landing page."""
    text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", re.sub(r"(?s)<script.*?</script>|<style.*?</style>", "", html)))
    out: dict = {}
    m = re.search(r"Date Deposited:\s*(\d{1,2} [A-Z][a-z]{2} \d{4} \d{2}:\d{2})", text)
    if m:
        local = datetime.strptime(m.group(1), "%d %b %Y %H:%M").replace(tzinfo=ZoneInfo("Europe/London"))
        out["deposited"] = local.astimezone(timezone.utc).isoformat()
    lic = re.search(r"creativecommons\.org/licenses/([a-z-]+)/(\d\.\d)", html)
    if lic:
        out["licence"] = f"https://creativecommons.org/licenses/{lic.group(1)}/{lic.group(2)}/"
    return out


def box_rate_rows(values: list[dict], observed_at: str, available_at: str) -> list[dict]:
    av = _cap(available_at, observed_at)
    return [{"source": "lbs-research-online", "kind": "box-rate", "key": v["day"], "observed_at": observed_at, "available_at": av,
             "value": v} for v in values]


def fetch_box_rates_bdg(store=None, limit: Optional[int] = None, dry_run: bool = False, client: Optional[httpx.Client] = None,
                        min_interval: float = 1.0, retries: int = 4, backoff: float = 1.0, report: Optional[dict] = None,
                        in_git: bool = False) -> list[dict]:
    """Bundle ``box-rates-bdg``: 2 calls (landing page for the deposit date and licence, then the xlsx). Returns archive rows.
    ``limit`` returns only the first N days and then nothing is stored (a preview never replaces the full bundle)."""
    rep = report if report is not None else {}
    if dry_run:
        rep["_plan"] = {"calls": 2, "urls": [BDG_PAGE, BDG_XLSX], "est_rows": 3548, "est_raw_kb": 321, "est_gz_kb": 160}
        return []
    own = client is None
    c = client or _client()
    th = Throttle(min_interval)
    try:
        page, pinfo = _get(c, BDG_PAGE, th, retries=retries, backoff=backoff)
        meta = parse_bdg_page(page.text) if page is not None else {}
        rep["page"] = {**pinfo, **meta}
        r, info = _get(c, BDG_XLSX, th, retries=retries, backoff=backoff)
        rep["xlsx"] = {**info, "bytes": len(r.content) if r is not None else 0}
        if r is None:
            return []
        observed = now_iso()
    finally:
        if own:
            c.close()
    values = parse_box_rates(r.content)
    available = meta.get("deposited") or BDG_DEPOSITED
    licence = meta.get("licence") or BDG_LICENCE
    rows = box_rate_rows(values[:limit] if limit else values, observed, available)
    rep["coverage"] = {"rows": len(values), "first": values[0]["day"] if values else None, "last": values[-1]["day"] if values else None,
                       "available_at": available, "licence": licence}
    archive = _archive(store)
    if archive is not None and rows and not limit:
        e = archive.store("box-rates-bdg", kind="data", rows=rows,
                          recipe=(f"GET {BDG_XLSX} (openpyxl, sheet '{BDG_SHEET}') and {BDG_PAGE} (deposit date, licence) via "
                                  f"mandel.data.options_misc.fetch_box_rates_bdg; httpx, {min_interval}s pacing, {retries} tries with exponential "
                                  f"backoff on 429/5xx. Published research series (file dated 2019-07-30, observations 2004-01-02..2018-03-19, "
                                  f"never updated): not public in real time, so available_at = the LBS Research Online deposit time "
                                  f"{available} for every row; observed_at = fetch time. {BDG_CITATION}. Licence {licence} (non-commercial, "
                                  f"no derivatives): kept out of git"),
                          tags=["options", "rates", "box", "research-series", "history", "not-point-in-time", "T-188"],
                          schema=BOX_SCHEMA, in_git=in_git)
        rep["stored"] = {"rows": e.rows, "bytes": e.bytes, "in_git": e.in_git}
    return rows


# ====================================================================== 2. OPEC announcements (T-518) ==

KANZIG_VINTAGE = "2025M12"
KANZIG_FILE = "oilSupplyNewsShocks_{vintage}.xlsx"
KANZIG_URL = "https://raw.githubusercontent.com/dkaenzig/oilsupplynews/master/" + KANZIG_FILE
KANZIG_COMMITS = "https://api.github.com/repos/dkaenzig/oilsupplynews/commits"
KANZIG_CITATION = ("Känzig, D.R. (2021) The Macroeconomic Effects of Oil Supply News: Evidence from OPEC Announcements, "
                   "American Economic Review 111(4), 1092-1125, doi:10.1257/aer.20190964; data CC BY 4.0, github.com/dkaenzig/oilsupplynews")
# Last commit touching each vintage file (git log, read 2026-10-15); used when the GitHub API is unreachable.
KANZIG_PUBLISHED_GIT = {
    "2017M12": "2021-12-21T01:06:40+01:00", "2020M12": "2021-12-21T01:06:40+01:00", "2021M08": "2022-04-20T12:14:59+01:00",
    "2021M12": "2022-04-20T12:00:40+01:00", "2022M06": "2022-10-14T15:36:27-05:00", "2022M12": "2023-05-22T09:46:31+02:00",
    "2023M06": "2023-12-06T23:38:35+01:00", "2023M12": "2024-07-16T20:06:09+02:00", "2024M06": "2024-12-02T19:19:04-06:00",
    "2024M12": "2025-05-20T22:07:52-05:00", "2025M06": "2025-12-15T23:42:26+01:00", "2025M12": "2026-06-10T01:02:14+02:00",
}
KANZIG_PUBLISHED = {k: datetime.fromisoformat(v).astimezone(timezone.utc).isoformat() for k, v in KANZIG_PUBLISHED_GIT.items()}
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
OPEC_SCHEMA = ("kind opec-announcement: value = {day, weekday, effective_weekday (next Monday for a weekend date; holidays not adjusted), "
               "vintage, in_pre_covid_sheet}, available_at = the announcement date 00:00 UTC | kind oil-supply-surprise: value = {day, vintage, "
               "surprise{Front, 1M..12M, PC} as published, pc_pre_covid}, available_at = the vintage's GitHub commit time")


def _kanzig_sheet(wb, sheet: str) -> list[dict]:
    if sheet not in wb.sheetnames:
        return []
    header: Optional[list] = None
    out: dict[str, dict] = {}
    for row in wb[sheet].iter_rows(values_only=True):
        if header is None:
            if row and str(row[0]).strip().startswith("Date"):
                header = [str(c).strip() if c is not None else "" for c in row]
            continue
        d = row[0] if row else None
        if isinstance(d, datetime):
            d = d.date()
        elif isinstance(d, str):
            try:
                d = date.fromisoformat(d.strip()[:10])
            except ValueError:
                continue
        if not isinstance(d, date):
            continue
        s = {header[i]: float(x) for i, x in enumerate(row) if 0 < i < len(header) and header[i] and isinstance(x, (int, float)) and not isinstance(x, bool)}
        out[d.isoformat()] = s
    return [{"day": k, "surprise": out[k]} for k in sorted(out)]


def parse_kanzig(content: bytes, sheet: str = "Daily", pre_covid_sheet: str = "Daily (pre-Covid)") -> list[dict]:
    """One dict per announcement date in ``sheet``: {day, surprise{column: value}, pc_pre_covid, in_pre_covid_sheet}."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        full = _kanzig_sheet(wb, sheet)
        pre = {r["day"]: r["surprise"] for r in _kanzig_sheet(wb, pre_covid_sheet)}
    finally:
        wb.close()
    return [{**r, "pc_pre_covid": pre.get(r["day"], {}).get("PC"), "in_pre_covid_sheet": r["day"] in pre} for r in full]


def effective_weekday(day: str) -> str:
    d = date.fromisoformat(day)
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d.isoformat()


def opec_rows(parsed: list[dict], vintage: str, observed_at: str, published_at: str) -> list[dict]:
    out = []
    for r in parsed:
        day = r["day"]
        out.append({"source": "kanzig-oilsupplynews", "kind": "opec-announcement", "key": day, "observed_at": observed_at,
                    "available_at": _cap(day + "T00:00:00+00:00", observed_at),
                    "value": {"day": day, "weekday": WEEKDAYS[date.fromisoformat(day).weekday()], "effective_weekday": effective_weekday(day),
                              "vintage": vintage, "in_pre_covid_sheet": r["in_pre_covid_sheet"]}})
    for r in parsed:
        day = r["day"]
        out.append({"source": "kanzig-oilsupplynews", "kind": "oil-supply-surprise", "key": f"{day}:{vintage}", "observed_at": observed_at,
                    "available_at": _cap(max(published_at, day + "T23:59:59+00:00", key=_ms), observed_at),
                    "value": {"day": day, "vintage": vintage, "surprise": r["surprise"], "pc_pre_covid": r["pc_pre_covid"]}})
    return out


def kanzig_published_at(vintage: str, client: httpx.Client, throttle: Throttle, retries: int = 2, backoff: float = 1.0) -> tuple[Optional[str], str]:
    """Newest commit time of the vintage file from the GitHub API, else the git-log constant. Returns (iso or None, basis)."""
    r, info = _get(client, KANZIG_COMMITS, throttle, params={"path": KANZIG_FILE.format(vintage=vintage), "per_page": 1},
                   headers={"Accept": "application/vnd.github+json"}, retries=retries, backoff=backoff)
    if r is not None:
        try:
            d = r.json()[0]["commit"]["committer"]["date"]
            return datetime.fromisoformat(d.replace("Z", "+00:00")).astimezone(timezone.utc).isoformat(), "github-api"
        except (ValueError, KeyError, IndexError, TypeError):
            pass
    if vintage in KANZIG_PUBLISHED:
        return KANZIG_PUBLISHED[vintage], f"git-log constant (API status {info.get('status')})"
    return None, f"unknown (API status {info.get('status')})"


def fetch_opec_announcements(store=None, vintage: str = KANZIG_VINTAGE, limit: Optional[int] = None, dry_run: bool = False,
                             client: Optional[httpx.Client] = None, min_interval: float = 1.0, retries: int = 4, backoff: float = 1.0,
                             report: Optional[dict] = None, in_git: bool = True) -> list[dict]:
    """Bundle ``opec-announcements``: the xlsx plus one GitHub API call for the vintage's commit time. ``limit`` keeps only
    the most recent N dates and then nothing is stored."""
    rep = report if report is not None else {}
    url = KANZIG_URL.format(vintage=vintage)
    if dry_run:
        rep["_plan"] = {"calls": 2, "urls": [url, KANZIG_COMMITS], "est_rows": 2 * 169, "est_gz_kb": 25}
        return []
    own = client is None
    c = client or _client()
    th = Throttle(min_interval)
    try:
        r, info = _get(c, url, th, retries=retries, backoff=backoff)
        rep["xlsx"] = {**info, "bytes": len(r.content) if r is not None else 0}
        if r is None:
            return []
        observed = now_iso()
        published, basis = kanzig_published_at(vintage, c, th, backoff=backoff)
    finally:
        if own:
            c.close()
    parsed = parse_kanzig(r.content)
    if limit:
        parsed = parsed[-limit:]
    pub = published or observed
    rows = opec_rows(parsed, vintage, observed, pub)
    rep["coverage"] = {"dates": len(parsed), "first": parsed[0]["day"] if parsed else None, "last": parsed[-1]["day"] if parsed else None,
                       "since_2009_09": sum(1 for p in parsed if p["day"] >= "2009-09-01"), "weekend": sum(1 for p in parsed if date.fromisoformat(p["day"]).weekday() >= 5),
                       "published_at": pub, "published_basis": basis if published else "fetch time (no commit date)"}
    archive = _archive(store)
    if archive is not None and parsed and not limit:
        e = archive.store("opec-announcements", kind="data", rows=rows,
                          recipe=(f"GET {url} (openpyxl, sheets 'Daily' and 'Daily (pre-Covid)'; column 'Date/Contract' = OPEC announcement dates) and "
                                  f"{KANZIG_COMMITS}?path={KANZIG_FILE.format(vintage=vintage)} (vintage commit time) via "
                                  f"mandel.data.options_misc.fetch_opec_announcements; {min_interval}s pacing, {retries} tries with backoff. "
                                  f"Event rows: available_at = the announcement date 00:00 UTC (scheduled meetings are known in advance; unscheduled ones "
                                  f"are not marked, so the date itself is used for all). Surprise rows: available_at = vintage commit {pub} ({basis}). "
                                  f"{KANZIG_CITATION}"),
                          tags=["calendar", "opec", "oil", "events", "research-series", "T-518"], schema=OPEC_SCHEMA, in_git=in_git)
        rep["stored"] = {"rows": e.rows, "bytes": e.bytes, "in_git": e.in_git}
    return rows


# ======================================================== 3. Deribit option trades around catalysts (T-222) ==

DERIBIT_HISTORY = "https://history.deribit.com/api/v2/public/"
CATALYST_BUNDLE = "deribit-option-trades-catalysts"
_INSTRUMENT = re.compile(r"^(?P<cur>[A-Z]+)-(?P<d>\d{1,2})(?P<mon>[A-Z]{3})(?P<yy>\d{2})-(?P<strike>\d+(?:d\d+)?)-(?P<cp>[CP])$")

# Frozen candidate calendar (unblock sweep T-222 recheck list). ``time`` = the event moment in UTC, checked on chain where
# there is a chain event; ``known_since`` is set only when the date was fixed by statute or protocol and the fixing date is
# verified; ``known_since_hint`` is unverified and must be checked by the harness before the sample is frozen.
CATALYSTS: list[dict] = [
    {"id": "btc-bch-fork-2017", "currencies": ["BTC"], "category": "fork", "event": "Bitcoin Cash hard fork (UAHF)", "date": "2017-08-01",
     "time": "2017-08-01T13:16:14+00:00", "time_basis": "BTC block 478558, the last common block (mempool.space)",
     "known_since": None, "known_since_hint": "UAHF activation time 2017-08-01 12:20 UTC published by Bitcoin ABC in summer 2017 (unverified)"},
    {"id": "btc-btg-fork-2017", "currencies": ["BTC"], "category": "fork", "event": "Bitcoin Gold fork", "date": "2017-10-24",
     "time": "2017-10-24T01:20:39+00:00", "time_basis": "BTC block 491407 (mempool.space)",
     "known_since": None, "known_since_hint": "fork height 491407 announced in 2017 (unverified)"},
    {"id": "bch-hard-fork-2018", "currencies": ["BTC"], "category": "fork", "event": "Bitcoin Cash November upgrade (ABC/SV split)", "date": "2018-11-15",
     "time": "2018-11-15T16:40:00+00:00", "time_basis": "BCH upgrade activation at median time past 1542300000",
     "known_since": None, "known_since_hint": "BCH semiannual upgrade date (Nov 15) from the 2017 upgrade schedule (unverified)"},
    {"id": "eth-istanbul-2019", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Istanbul hard fork", "date": "2019-12-08",
     "time": "2019-12-08T00:25:09+00:00", "time_basis": "ETH block 9,069,000 (public RPC)",
     "known_since": None, "known_since_hint": "block chosen on an All Core Devs call in autumn 2019 (unverified)"},
    {"id": "btc-halving-2020", "currencies": ["BTC"], "category": "halving", "event": "Bitcoin halving (block 630,000)", "date": "2020-05-11",
     "time": "2020-05-11T19:23:43+00:00", "time_basis": "BTC block 630000 (mempool.space)",
     "known_since": "2020-03-12", "known_since_basis": "height fixed by protocol since 2009; the date was a few-day estimate; set 60 days ahead"},
    {"id": "us-election-2020", "currencies": ["BTC", "ETH"], "category": "macro", "event": "US presidential election", "date": "2020-11-03",
     "time": "2020-11-04T05:00:00+00:00", "time_basis": "midnight US Eastern after Election Day",
     "known_since": "2019-11-03", "known_since_basis": "date fixed by statute (3 U.S.C. 1); set one year ahead"},
    {"id": "eth-beacon-genesis-2020", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum beacon chain genesis", "date": "2020-12-01",
     "time": "2020-12-01T12:00:23+00:00", "time_basis": "beacon genesis_time 1606824023",
     "known_since": "2020-11-24", "known_since_basis": "the 16,384-validator threshold was met on 2020-11-24; from 2020-11-04 the date was conditional"},
    {"id": "eth-berlin-2021", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Berlin hard fork", "date": "2021-04-15",
     "time": "2021-04-15T10:07:03+00:00", "time_basis": "ETH block 12,244,000 (public RPC)",
     "known_since": None, "known_since_hint": "Ethereum Foundation blog announcement 2021-03-08 (unverified)"},
    {"id": "eth-london-2021", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum London hard fork (EIP-1559)", "date": "2021-08-05",
     "time": "2021-08-05T12:33:42+00:00", "time_basis": "ETH block 12,965,000 (public RPC)",
     "known_since": None, "known_since_hint": "Ethereum Foundation blog announcement 2021-07-15, 21 days ahead (unverified)"},
    {"id": "btc-taproot-2021", "currencies": ["BTC"], "category": "protocol", "event": "Bitcoin Taproot activation", "date": "2021-11-14",
     "time": "2021-11-14T05:15:27+00:00", "time_basis": "BTC block 709632 (mempool.space)",
     "known_since": "2021-06-12", "known_since_basis": "Speedy Trial lock-in at block 687284 (2021-06-12 12:04 UTC, mempool.space) fixed the activation height"},
    {"id": "eth-merge-2022", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Merge (terminal total difficulty)", "date": "2022-09-15",
     "time": "2022-09-15T06:42:42+00:00", "time_basis": "ETH block 15,537,393 reached TTD (public RPC)",
     "known_since": None, "known_since_hint": "Ethereum Foundation mainnet announcement with the TTD 2022-08-24; tentative mid-September date from mid-July (unverified)"},
    {"id": "eth-shapella-2023", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Shanghai/Capella upgrade", "date": "2023-04-12",
     "time": "2023-04-12T22:27:35+00:00", "time_basis": "beacon epoch 194048 = genesis + 194048 x 384 s",
     "known_since": None, "known_since_hint": "date set on an All Core Devs call in March 2023 (unverified)"},
    {"id": "btc-spot-etf-2024", "currencies": ["BTC"], "category": "etf", "event": "US spot bitcoin ETF decision (ARK 21Shares final deadline)", "date": "2024-01-10",
     "time": "2024-01-10T21:00:00+00:00", "time_basis": "SEC approval order after the US close (approximate)",
     "known_since": "2023-12-01", "known_since_basis": "statutory 240-day final deadline of the 19b-4 filing, widely reported from August 2023; set conservatively"},
    {"id": "eth-dencun-2024", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Dencun upgrade", "date": "2024-03-13",
     "time": "2024-03-13T13:55:35+00:00", "time_basis": "beacon epoch 269568 = genesis + 269568 x 384 s",
     "known_since": None, "known_since_hint": "date set on the 2024-02-08 All Core Devs call (unverified)"},
    {"id": "btc-halving-2024", "currencies": ["BTC"], "category": "halving", "event": "Bitcoin halving (block 840,000)", "date": "2024-04-20",
     "time": "2024-04-20T00:09:27+00:00", "time_basis": "BTC block 840000 (mempool.space)",
     "known_since": "2024-02-20", "known_since_basis": "height fixed by protocol since 2009; the date was a few-day estimate; set 60 days ahead"},
    {"id": "eth-spot-etf-2024", "currencies": ["ETH"], "category": "etf", "event": "US spot ether ETF 19b-4 decision (VanEck final deadline)", "date": "2024-05-23",
     "time": "2024-05-23T21:00:00+00:00", "time_basis": "SEC approval order after the US close (approximate)",
     "known_since": "2024-04-01", "known_since_basis": "statutory 240-day final deadline of the 19b-4 filing; set conservatively"},
    {"id": "us-election-2024", "currencies": ["BTC", "ETH"], "category": "macro", "event": "US presidential election", "date": "2024-11-05",
     "time": "2024-11-06T05:00:00+00:00", "time_basis": "midnight US Eastern after Election Day",
     "known_since": "2023-11-05", "known_since_basis": "date fixed by statute (3 U.S.C. 1); set one year ahead"},
    {"id": "eth-pectra-2025", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Pectra upgrade", "date": "2025-05-07",
     "time": "2025-05-07T10:05:11+00:00", "time_basis": "beacon epoch 364032 = genesis + 364032 x 384 s",
     "known_since": None, "known_since_hint": "date set on the 2025-04-03 All Core Devs call (unverified)"},
    {"id": "eth-fusaka-2025", "currencies": ["ETH"], "category": "protocol", "event": "Ethereum Fusaka upgrade", "date": "2025-12-03",
     "time": "2025-12-03T21:49:11+00:00", "time_basis": "beacon epoch 411392 = genesis + 411392 x 384 s",
     "known_since": None, "known_since_hint": "mainnet date set in autumn 2025 (unverified)"},
]
BEACON_GENESIS = 1_606_824_023                       # 2020-12-01T12:00:23Z
BEACON_EPOCHS = {"eth-shapella-2023": 194048, "eth-dencun-2024": 269568, "eth-pectra-2025": 364032, "eth-fusaka-2025": 411392}

# Rough option trades per UTC day (all expiries), from 2-hour samples at 12:00 UTC on each window's entry and exit day
# (2026-10-15); used only for the dry-run call estimate. Bursts make single samples noisy (ETH 2024-02-21 extrapolated to 84k).
TRADES_PER_DAY = {("BTC", 2017): 250, ("BTC", 2018): 500, ("BTC", 2019): 600, ("BTC", 2020): 5000, ("BTC", 2021): 3000,
                  ("BTC", 2022): 8000, ("BTC", 2023): 10000, ("BTC", 2024): 16000, ("BTC", 2025): 15000,
                  ("ETH", 2019): 200, ("ETH", 2020): 4000, ("ETH", 2021): 9000, ("ETH", 2022): 13000, ("ETH", 2023): 9000,
                  ("ETH", 2024): 16000, ("ETH", 2025): 8000}
TRADE_SCHEMA = ("kind option-trade: value = raw Deribit trade {trade_id, trade_seq, timestamp (ms), instrument_name, price (coin), mark_price, iv, "
                "index_price, amount, contracts?, direction, tick_direction, block_trade_id?} + {catalyst, currency, expiry, expiry_code, strike, type}; "
                "available_at = trade time | kind catalyst-window: value = {catalyst, currency, event, category, event_date, event_time, entry_day, "
                "exit_day, window_start, window_end, target_expiry, target_expiry_code, target_listed_at, trades_seen, trades_kept, pages, complete, "
                "known_since, known_since_basis|known_since_hint}; available_at = known_since if verified else the event time")


def parse_instrument(name: str) -> Optional[dict]:
    m = _INSTRUMENT.match(name or "")
    if not m:
        return None
    try:
        exp = date(2000 + int(m.group("yy")), MONTHS3.index(m.group("mon")) + 1, int(m.group("d")))
    except ValueError:
        return None
    return {"currency": m.group("cur"), "expiry": exp.isoformat(), "expiry_code": f"{m.group('d')}{m.group('mon')}{m.group('yy')}",
            "strike": float(m.group("strike").replace("d", ".")), "type": "call" if m.group("cp") == "C" else "put"}


def expiry_code(ts_ms: int) -> str:
    d = datetime.fromtimestamp(ts_ms / 1000, tz=timezone.utc)
    return f"{d.day}{MONTHS3[d.month - 1]}{d.year % 100:02d}"


def compact_instrument(x: dict) -> dict:
    return {k: x.get(k) for k in ("instrument_name", "strike", "option_type", "expiration_timestamp", "creation_timestamp", "settlement_period")}


def deribit_instruments(currency: str, client: httpx.Client, throttle: Throttle, expired: bool = True, retries: int = 3,
                        backoff: float = 2.0) -> list[dict]:
    """Compact option instruments (name, strike, type, expiration and creation timestamps) from history.deribit.com.
    ``expired=True`` returns every expired option since 2016 (BTC: ~124k instruments, ~77 MB, ~25 s)."""
    r, _ = _get(client, DERIBIT_HISTORY + "get_instruments", throttle, params={"currency": currency, "kind": "option", "expired": str(expired).lower()},
                retries=retries, backoff=backoff)
    if r is None:
        return []
    try:
        return [compact_instrument(x) for x in (r.json().get("result") or [])]
    except ValueError:
        return []


def deribit_option_instruments(currency: str, client: httpx.Client, throttle: Throttle, retries: int = 3, backoff: float = 2.0) -> list[dict]:
    """Expired plus live option instruments (live ones are needed for expiries that have not expired yet), de-duplicated by name."""
    seen: dict[str, dict] = {}
    for expired in (True, False):
        for x in deribit_instruments(currency, client, throttle, expired=expired, retries=retries, backoff=backoff):
            seen.setdefault(x["instrument_name"], x)
    return list(seen.values())


def expiry_listings(instruments: Iterable[dict]) -> dict[int, dict]:
    """{expiration_ts: {listed_at (earliest instrument creation, ms), n_instruments, settlement_period}}."""
    out: dict[int, dict] = {}
    periods: dict[int, Counter] = defaultdict(Counter)
    for x in instruments:
        e, c = x.get("expiration_timestamp"), x.get("creation_timestamp")
        if e is None or c is None:
            continue
        d = out.setdefault(int(e), {"listed_at": int(c), "n_instruments": 0})
        d["listed_at"] = min(d["listed_at"], int(c))
        d["n_instruments"] += 1
        periods[int(e)][x.get("settlement_period")] += 1
    for e, d in out.items():
        d["settlement_period"] = periods[e].most_common(1)[0][0]
    return out


def catalyst_windows(catalysts: Optional[list[dict]] = None, listings: Optional[dict[str, dict[int, dict]]] = None,
                     currencies: Optional[Iterable[str]] = None, before_days: int = 22, after_days: int = 1, entry_days: int = 21,
                     n_expiries: int = 1) -> list[dict]:
    """One window per (catalyst, currency): [date - before_days, date + after_days + 1) in UTC days. With ``listings``
    ({currency: expiry_listings(...)}) each window gets its target expiries: the first ``n_expiries`` expiring after the
    event time whose first instrument was created before the end of the entry day (date - entry_days)."""
    cats = catalysts if catalysts is not None else CATALYSTS
    want = set(currencies) if currencies else None
    out = []
    for cat in cats:
        d = date.fromisoformat(cat["date"])
        ev_ms = _ms(cat.get("time") or (d + timedelta(days=1)).isoformat())
        entry, exit_ = d - timedelta(days=entry_days), d - timedelta(days=1)
        entry_end = _ms((entry + timedelta(days=1)).isoformat())
        for cur in cat["currencies"]:
            if want and cur not in want:
                continue
            w = {"catalyst": cat["id"], "currency": cur, "event": cat["event"], "category": cat.get("category"), "event_date": cat["date"],
                 "event_time": _iso_ms(ev_ms), "entry_day": entry.isoformat(), "exit_day": exit_.isoformat(),
                 "start_ms": _ms((d - timedelta(days=before_days)).isoformat()), "end_ms": _ms((d + timedelta(days=after_days + 1)).isoformat()),
                 "known_since": cat.get("known_since"), "targets": None}
            if listings is not None and cur in listings:
                cand = sorted(e for e, L in listings[cur].items() if e > ev_ms and L["listed_at"] < entry_end)
                w["targets"] = [{"expiry_ts": e, "expiry": _iso_ms(e)[:10], "expiry_code": expiry_code(e),
                                 "listed_at": _iso_ms(listings[cur][e]["listed_at"]), "settlement_period": listings[cur][e].get("settlement_period")}
                                for e in cand[:n_expiries]]
            out.append(w)
    return out


def trades_in_window(currency: str, start_ms: int, end_ms: int, client: httpx.Client, throttle: Throttle,
                     keep: Callable[[dict], bool] = lambda t: True, max_pages: Optional[int] = None, retries: int = 4,
                     backoff: float = 1.0) -> tuple[list[dict], dict]:
    """All option trades of ``currency`` in [start_ms, end_ms), paged forward; returns (kept trades, info{pages, seen,
    complete, error?}). ``complete`` is False when a page failed after retries or ``max_pages`` stopped the walk."""
    kept: list[dict] = []
    info: dict = {"pages": 0, "seen": 0, "complete": False}
    cursor, boundary = start_ms, set()
    while True:
        if max_pages is not None and info["pages"] >= max_pages:
            info["stopped"] = "max_pages"
            break
        r, ginfo = _get(client, DERIBIT_HISTORY + "get_last_trades_by_currency_and_time", throttle,
                        params={"currency": currency, "kind": "option", "start_timestamp": cursor, "end_timestamp": end_ms - 1,
                                "count": 1000, "include_old": "true", "sorting": "asc"}, retries=retries, backoff=backoff)
        if r is None:
            info["error"] = {k: ginfo.get(k) for k in ("status", "error", "body", "attempts", "retry_after") if ginfo.get(k) is not None}
            break
        try:
            res = r.json().get("result") or {}
        except ValueError:
            info["error"] = {"status": r.status_code, "error": "invalid JSON"}
            break
        trades = res.get("trades") or []
        info["pages"] += 1
        new = [t for t in trades if t.get("trade_id") not in boundary and start_ms <= int(t.get("timestamp", -1)) < end_ms]
        info["seen"] += len(new)
        kept.extend(t for t in new if keep(t))
        if not res.get("has_more") or not trades:
            info["complete"] = True
            break
        last = max(int(t["timestamp"]) for t in trades)
        nxt = last if (new and last > cursor) else last + 1        # a page with nothing new (> 1000 trades in one ms) steps past it
        if nxt >= end_ms:
            info["complete"] = True
            break
        boundary = {t.get("trade_id") for t in trades if int(t["timestamp"]) >= nxt}
        cursor = nxt
    return kept, info


def _trade_row(t: dict, catalyst: str, currency: str, observed_at: str) -> dict:
    p = parse_instrument(t.get("instrument_name", "")) or {}
    v = {"catalyst": catalyst, "currency": currency, "expiry": p.get("expiry"), "expiry_code": p.get("expiry_code"), "strike": p.get("strike"),
         "type": p.get("type"), **t}
    return {"source": "deribit-history", "kind": "option-trade", "key": f"{currency}:{t.get('trade_id')}", "observed_at": observed_at,
            "available_at": _cap(_iso_ms(int(t["timestamp"])), observed_at), "value": v}


def _window_row(w: dict, cat: dict, info: dict, kept: int, observed_at: str) -> dict:
    tgt = (w.get("targets") or [None])[0] or {}
    av = cat.get("known_since") + "T00:00:00+00:00" if cat.get("known_since") else w["event_time"]
    v = {"catalyst": w["catalyst"], "currency": w["currency"], "event": w["event"], "category": w["category"], "event_date": w["event_date"],
         "event_time": w["event_time"], "time_basis": cat.get("time_basis"), "entry_day": w["entry_day"], "exit_day": w["exit_day"],
         "window_start": _iso_ms(w["start_ms"]), "window_end": _iso_ms(w["end_ms"]), "target_expiry": tgt.get("expiry"),
         "target_expiry_code": tgt.get("expiry_code"), "target_listed_at": tgt.get("listed_at"), "target_settlement_period": tgt.get("settlement_period"),
         "targets": [x["expiry_code"] for x in (w.get("targets") or [])], "trades_seen": info.get("seen"), "trades_kept": kept,
         "pages": info.get("pages"), "complete": bool(info.get("complete")), "known_since": cat.get("known_since"),
         "known_since_basis": cat.get("known_since_basis"), "known_since_hint": cat.get("known_since_hint")}
    return {"source": "mandel-catalysts", "kind": "catalyst-window", "key": f"{w['catalyst']}:{w['currency']}", "observed_at": observed_at,
            "available_at": _cap(av, observed_at), "value": v}


def _plan_estimate(windows: list[dict]) -> dict:
    per, calls, trades = [], 0, 0
    for w in windows:
        days = (w["end_ms"] - w["start_ms"]) / DAY_MS
        yr = int(w["event_date"][:4])
        tpd = TRADES_PER_DAY.get((w["currency"], yr)) or TRADES_PER_DAY.get((w["currency"], min(max(yr, 2017), 2025)), 5000)
        n = int(days * tpd)
        c = max(1, -(-n // 1000))
        per.append({"window": f"{w['catalyst']}:{w['currency']}", "days": days, "est_trades": n, "est_calls": c})
        calls += c
        trades += n
    return {"windows": len(windows), "est_trades_seen": trades, "est_calls": calls, "per_window": per}


def fetch_deribit_catalyst_trades(store=None, catalysts: Optional[list[dict]] = None, currencies: Optional[Iterable[str]] = None,
                                  limit: Optional[int] = None, dry_run: bool = False, max_pages: Optional[int] = None,
                                  client: Optional[httpx.Client] = None, instruments: Optional[dict[str, list[dict]]] = None,
                                  n_expiries: int = 1, before_days: int = 22, after_days: int = 1, min_interval: float = 0.25,
                                  retries: int = 4, backoff: float = 1.0, report: Optional[dict] = None, skip_stored: bool = False) -> list[dict]:
    """Bundle ``deribit-option-trades-catalysts`` (local disk only). Returns the archive rows of this run.

    ``limit`` caps the number of (catalyst, currency) windows; ``max_pages`` caps pages per window (a preview: such
    windows are marked incomplete and are never stored); ``instruments`` ({currency: compact instrument list}) skips the
    large get_instruments download; ``dry_run`` makes no request and puts the plan and call estimate in ``report``.
    With ``store``, complete windows replace the same windows in the stored bundle and other stored windows are kept; the
    store also happens when the run is interrupted. ``skip_stored`` skips windows already stored complete (resume a long run)."""
    cats = catalysts if catalysts is not None else CATALYSTS
    by_id = {c["id"]: c for c in cats}
    rep = report if report is not None else {}
    archive = _archive(store)
    windows = catalyst_windows(cats, None, currencies, before_days, after_days)
    if archive is not None and skip_stored:
        have = {r["key"] for r in _load_existing(archive, CATALYST_BUNDLE) if r["kind"] == "catalyst-window" and r["value"].get("complete")}
        rep["_already_stored"] = sorted(k for k in have if any(f"{w['catalyst']}:{w['currency']}" == k for w in windows))
        windows = [w for w in windows if f"{w['catalyst']}:{w['currency']}" not in have]
    if limit is not None:
        windows = windows[:limit]
    if dry_run:
        rep["_plan"] = {**_plan_estimate(windows), "instrument_calls": 2 * len({w["currency"] for w in windows}) if instruments is None else 0,
                        "instrument_mb_each": 77, "min_interval": min_interval}
        return []
    own = client is None
    c = client or _client()
    th = Throttle(min_interval)
    th_inst = Throttle(max(min_interval, 1.1))              # get_instruments: 1 request/second sustained
    rows: list[dict] = []
    done: set[str] = set()
    try:
        insts = dict(instruments or {})
        for cur in sorted({w["currency"] for w in windows} - set(insts)):
            insts[cur] = deribit_option_instruments(cur, c, th_inst, retries=retries, backoff=max(backoff, 1.0) if backoff else 0)
        listings = {cur: expiry_listings(xs) for cur, xs in insts.items() if xs}
        windows = [w2 for w in windows for w2 in catalyst_windows([by_id[w["catalyst"]]], listings, [w["currency"]], before_days, after_days,
                                                                   n_expiries=n_expiries)]
        for w in windows:
            key = f"{w['catalyst']}:{w['currency']}"
            if not w.get("targets"):
                rep[key] = {"skipped": "no listed expiry after the event (instrument list missing or empty)"}
                continue
            codes = {t["expiry_code"] for t in w["targets"]}
            kept, info = trades_in_window(w["currency"], w["start_ms"], w["end_ms"], c, th,
                                          keep=lambda t, codes=codes: (parse_instrument(t.get("instrument_name", "")) or {}).get("expiry_code") in codes,
                                          max_pages=max_pages, retries=retries, backoff=backoff)
            observed = now_iso()
            wrows = [_trade_row(t, w["catalyst"], w["currency"], observed) for t in kept]
            wrows.append(_window_row(w, by_id[w["catalyst"]], info, len(kept), observed))
            rows.extend(wrows)
            rep[key] = {**info, "kept": len(kept), "targets": sorted(codes)}
            if info["complete"]:
                done.add(key)
            elif (info.get("error") or {}).get("retry_after"):
                rep["_aborted"] = {"window": key, "retry_after_s": info["error"]["retry_after"], "body": info["error"].get("body")}
                break
    finally:                                                # also on an interrupt: complete windows are kept
        if own:
            c.close()
        if archive is not None and done:
            _store_catalyst_bundle(archive, rows, done, before_days, after_days, min_interval, retries, rep)
    return rows


def _store_catalyst_bundle(archive, rows: list[dict], done: set[str], before_days: int, after_days: int, min_interval: float,
                           retries: int, rep: dict) -> None:
    name = CATALYST_BUNDLE
    fresh = [r for r in rows if f"{r['value']['catalyst']}:{r['value']['currency']}" in done]
    old = [r for r in _load_existing(archive, name) if f"{r['value'].get('catalyst')}:{r['value'].get('currency')}" not in done]
    allrows = sorted(old + fresh, key=lambda r: (r["value"].get("catalyst") or "", r["value"].get("currency") or "", r["kind"] != "catalyst-window",
                                                  r["value"].get("timestamp") or 0, str(r["value"].get("trade_id") or "")))
    windows_stored = sorted({f"{r['value']['catalyst']}:{r['value']['currency']}" for r in allrows if r["kind"] == "catalyst-window"})
    e = archive.store(name, kind="data", rows=allrows,
                      recipe=(f"GET {DERIBIT_HISTORY}get_last_trades_by_currency_and_time?currency=<BTC|ETH>&kind=option&start_timestamp&end_timestamp"
                              f"&count=1000&include_old=true&sorting=asc, paged forward, for each window in mandel.data.options_misc.CATALYSTS "
                              f"[event date - {before_days} d, event date + {after_days + 1} d) UTC, keeping only the first expiry after the event that was "
                              f"listed by the end of the entry day (event - 21 d; listing = min creation_timestamp from {DERIBIT_HISTORY}get_instruments"
                              f", expired=true and false); via mandel.data.options_misc.fetch_deribit_catalyst_trades; {min_interval}s pacing, {retries} tries with "
                              f"exponential backoff on 429/5xx. Trade available_at = trade time; window available_at = verified known_since else the "
                              f"event time. Windows stored: {', '.join(windows_stored)}. Deribit states no data licence: local disk only"),
                      tags=["options", "crypto", "deribit", "trades", "events", "T-222"], schema=TRADE_SCHEMA, in_git=False)
    rep["stored"] = {"rows": e.rows, "bytes": e.bytes, "windows": windows_stored}


# ====================================================================== 4. Tardis monthly boxes (T-260) ==

TARDIS_FEED = "https://api.tardis.dev/v1/data-feeds/deribit"
BOX_BUNDLE = "deribit-box-monthly"
TARDIS_FIRST_DAY = "2019-04-01"                       # Deribit options are on Tardis from 2019-03-30
DERIBIT_CHART = "https://www.deribit.com/api/v2/public/get_tradingview_chart_data"
TARDIS_LICENCE = ("Tardis.dev licence (FAQ 'What can I use the data for?' and Terms clause 9, read 2026-10-15): internal research use only, "
                  "free first-of-month data included; no redistribution or resale of raw data (aggregates >= 10 minutes that cannot be "
                  "reverse-engineered excepted); no training or evaluation of general-purpose AI models; bind external service providers "
                  "(e.g. cloud AI services) in writing before giving them the data or results derived from it")
BOX_TICKER_SCHEMA = ("kind option-ticker: value = {capture_day, instrument, expiry, expiry_code, dte, strike, type, snap ('first' = the 00:00 "
                     "resubscribe snapshot, 'last' = last message in minute 00:00-00:01), n_messages, local_timestamp, spot_hint, ticker (raw "
                     "Deribit ticker: best_bid_price, best_ask_price, best_bid_amount, best_ask_amount, mark_price, mark_iv, bid_iv, ask_iv, "
                     "underlying_price, underlying_index, index_price, interest_rate, open_interest, greeks, timestamp, ...)}; prices in BTC; "
                     "available_at = Tardis local receive time | kind capture: value = {capture_day, spot_hint, expiries, strikes, symbols, "
                     "received, status}; available_at = capture day 00:01 UTC")


def box_capture_days(first: str = TARDIS_FIRST_DAY, last: Optional[str] = None) -> list[str]:
    """First day of each month from ``first`` through ``last`` (default today), as YYYY-MM-01."""
    end = date.fromisoformat(last) if last else datetime.now(timezone.utc).date()
    d = date.fromisoformat(first)
    d = date(d.year, d.month, 1) if d.day == 1 else date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    out = []
    while d <= end:
        out.append(d.isoformat())
        d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return out


def spot_hint(day: str, client: httpx.Client, throttle: Throttle, retries: int = 4, backoff: float = 1.0) -> Optional[float]:
    """Close of the BTC-PERPETUAL 1-hour bar that ends at ``day`` 00:00 UTC (no later price is used)."""
    t = _ms(day)
    r, _ = _get(client, DERIBIT_CHART, throttle, params={"instrument_name": "BTC-PERPETUAL", "start_timestamp": t - 4 * 3_600_000,
                                                         "end_timestamp": t, "resolution": "60"}, retries=retries, backoff=backoff)
    if r is None:
        return None
    try:
        res = r.json().get("result") or {}
    except ValueError:
        return None
    bars = [(int(ts), cl) for ts, cl in zip(res.get("ticks") or [], res.get("close") or []) if cl is not None and int(ts) + 3_600_000 <= t]
    return float(max(bars)[1]) if bars else None


def select_box_legs(instruments: Iterable[dict], day: str, spot: float, dte_min: float = 30, dte_max: float = 90, n_side: int = 2) -> list[dict]:
    """For every expiry ``dte_min``..``dte_max`` days after ``day`` 00:00 UTC that was listed then: the ``n_side`` strikes at or
    below ``spot`` and the ``n_side`` above it that have both a call and a put listed at that moment."""
    t = _ms(day)
    by_exp: dict[int, dict[float, dict[str, str]]] = defaultdict(lambda: defaultdict(dict))
    for x in instruments:
        e, c = x.get("expiration_timestamp"), x.get("creation_timestamp")
        if e is None or c is None or not (int(c) <= t < int(e)):
            continue
        by_exp[int(e)][float(x["strike"])][x.get("option_type")] = x["instrument_name"]
    out = []
    for e in sorted(by_exp):
        dte = (e - t) / DAY_MS
        if not (dte_min <= dte <= dte_max):
            continue
        both = sorted(k for k, v in by_exp[e].items() if "call" in v and "put" in v)
        below = [k for k in both if k <= spot][-n_side:]
        above = [k for k in both if k > spot][:n_side]
        if not below or not above:
            continue
        strikes = below + above
        names = [by_exp[e][k][cp] for k in strikes for cp in ("call", "put")]
        out.append({"expiry": _iso_ms(e)[:10], "expiry_code": expiry_code(e), "expiry_ts": e, "dte": round(dte, 3), "strikes": strikes,
                    "instruments": names})
    return out


def tardis_filters(symbols: list[str]) -> str:
    return json.dumps([{"channel": "ticker", "symbols": list(symbols)}], separators=(",", ":"))


def parse_tardis(text: str) -> list[tuple[str, dict]]:
    """NDJSON lines ``<local ISO timestamp> <json>`` -> [(local timestamp as UTC ISO with microseconds, message)]."""
    out = []
    for ln in text.splitlines():
        ln = ln.strip()
        if not ln or " " not in ln:
            continue
        ts, payload = ln.split(" ", 1)
        m = re.match(r"^(\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2})(?:\.(\d+))?Z?$", ts)
        if not m:
            continue
        try:
            msg = json.loads(payload)
        except ValueError:
            continue
        frac = (m.group(2) or "")[:6].ljust(6, "0")
        out.append((f"{m.group(1)}.{frac}+00:00", msg))
    return out


def ticker_snapshots(lines: list[tuple[str, dict]]) -> dict[str, dict]:
    """{instrument: {first: (ts, data), last: (ts, data), n}} from ``ticker.<instrument>.<interval>`` messages."""
    out: dict[str, dict] = {}
    for ts, msg in lines:
        p = msg.get("params") or {}
        ch, data = p.get("channel") or "", p.get("data")
        if not ch.startswith("ticker.") or not isinstance(data, dict):
            continue
        name = data.get("instrument_name") or ch.split(".")[1]
        s = out.setdefault(name, {"first": (ts, data), "last": (ts, data), "n": 0})
        s["last"] = (ts, data)
        s["n"] += 1
    return out


def tardis_minute(day: str, symbols: list[str], client: httpx.Client, throttle: Throttle, retries: int = 4,
                  backoff: float = 1.0) -> tuple[Optional[str], dict]:
    """Minute 00:00-00:01 UTC of ``day`` (must be the first of a month) for the ticker channel of ``symbols``."""
    if not day.endswith("-01"):
        raise ValueError("Tardis serves keyless data only for the first day of each month")
    r, info = _get(client, TARDIS_FEED, throttle, params={"from": day, "offset": 0, "filters": tardis_filters(symbols)},
                   headers={"Accept-Encoding": "gzip"}, retries=retries, backoff=backoff)
    if r is not None:
        info["bytes"] = len(r.content)
        info["x_name"] = r.headers.get("x-name")
        return r.text, info
    return None, info


def box_rows(day: str, legs: list[dict], snaps: dict[str, dict], spot: Optional[float], observed_at: str, status: str) -> list[dict]:
    meta = {name: (L, parse_instrument(name) or {}) for L in legs for name in L["instruments"]}
    rows = []
    for name in sorted(snaps):
        if name not in meta:
            continue
        L, p = meta[name]
        for snap in ("first", "last"):
            ts, data = snaps[name][snap]
            rows.append({"source": "tardis", "kind": "option-ticker", "key": f"{day}:{name}:{snap}", "observed_at": observed_at,
                         "available_at": _cap(ts, observed_at),
                         "value": {"capture_day": day, "instrument": name, "expiry": L["expiry"], "expiry_code": L["expiry_code"], "dte": L["dte"],
                                   "strike": p.get("strike"), "type": p.get("type"), "snap": snap, "n_messages": snaps[name]["n"],
                                   "local_timestamp": ts, "spot_hint": spot, "ticker": data}})
    rows.append({"source": "tardis", "kind": "capture", "key": day, "observed_at": observed_at, "available_at": _cap(day + "T00:01:00+00:00", observed_at),
                 "value": {"capture_day": day, "spot_hint": spot, "expiries": [L["expiry_code"] for L in legs],
                           "strikes": {L["expiry_code"]: L["strikes"] for L in legs}, "symbols": sum(len(L["instruments"]) for L in legs),
                           "received": sum(1 for n in snaps if n in meta), "status": status}})
    return rows


def fetch_tardis_box_monthly(store=None, days: Optional[Iterable[str]] = None, limit: Optional[int] = None, dry_run: bool = False,
                             client: Optional[httpx.Client] = None, instruments: Optional[list[dict]] = None,
                             spots: Optional[dict[str, float]] = None, dte_min: float = 30, dte_max: float = 90, n_side: int = 2,
                             min_interval: float = 1.0, retries: int = 4, backoff: float = 1.0, report: Optional[dict] = None,
                             skip_stored: bool = False) -> list[dict]:
    """Bundle ``deribit-box-monthly`` (never in git: Tardis licence). Returns the archive rows of this run.

    ``days`` defaults to every first-of-month from 2019-04-01 to today; ``limit`` keeps the first N of them (each capture
    day is fetched whole). ``instruments`` (compact BTC option list) skips the get_instruments download; ``spots``
    ({day: price}) skips the chart call per day. A day whose feed call fails or returns no leg is reported, not stored;
    stored capture days not fetched in this run are kept, and the store also happens when the run is interrupted.
    ``skip_stored`` skips capture days already in the stored bundle (the monthly top-up: only the new first-of-month)."""
    want = list(days) if days is not None else box_capture_days()
    rep = report if report is not None else {}
    archive = _archive(store)
    if archive is not None and skip_stored:
        have = {r["value"].get("capture_day") for r in _load_existing(archive, BOX_BUNDLE) if r["kind"] == "capture"}
        rep["_already_stored"] = sorted(d for d in want if d in have)
        want = [d for d in want if d not in have]
    if limit is not None:
        want = want[:limit]
    if dry_run:
        rep["_plan"] = {"days": len(want), "first": want[0] if want else None, "last": want[-1] if want else None,
                        "calls": len(want) * (1 if spots else 2) + (0 if instruments is not None else 2),
                        "instrument_mb": 0 if instruments is not None else 77, "est_rows": round(len(want) * (1 + 2 * 1.5 * 4 * n_side)),
                        "est_gz_kb": round(len(want) * (1 + 2 * 1.5 * 4 * n_side) * 0.15), "est_seconds": round(len(want) * 2 * (min_interval + 0.5) + 30),
                        "est_tardis_transfer_kb": round(len(want) * 1.5 * 4 * n_side * 24 * 0.8 / 8)}
        return []
    own = client is None
    c = client or _client()
    th, th_feed = Throttle(min_interval), Throttle(min_interval)
    rows: list[dict] = []
    done: set[str] = set()
    try:
        insts = instruments if instruments is not None else deribit_option_instruments("BTC", c, Throttle(1.1), retries=retries, backoff=max(backoff, 1.0) if backoff else 0)
        for day in want:
            spot = (spots or {}).get(day)
            if spot is None:
                spot = spot_hint(day, c, th, retries=retries, backoff=backoff)
            if spot is None:
                rep[day] = {"skipped": "no spot hint"}
                continue
            legs = select_box_legs(insts, day, spot, dte_min, dte_max, n_side)
            symbols = [n for L in legs for n in L["instruments"]]
            if not symbols:
                rep[day] = {"skipped": "no 30-90 DTE expiry with strikes on both sides", "spot": spot}
                continue
            text, info = tardis_minute(day, symbols, c, th_feed, retries=retries, backoff=backoff)
            observed = now_iso()
            snaps = ticker_snapshots(parse_tardis(text)) if text else {}
            got = sum(1 for s in symbols if s in snaps)
            status = "ok" if got == len(symbols) else ("partial" if got else "empty")
            rep[day] = {**info, "spot": spot, "expiries": [L["expiry_code"] for L in legs], "symbols": len(symbols), "received": got, "status": status}
            if text is None and info.get("retry_after"):
                rep["_aborted"] = {"day": day, "retry_after_s": info["retry_after"], "body": info.get("body")}
                break
            if text is None or not got:
                continue
            rows.extend(box_rows(day, legs, snaps, spot, observed, status))
            done.add(day)
    finally:                                                # also on an interrupt: fetched days are kept
        if own:
            c.close()
        if archive is not None and done:
            _store_box_bundle(archive, rows, done, dte_min, dte_max, n_side, min_interval, retries, rep)
    return rows


def _store_box_bundle(archive, rows: list[dict], done: set[str], dte_min: float, dte_max: float, n_side: int, min_interval: float,
                      retries: int, rep: dict) -> None:
    name = BOX_BUNDLE
    old = [r for r in _load_existing(archive, name) if r["value"].get("capture_day") not in done]
    allrows = sorted(old + [r for r in rows if r["value"]["capture_day"] in done],
                     key=lambda r: (r["value"]["capture_day"], r["kind"] != "capture", r["key"]))
    stored_days = sorted({r["value"]["capture_day"] for r in allrows})
    e = archive.store(name, kind="data", rows=allrows,
                      recipe=(f"GET {TARDIS_FEED}?from=<YYYY-MM-01>&offset=0&filters=[{{\"channel\":\"ticker\",\"symbols\":[...]}}] (keyless, first day "
                              f"of each month only; minute 00:00-00:01 UTC) for BTC box legs: every expiry {dte_min}-{dte_max} DTE listed that day "
                              f"(creation_timestamp from {DERIBIT_HISTORY}get_instruments, expired=true and false), {n_side} strikes at/below and {n_side} above the "
                              f"spot hint ({DERIBIT_CHART} BTC-PERPETUAL 1h bar ending 00:00 UTC) with both call and put listed; first and last ticker "
                              f"message per leg kept raw. via mandel.data.options_misc.fetch_tardis_box_monthly; {min_interval}s pacing, {retries} tries "
                              f"with backoff (Retry-After honoured). available_at = Tardis local receive time. Capture days stored: {stored_days[0]}.."
                              f"{stored_days[-1]} ({len(stored_days)}). {TARDIS_LICENCE}: never in git"),
                      tags=["options", "crypto", "deribit", "tardis", "box", "monthly", "T-260"], schema=BOX_TICKER_SCHEMA, in_git=False)
    rep["stored"] = {"rows": e.rows, "bytes": e.bytes, "days": len(stored_days)}

